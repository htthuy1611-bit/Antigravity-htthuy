import os
import sys
import subprocess
import docx
from docx.shared import Pt, Cm, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

if sys.stdout.encoding.lower() != 'utf-8':
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except Exception:
        pass

# File LaTeX được thiết kế chuẩn xác từng chi tiết theo PDF gốc
tex_content = r"""\documentclass[12pt,a4paper]{article}
\usepackage[utf8]{vietnam}
\usepackage{amsmath,amssymb}
\usepackage{graphicx}
\usepackage{tcolorbox}
\usepackage[top=1.6cm,bottom=1.6cm,left=1.8cm,right=1.5cm]{geometry}

\begin{document}

% 1. HEADER CHUYÊN NGHIỆP LỚP TOÁN CÔ THÚY
\noindent
\begin{tabular}{|p{8.5cm}|p{8.5cm}|}
\hline
\centering \textbf{\large LỚP TOÁN CÔ THÚY} \tabularnewline
\centering \textbf{SĐT: 0935.322.328} \tabularnewline
\centering \textbf{Đ/c: 50/2C Phạm Thị Liên} &
\centering \textbf{\large ĐỀ ÔN TẬP TOÁN} \tabularnewline
\centering \textbf{MÔN: TOÁN -- KHỐI LỚP: 12} \tabularnewline
\centering \textit{Thời gian làm bài: 90 phút (không kể phát đề)} \tabularnewline
\hline
\end{tabular}

\vspace{0.15cm}
\noindent \hfill \fbox{\textbf{MÃ ĐỀ THI 2009}}

\vspace{0.25cm}
% TIÊU ĐỀ PHẦN 1
\noindent\fbox{\textbf{\large PHẦN 1. Câu trắc nghiệm 4 phương án}}

\vspace{0.15cm}
\noindent\textit{Thí sinh trả lời từ câu 1 đến câu 12. Mỗi câu hỏi thí sinh chỉ chọn một phương án.}

\vspace{0.25cm}
% CÂU 1 (4 dòng riêng biệt cho A, B, C, D)
\textbf{Câu 1.} Cho hàm số $y = f(x)$ có đạo hàm trên $\mathbb{R}$ thỏa $f'(x) < 0$, $\forall x \in (1; 2)$ và $f'(x) > 0$, $\forall x \in (2; 3)$. Phát biểu nào sau đây là đúng?

\vspace{0.1cm}
\noindent\textbf{A.} Hàm số $y = f(x)$ đồng biến trên cả hai khoảng $(1; 2)$ và $(2; 3)$.

\vspace{0.1cm}
\noindent\textbf{B.} Hàm số $y = f(x)$ nghịch biến trên cả hai khoảng $(1; 2)$ và $(2; 3)$.

\vspace{0.1cm}
\noindent\textbf{C.} Hàm số $y = f(x)$ đồng biến trên khoảng $(1; 2)$ và nghịch biến trên khoảng $(2; 3)$.

\vspace{0.1cm}
\noindent\textbf{D.} Hàm số $y = f(x)$ nghịch biến trên khoảng $(1; 2)$ và đồng biến trên khoảng $(2; 3)$.

\vspace{0.25cm}
% CÂU 2 (1 dòng 4 phương án căn cột đều tăm tắp)
\textbf{Câu 2.} Giá trị cực đại của hàm số $f(x) = 2x^3 - 9x^2 - 24x + 1$ là

\vspace{0.1cm}
\noindent
\begin{tabular}{@{}p{4.2cm}p{4.2cm}p{4.2cm}p{4.2cm}@{}}
\textbf{A.} $-1$. & \textbf{B.} $14$. & \textbf{C.} $4$. & \textbf{D.} $-111$.
\end{tabular}

\vspace{0.25cm}
% CÂU 3 (2 dòng, mỗi dòng 2 phương án)
\textbf{Câu 3.} Cho lăng trụ $ABC.A'B'C'$. Khẳng định nào sau đây đúng?

\vspace{0.1cm}
\noindent
\begin{tabular}{@{}p{8.5cm}p{8.5cm}@{}}
\textbf{A.} $\overrightarrow{BA} + \overrightarrow{A'C'} = \overrightarrow{BC}$. & \textbf{B.} $\overrightarrow{BA} + \overrightarrow{A'C'} = \overrightarrow{BC'}$. \tabularnewline
\textbf{C.} $\overrightarrow{BA} + \overrightarrow{A'C'} = \overrightarrow{C'B}$. & \textbf{D.} $\overrightarrow{BA} + \overrightarrow{A'C'} = \overrightarrow{B'C}$.
\end{tabular}

\vspace{0.25cm}
% CÂU 4 (Bên trái: Đề bài + 4 phương án, Bên phải: Hình ảnh)
\noindent
\begin{tabular}{@{}p{11.8cm}p{5.2cm}@{}}
\textbf{Câu 4.} Hàm số $y = f(x)$ xác định trên đoạn $[-1; 6]$ và có đồ thị như hình vẽ. Hàm số đã cho nghịch biến trên khoảng nào sau đây?

\vspace{0.15cm}
\begin{tabular}{@{}p{5.8cm}p{5.8cm}@{}}
\textbf{A.} $(-1; 2)$. & \textbf{B.} $(0; 2)$. \tabularnewline
\textbf{C.} $(2; 6)$. & \textbf{D.} $(-2; 0)$.
\end{tabular}
&
\centering \includegraphics[width=4.6cm]{fig_cau4.png}
\end{tabular}

\vspace{0.25cm}
% CÂU 5 (2 dòng x 2 phương án)
\textbf{Câu 5.} Cho tứ diện $ABCD$. Lấy $G$ là trọng tâm của tam giác $ABC$. Phát biểu nào sau đây là sai?

\vspace{0.1cm}
\noindent
\begin{tabular}{@{}p{8.5cm}p{8.5cm}@{}}
\textbf{A.} $\overrightarrow{GA} + \overrightarrow{GB} + \overrightarrow{GC} = \overrightarrow{0}$. & \textbf{B.} $\overrightarrow{GA} + \overrightarrow{GB} + \overrightarrow{GC} + \overrightarrow{GD} = \overrightarrow{0}$. \tabularnewline
\textbf{C.} $\overrightarrow{GD} - \overrightarrow{GA} = \overrightarrow{AD}$. & \textbf{D.} $\overrightarrow{DA} + \overrightarrow{DB} + \overrightarrow{DC} = 3\overrightarrow{DG}$.
\end{tabular}

\vspace{0.25cm}
% CÂU 6 (1 dòng 4 phương án)
\textbf{Câu 6.} Một chiếc hộp hình lập phương $ABCD.A'B'C'D'$ có cạnh bằng $12\text{ cm}$, mặt trên $A'B'C'D'$ không nắp. Có một con kiến ở đỉnh $A$ bên ngoài hộp và một miếng mồi của kiến tại điểm $O$ là tâm đáy $ABCD$ ở bên trong hộp. Quãng đường ngắn nhất mà con kiến tìm đến miếng mồi (làm tròn đến hai chữ số thập phân) là

\vspace{0.1cm}
\noindent
\begin{tabular}{@{}p{4.2cm}p{4.2cm}p{4.2cm}p{4.2cm}@{}}
\textbf{A.} $32{,}49\text{ (cm)}$. & \textbf{B.} $36{,}29\text{ (cm)}$. & \textbf{C.} $12\text{ (cm)}$. & \textbf{D.} $30{,}59\text{ (cm)}$.
\end{tabular}

\vspace{0.25cm}
% CÂU 7 (1 dòng 4 phương án)
\textbf{Câu 7.} Giá trị nhỏ nhất của hàm số $f(x) = x^4 - 8x^2 + a$, ($a \in \mathbb{R}$) trên đoạn $[-1; 3]$ bằng

\vspace{0.1cm}
\noindent
\begin{tabular}{@{}p{4.2cm}p{4.2cm}p{4.2cm}p{4.2cm}@{}}
\textbf{A.} $-6$. & \textbf{B.} $a$. & \textbf{C.} $-16 + a$. & \textbf{D.} $9 + a$.
\end{tabular}

\vspace{0.25cm}
% CÂU 8 (Bên trái đề + phương án, Bên phải hình)
\noindent
\begin{tabular}{@{}p{11.8cm}p{5.2cm}@{}}
\textbf{Câu 8.} Hàm số $y = f(x)$ xác định trên đoạn $[-1; 5]$ và có đồ thị như hình vẽ. Tập giá trị của hàm số $y = f(x)$ trên đoạn $[-1; 5]$ là

\vspace{0.15cm}
\begin{tabular}{@{}p{5.8cm}p{5.8cm}@{}}
\textbf{A.} $[-1; 5]$. & \textbf{B.} $[1; 3]$. \tabularnewline
\textbf{C.} $[-1; 3]$. & \textbf{D.} $[1; 5]$.
\end{tabular}
&
\centering \includegraphics[width=4.6cm]{fig_cau8.png}
\end{tabular}

\newpage
% TRANG 2
% CÂU 9
\textbf{Câu 9.} Mỗi ngày bác Bình đều đi bộ để rèn luyện sức khoẻ. Quãng đường đi bộ mỗi ngày (đơn vị: $\text{km}$) của bác Bình trong 20 ngày được thống kê lại ở bảng sau:

\begin{center}
\begin{tabular}{|l|c|c|c|c|c|}
\hline
\textbf{Quãng đường} & $[2{,}7; 3{,}0)$ & $[3{,}0; 3{,}3)$ & $[3{,}3; 3{,}6)$ & $[3{,}6; 3{,}9)$ & $[3{,}9; 4{,}2)$ \\
\hline
\textbf{Số ngày} & $3$ & $6$ & $5$ & $4$ & $2$ \\
\hline
\end{tabular}
\end{center}

Khoảng biến thiên của mẫu số liệu ghép nhóm trên bằng

\vspace{0.1cm}
\noindent
\begin{tabular}{@{}p{4.2cm}p{4.2cm}p{4.2cm}p{4.2cm}@{}}
\textbf{A.} $1{,}2$. & \textbf{B.} $0{,}362$. & \textbf{C.} $3{,}39$. & \textbf{D.} $1{,}5$.
\end{tabular}

\vspace{0.25cm}
% CÂU 10
\textbf{Câu 10.} Trong không gian $Oxyz$, cho hai điểm $A(1; 1; 2)$ và $B(3; 1; 0)$. Trung điểm của đoạn thẳng $AB$ có toạ độ là

\vspace{0.1cm}
\noindent
\begin{tabular}{@{}p{4.2cm}p{4.2cm}p{4.2cm}p{4.2cm}@{}}
\textbf{A.} $(2; 1; 1)$. & \textbf{B.} $(4; 2; 2)$. & \textbf{C.} $(2; 0; -2)$. & \textbf{D.} $(1; 0; -1)$.
\end{tabular}

\vspace{0.25cm}
% CÂU 11
\textbf{Câu 11.} Đường tiệm cận xiên của đồ thị hàm số $y = \dfrac{x^2 + 2x - 2}{x - 2}$ là

\vspace{0.1cm}
\noindent
\begin{tabular}{@{}p{4.2cm}p{4.2cm}p{4.2cm}p{4.2cm}@{}}
\textbf{A.} $y = -x + 3$. & \textbf{B.} $y = x + 3$. & \textbf{C.} $y = x - 3$. & \textbf{D.} $y = x + 4$.
\end{tabular}

\vspace{0.25cm}
% CÂU 12
\textbf{Câu 12.} Trong không gian $Oxyz$, cho điểm $A(-3; 1; -4)$, $B(1; -5; 2)$. Đường thẳng $AB$ cắt mặt phẳng $(Oxy)$ tại điểm

\vspace{0.1cm}
\noindent
\begin{tabular}{@{}p{8.5cm}p{8.5cm}@{}}
\textbf{A.} $M\left(-\dfrac{1}{3}; -3; 0\right)$. & \textbf{B.} $N\left(\dfrac{1}{3}; 3; 0\right)$. \tabularnewline
\textbf{C.} $P(0; 3; 1)$. & \textbf{D.} $Q(-3; 1; 0)$.
\end{tabular}

\vspace{0.35cm}
% TIÊU ĐỀ PHẦN 2
\noindent\fbox{\textbf{\large PHẦN 2. Câu trắc nghiệm đúng sai}}

\vspace{0.15cm}
\noindent\textit{Thí sinh trả lời từ câu 1 đến câu 4. Trong mỗi ý a), b), c), d) ở mỗi câu, thí sinh chọn đúng hoặc sai.}

\vspace{0.25cm}
% CÂU 1 PHẦN 2
\textbf{Câu 1.} Hàm số $y = f(x)$ liên tục trên $\mathbb{R}$ và có bảng biến thiên như sau:

\begin{center}
\includegraphics[width=11cm]{fig_bbt.png}
\end{center}

\noindent\textbf{a)} Đồ thị hàm số đã cho có hai đường tiệm cận ngang.

\vspace{0.08cm}
\noindent\textbf{b)} Giá trị nhỏ nhất của hàm số trên $(-\infty; +\infty)$ bằng $8$.

\vspace{0.08cm}
\noindent\textbf{c)} Hàm số đồng biến trên $(8; 38)$.

\vspace{0.08cm}
\noindent\textbf{d)} Giá trị lớn nhất của hàm số trên $\mathbb{R}$ bằng $142$.

\vspace{0.25cm}
% CÂU 2 PHẦN 2 (Tam giác)
\noindent
\begin{tabular}{@{}p{12.2cm}p{4.8cm}@{}}
\textbf{Câu 2.} Xét tam giác $ABC$ có $AC = 2AB$ và $BC = 10\text{ cm}$. Trên cạnh $AC$ lấy điểm $D$ sao cho $AD = \dfrac{1}{4}AC$, trên cạnh $AB$ lấy điểm $E$ sao cho $AE = \dfrac{1}{4}AB$, trên cạnh $AD$ lấy điểm $F$ sao cho $AF = \dfrac{1}{4}AD$ và tiếp tục lấy các điểm $G, H, I, J\dots$ (vô hạn lần) theo quy luật đó. Xét tính đúng sai các mệnh đề sau:

\vspace{0.1cm}
\noindent\textbf{a)} $\dfrac{AB}{AC} = \dfrac{AD}{AB}$.

\vspace{0.08cm}
\noindent\textbf{b)} Tam giác $ABD$ đồng dạng với tam giác $ABC$.

\vspace{0.08cm}
\noindent\textbf{c)} $BD = 5\text{ cm}$; $DE = 3\text{ cm}$.

\vspace{0.08cm}
\noindent\textbf{d)} Độ dài đường gấp khúc $CBDEFGH\dots$ bằng $20\text{ cm}$. &
\centering \includegraphics[width=4.4cm]{fig_tamgiac.png}
\end{tabular}

\newpage
% TRANG 3
% CÂU 3 PHẦN 2 (Máy bay)
\noindent
\begin{tabular}{@{}p{11.8cm}p{5.2cm}@{}}
\textbf{Câu 3.} Hình vẽ sau mô tả vị trí của máy bay vào thời điểm 9h30 phút. Biết các đơn vị trên hình tính theo đơn vị $\text{km}$. Trong các khẳng định sau đây, khẳng định nào đúng, khẳng định nào sai?

\vspace{0.1cm}
\noindent\textbf{a)} Máy bay đang ở độ cao $9\text{ km}$.

\vspace{0.08cm}
\noindent\textbf{b)} Tọa độ của máy bay $(300; 150; 9)$.

\vspace{0.08cm}
\noindent\textbf{c)} Phi công để máy bay ở chế độ tự động với vận tốc theo hướng đông là $750\text{ km/h}$, độ cao không đổi. Biết rằng gió thổi theo hướng đông với vận tốc $10\text{ m/s}$. Giả sử vận tốc và hướng gió không đổi thì lúc 10h30 phút máy bay ở tọa độ $(150; 1086; 9)$.

\vspace{0.08cm}
\noindent\textbf{d)} Sau khi bay đến vị trí lúc 10h30 thì máy bay bay theo hướng ngược lại với vận tốc $800\text{ km/h}$ với độ cao không đổi, biết lúc đó trời lặng gió thì lúc 11h máy bay ở tọa độ $(686; 150; 9)$. &
\centering \includegraphics[width=5cm]{fig_maybay.png}
\end{tabular}

\vspace{0.25cm}
% CÂU 4 PHẦN 2
\textbf{Câu 4.} Khảo sát chiều cao của 20 học sinh nam lớp 12A của một trường THPT X, người ta được kết quả thống kê trong bảng sau:

\begin{center}
\begin{tabular}{|l|c|c|c|c|c|}
\hline
\textbf{Chiều cao (cm)} & $[160; 165)$ & $[165; 170)$ & $[170; 175)$ & $[175; 180)$ & $[180; 185)$ \\
\hline
\textbf{Số học sinh} & $3$ & $5$ & $7$ & $4$ & $1$ \\
\hline
\end{tabular}
\end{center}

\noindent\textbf{a)} Gọi $x_1; x_2; \dots; x_{20}$ là mẫu số liệu gốc gồm chiều cao của 20 học sinh trên được xếp theo thứ tự không giảm. Khi đó, $x_3 \in [165; 170)$ và $x_9 \in [170; 175)$.

\vspace{0.08cm}
\noindent\textbf{b)} Tứ phân vị thứ ba của mẫu số liệu ghép nhóm đã cho bằng $175$.

\vspace{0.08cm}
\noindent\textbf{c)} Khoảng tứ phân vị của mẫu số liệu ghép nhóm đã cho là $\Delta Q = Q_3 - Q_1 = 8{,}5$.

\vspace{0.08cm}
\noindent\textbf{d)} Chọn ngẫu nhiên một học sinh trong nhóm khảo sát nói trên, xác suất chọn được học sinh có chiều cao từ $175\text{ cm}$ trở lên bằng $0{,}25$.

\vspace{0.35cm}
% TIÊU ĐỀ PHẦN 3
\noindent\fbox{\textbf{\large PHẦN 3. Câu trắc nghiệm trả lời ngắn}}

\vspace{0.15cm}
\noindent\textit{Thí sinh trả lời từ câu 1 đến câu 6.}

\vspace{0.25cm}
% CÂU 1 PHẦN 3
\textbf{Câu 1.} Một doanh nghiệp dự kiến sản xuất không quá $2000$ sản phẩm cùng loại. Giả sử rằng doanh thu của doanh nghiệp khi sản xuất $x$ sản phẩm ($x \in \mathbb{N}$; $0 \leqslant x \leqslant 2000$) là $R(x) = 16x - 0{,}005x^2$ (triệu đồng) và doanh nghiệp phải nộp một khoản thuế là $5\%$ của doanh thu $R(x)$. Chi phí để sản xuất mỗi sản phẩm là bốn triệu đồng. Lợi nhuận của doanh nghiệp (khi sản xuất $x$ sản phẩm) được tính theo công thức: Lợi nhuận bằng doanh thu trừ đi thuế và chi phí. Doanh nghiệp phải sản xuất bao nhiêu sản phẩm để thu được lợi nhuận lớn nhất?

\vspace{0.1cm}
\noindent\textbf{KQ:} \begin{tabular}{|p{0.55cm}|p{0.55cm}|p{0.55cm}|p{0.55cm}|} \hline \phantom{0} & \phantom{0} & \phantom{0} & \phantom{0} \tabularnewline \hline \end{tabular}

\vspace{0.25cm}
% CÂU 2 PHẦN 3 (Shipper)
\noindent
\begin{tabular}{@{}p{11.8cm}p{5.2cm}@{}}
\textbf{Câu 2.} Một bác Shipper giao hàng xuất phát từ kho $A$ để lấy hàng và đi giao tất cả các con đường sau đó lại trở về kho $A$ để trả lại những hàng hóa mà khách hàng chưa nhận. Con đường có sơ đồ và thời gian giao hàng (phút) trên mỗi con đường được mô tả trong hình sau. Thời gian ngắn nhất để bác Shipper hoàn thành công việc trên là bao nhiêu phút?

\vspace{0.15cm}
\noindent\textbf{KQ:} \begin{tabular}{|p{0.55cm}|p{0.55cm}|p{0.55cm}|p{0.55cm}|} \hline \phantom{0} & \phantom{0} & \phantom{0} & \phantom{0} \tabularnewline \hline \end{tabular} &
\centering \includegraphics[width=5cm]{fig_shipper.png}
\end{tabular}

\vspace{0.25cm}
% CÂU 3 PHẦN 3
\textbf{Câu 3.} Trong không gian $Oxyz$, cho tam giác $ABC$ có $A(-4; -1; 2)$, $B(3; 5; -6)$ và $C(a; b; c)$. Biết trung điểm cạnh $AC$ thuộc trục tung, trung điểm cạnh $BC$ thuộc mặt phẳng $(Oxz)$. Tính $T = 2a + b - c$.

\vspace{0.1cm}
\noindent\textbf{KQ:} \begin{tabular}{|p{0.55cm}|p{0.55cm}|p{0.55cm}|p{0.55cm}|} \hline \phantom{0} & \phantom{0} & \phantom{0} & \phantom{0} \tabularnewline \hline \end{tabular}

\newpage
% TRANG 4
% CÂU 4 PHẦN 3
\textbf{Câu 4.} Cho tập hợp $X = \{1; 2; 3; 4; 5; 6; 7; 8\}$. Gọi $S$ là tập hợp tất cả các số tự nhiên có 4 chữ số được lập từ các chữ số thuộc tập $X$. Chọn ngẫu nhiên một số từ tập hợp $S$. Xác suất để chọn được một số chia hết cho 3 bằng $\dfrac{a}{b}$ (với $a, b \in \mathbb{N}^*$, $\dfrac{a}{b}$ là phân số tối giản). Tính $T = a + b$.

\vspace{0.1cm}
\noindent\textbf{KQ:} \begin{tabular}{|p{0.55cm}|p{0.55cm}|p{0.55cm}|p{0.55cm}|} \hline \phantom{0} & \phantom{0} & \phantom{0} & \phantom{0} \tabularnewline \hline \end{tabular}

\vspace{0.25cm}
% CÂU 5 PHẦN 3
\textbf{Câu 5.} Một quần thể vi khuẩn được nuôi cấy trong phòng thí nghiệm. Nồng độ dinh dưỡng $S$ (đơn vị: $\text{mg/}\ell$) thay đổi theo thời gian $t$ giờ ($t \geqslant 0$) được mô hình hóa bởi hàm số: $S(t) = \dfrac{10t + 5}{t + 1}$. Biết tốc độ sinh trưởng $V$ của vi khuẩn phụ thuộc vào nồng độ dinh dưỡng theo hàm số $V(S) = \dfrac{5S}{S + 2}$. Khi thời gian $t$ kéo dài, tốc độ sinh trưởng $V$ tăng dần và ổn định quanh một ngưỡng $K$ nhất định. Hỏi sau bao nhiêu phút thì tốc độ sinh trưởng của vi khuẩn đạt $90\%$ ngưỡng $K$?

\vspace{0.1cm}
\noindent\textbf{KQ:} \begin{tabular}{|p{0.55cm}|p{0.55cm}|p{0.55cm}|p{0.55cm}|} \hline \phantom{0} & \phantom{0} & \phantom{0} & \phantom{0} \tabularnewline \hline \end{tabular}

\vspace{0.25cm}
% CÂU 6 PHẦN 3
\textbf{Câu 6.} Trong không gian với hệ trục tọa độ $Oxyz$, mặt phẳng $(P)\colon bcx + acy + abz - abc = 0$ qua điểm $M(2; 4; 8)$ và cắt các tia $Ox, Oy, Oz$ lần lượt tại $A, B, C$ sao cho $OA = 2OB = 4OC$. Tính $T = a + b + c$.

\vspace{0.1cm}
\noindent\textbf{KQ:} \begin{tabular}{|p{0.55cm}|p{0.55cm}|p{0.55cm}|p{0.55cm}|} \hline \phantom{0} & \phantom{0} & \phantom{0} & \phantom{0} \tabularnewline \hline \end{tabular}

\vspace{0.6cm}
\begin{center}
\textbf{---------- HẾT ----------}
\end{center}

\end{document}
"""

with open("De_Thi_2009_perfect.tex", "w", encoding="utf-8") as f:
    f.write(tex_content)

pandoc_exe = os.path.expandvars(r"%LOCALAPPDATA%\Pandoc\pandoc.exe")
cmd = [pandoc_exe, "De_Thi_2009_perfect.tex", "-o", "De_Thi_2009_temp.docx"]
subprocess.run(cmd, check=True)

# Post-processing với python-docx:
doc = docx.Document("De_Thi_2009_temp.docx")

# 1. Khổ giấy và Căn lề A4 chuẩn Quốc gia
for s in doc.sections:
    s.top_margin = Cm(1.6)
    s.bottom_margin = Cm(1.6)
    s.left_margin = Cm(2.0)
    s.right_margin = Cm(1.5)
    s.page_width = Cm(21.0)
    s.page_height = Cm(29.7)
    
    # Footer
    footer = s.footer
    f_p = footer.paragraphs[0]
    f_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    f_p.text = ""
    run_l = f_p.add_run("Lớp Toán Cô Thúy  •  ĐT: 0935.322.328  •  50/2C Phạm Thị Liên                                   ")
    run_l.font.name = "Times New Roman"
    run_l.font.size = Pt(9.5)
    run_l.italic = True
    run_l.font.color.rgb = RGBColor(100, 100, 100)
    
    run_r = f_p.add_run("Trang Mã đề 2009")
    run_r.font.name = "Times New Roman"
    run_r.font.size = Pt(9.5)
    run_r.italic = True
    run_r.font.color.rgb = RGBColor(80, 80, 80)

# 2. Chuẩn hoá Style Normal
style = doc.styles['Normal']
style.font.name = 'Times New Roman'
style.font.size = Pt(12)
style.font.color.rgb = RGBColor(0, 0, 0)
rPr = style.element.get_or_add_rPr()
rFonts = parse_xml(r'<w:rFonts %s w:ascii="Times New Roman" w:hAnsi="Times New Roman" w:cs="Times New Roman"/>' % nsdecls('w'))
rPr.append(rFonts)

def format_paragraph(p):
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(2.5)
    p.paragraph_format.space_before = Pt(0)
    
    text = p.text.strip()
    # CĂN ĐỀU 2 BÊN (JUSTIFY) CHO CÂU HỎI VÀ CÁC MỆNH ĐỀ ĐÚNG SAI
    if text.startswith("Câu ") or text.startswith("a)") or text.startswith("b)") or text.startswith("c)") or text.startswith("d)"):
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    elif text.startswith("PHẦN"):
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    elif "HẾT" in text or "MÃ ĐỀ" in text:
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif text.startswith("A.") or text.startswith("B.") or text.startswith("C.") or text.startswith("D."):
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    # Đảm bảo font Times New Roman 12pt
    for r in p.runs:
        r.font.name = 'Times New Roman'
        r.font.size = Pt(12)
        r_rPr = r._r.get_or_add_rPr()
        f = parse_xml(r'<w:rFonts %s w:ascii="Times New Roman" w:hAnsi="Times New Roman" w:cs="Times New Roman"/>' % nsdecls('w'))
        r_rPr.append(f)

for p in doc.paragraphs:
    format_paragraph(p)

# 3. Xử lý các bảng
is_header_table = True
for table in doc.tables:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    is_data_table = False
    is_answer_box = False
    for row in table.rows:
        for cell in row.cells:
            text = cell.text.strip()
            if 'Quãng đường' in text or 'Số ngày' in text or 'Chiều cao' in text or 'Số học sinh' in text:
                is_data_table = True
            if len(row.cells) == 4 and all(c.text.strip() == '' for c in row.cells):
                is_answer_box = True

    if is_header_table:
        # Bảng Header: viền xanh navy sang trọng
        tblPr = table._tbl.tblPr
        tblBorders = parse_xml(
            r'<w:tblBorders %s>'
            r'<w:top w:val="single" w:sz="12" w:space="0" w:color="1F497D"/>'
            r'<w:left w:val="single" w:sz="12" w:space="0" w:color="1F497D"/>'
            r'<w:bottom w:val="single" w:sz="12" w:space="0" w:color="1F497D"/>'
            r'<w:right w:val="single" w:sz="12" w:space="0" w:color="1F497D"/>'
            r'<w:insideH w:val="none"/>'
            r'<w:insideV w:val="single" w:sz="6" w:space="0" w:color="1F497D"/>'
            r'</w:tblBorders>' % nsdecls('w')
        )
        tblPr.append(tblBorders)
        is_header_table = False
    elif is_answer_box:
        # Ô trả lời ngắn: giữ viền đen chuẩn
        tblPr = table._tbl.tblPr
        tblBorders = parse_xml(
            r'<w:tblBorders %s>'
            r'<w:top w:val="single" w:sz="8" w:space="0" w:color="000000"/>'
            r'<w:left w:val="single" w:sz="8" w:space="0" w:color="000000"/>'
            r'<w:bottom w:val="single" w:sz="8" w:space="0" w:color="000000"/>'
            r'<w:right w:val="single" w:sz="8" w:space="0" w:color="000000"/>'
            r'<w:insideH w:val="single" w:sz="8" w:space="0" w:color="000000"/>'
            r'<w:insideV w:val="single" w:sz="8" w:space="0" w:color="000000"/>'
            r'</w:tblBorders>' % nsdecls('w')
        )
        tblPr.append(tblBorders)
    elif not is_data_table:
        # Xóa viền cho bảng layout ảnh và bảng phương án A, B, C, D
        tblPr = table._tbl.tblPr
        tblBorders = parse_xml(
            r'<w:tblBorders %s>'
            r'<w:top w:val="none"/>'
            r'<w:left w:val="none"/>'
            r'<w:bottom w:val="none"/>'
            r'<w:right w:val="none"/>'
            r'<w:insideH w:val="none"/>'
            r'<w:insideV w:val="none"/>'
            r'</w:tblBorders>' % nsdecls('w')
        )
        tblPr.append(tblBorders)

    # Format text trong cell
    for row in table.rows:
        for cell in row.cells:
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            for p in cell.paragraphs:
                format_paragraph(p)

# Tắt Compatibility Mode
settings = doc.settings.element
compat = parse_xml(
    r'<w:compat %s>'
    r'<w:compatSetting w:name="compatibilityMode" w:uri="http://schemas.microsoft.com/office/word" w:val="15"/>'
    r'</w:compat>' % nsdecls('w')
)
settings.append(compat)

output_file = "De_Thi_2009.docx"
doc.save(output_file)
print(f"[HOÀN TẤT 100%] Đã lưu file Word chuẩn in ấn: {output_file}")
