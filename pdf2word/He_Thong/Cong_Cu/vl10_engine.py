# -*- coding: utf-8 -*-
"""
BUILD ĐỀ KIỂM TRA GIỮA KỲ I - VẬT LÍ 10 - THPT HAI BÀ TRƯNG (HUẾ) - MÃ ĐỀ 209
Bản quyền: LỚP TOÁN CÔ THÚY - GV: HỒ THỊ THÚY - SĐT: 0935.322.328

Một nguồn dữ liệu duy nhất (QUESTIONS) -> sinh ra:
  1. LaTeX Đề  (ex_test [dethi])   -> pdflatex -> PDF
  2. LaTeX HDG (ex_test [loigiai]) -> pdflatex -> PDF
  3. Word Đề / HDG (Pandoc -> OMML 100% + python-docx hậu xử lý chuẩn BTPro)
"""
import os
import re
import sys
import shutil
import subprocess
import docx
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

if sys.stdout.encoding.lower() != 'utf-8':
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except Exception:
        pass

CURRENT_DIR = os.path.dirname(os.path.abspath(__file__))
BASE_DIR = os.path.abspath(os.path.join(CURRENT_DIR, "..", ".."))
FIG_DIR = os.path.join(BASE_DIR, "He_Thong", "Hinh_Anh", "HBT_10")
LATEX_DIR = os.path.join(BASE_DIR, "He_Thong", "LaTeX", "HBT_10")
SCRATCH_DIR = os.path.join(CURRENT_DIR, "scratch_hbt_10")
SAN_PHAM_DIR = os.path.join(BASE_DIR, "San_Pham")
PANDOC = os.path.expandvars(r"%LOCALAPPDATA%\Pandoc\pandoc.exe")
for d in (LATEX_DIR, SCRATCH_DIR, SAN_PHAM_DIR):
    os.makedirs(d, exist_ok=True)

MADE = "209"
NAME_DE = f"VL10_HBT_De_{MADE}"
NAME_HDG = f"VL10_HBT_HDG_{MADE}"
SCHOOL = "THPT HAI BÀ TRƯNG"
YEAR = "2025 -- 2026"
YEAR_W = "2025 – 2026"
BRAND = "Lớp Toán Cô Thúy -- SĐT: 0935.322.328 -- 50/2C Phạm Thị Liên"
BRAND_W = "Lớp Toán Cô Thúy  •  SĐT: 0935.322.328  •  50/2C Phạm Thị Liên"
LETTERS = "ABCD"
FOLDER = "HBT_10"

# ============================================================================
# DỮ LIỆU ĐỀ THI
# ============================================================================
P1, P2, P3, P4, SECTIONS = [], [], [], [], []

# ============================================================================
# 1. SINH LATEX (ex_test)
# ============================================================================
def tex_fig(fig):
    name, w = fig
    return "\t\\begin{center}\n\t\t\\includegraphics[width=%.1fcm]{%s.pdf}\n\t\\end{center}\n" % (w, name)


def tex_sol(lines):
    out = []
    for ln in lines:
        if ln.startswith(r"\["):
            out.append("\t\t" + ln)
        else:
            out.append("\t\t" + ln + "\\par")
    return "\n".join(out)


def tex_section(idx):
    title, note = SECTIONS[idx]
    return ("\n\\vspace{0.15cm}\n\\begin{flushleft}\n"
            "\t{\\color{blue!50!green}\\fbox{\\fontfamily{qag}\\bfseries\\selectfont %s}}\n"
            "\t\\textit{\\small (%s)}\n\\end{flushleft}\n\\setcounter{ex}{0}\n\\setcounter{bt}{0}\n" % (title, note))


def tex_header(is_sol):
    if is_sol:
        top = r"{\bfseries\color{red!80!black} HƯỚNG DẪN GIẢI CHI TIẾT GIỮA KỲ I}\\[3pt]" + "\n" + \
              r"\textbf{Môn: VẬT LÍ 10 -- %s}\\[2pt]" % SCHOOL + "\n" + \
              r"\textit{Năm học %s -- Mã đề: %s}" % (YEAR, MADE)
        row2 = (r"\rule{0pt}{14pt}\textbf{Giáo viên:} HỒ THỊ THÚY &" + "\n" +
                r"\rule{0pt}{14pt}\textbf{Môn học:} VẬT LÍ 10 &" + "\n" +
                r"\rule{0pt}{14pt}\textbf{Mã đề thi %s}\rule[-4pt]{0pt}{4pt} \\" % MADE)
    else:
        top = r"{\bfseries\color{red!80!black} KIỂM TRA GIỮA KỲ I -- NĂM HỌC %s}\\[3pt]" % YEAR + "\n" + \
              r"\textbf{Môn: VẬT LÍ, Lớp 10 -- %s}\\[2pt]" % SCHOOL + "\n" + \
              r"\textit{Thời gian làm bài: 45 phút (Không kể thời gian phát đề)}"
        row2 = (r"\rule{0pt}{13pt}Họ và tên thí sinh: \dotfill\rule[-3pt]{0pt}{3pt} &" + "\n" +
                r"\rule{0pt}{13pt}Số báo danh: \dotfill\rule[-3pt]{0pt}{3pt} &" + "\n" +
                r"\rule{0pt}{13pt}\textbf{Mã đề thi %s}\rule[-3pt]{0pt}{3pt} \\" % MADE)
    return r"""\noindent
\begin{tabular}{|p{6.6cm}|p{6.6cm}|>{\centering\arraybackslash}p{3.2cm}|}
\hline
\begin{minipage}{6.6cm}
\vspace{3pt}
\centering
{\large\bfseries\color{blue!80!black} LỚP TOÁN CÔ THÚY}\\[3pt]
\textbf{SĐT:} 0935.322.328\\[2pt]
\textbf{Địa chỉ:} 50/2C Phạm Thị Liên
\vspace{3pt}
\end{minipage}
&
\multicolumn{2}{p{10.2cm}|}{
\begin{minipage}{10.2cm}
\vspace{3pt}
\centering
""" + top + r"""
\vspace{3pt}
\end{minipage}
} \\
\hline
""" + row2 + r"""
\hline
\end{tabular}
\vspace{0.1cm}
"""


def tex_answer_key():
    s = "\n\\vspace{0.3cm}\n\\noindent\\begin{minipage}{\\linewidth}\n\\begin{center}{\\large\\bfseries\\color{red!80!black} BẢNG ĐÁP ÁN -- MÃ ĐỀ %s}\\end{center}\n" % MADE
    s += "\\noindent\\textbf{Phần I.}\\par\\smallskip\n\\begin{center}\\begin{tabular}{|c|" + "c|" * len(P1) + "}\\hline\n"
    s += "\\textbf{Câu} & " + " & ".join(str(i + 1) for i in range(len(P1))) + " \\\\ \\hline\n"
    s += "\\textbf{Đáp án} & " + " & ".join("\\textbf{%s}" % LETTERS[q["ans"]] for q in P1) + " \\\\ \\hline\n\\end{tabular}\\end{center}\n"
    s += "\\noindent\\textbf{Phần II.}\\par\\smallskip\n\\begin{center}\\begin{tabular}{|c|c|c|c|c|}\\hline\n"
    s += "\\textbf{Câu} & \\textbf{a)} & \\textbf{b)} & \\textbf{c)} & \\textbf{d)} \\\\ \\hline\n"
    for i, q in enumerate(P2):
        s += "%d & " % (i + 1) + " & ".join("Đ" if t else "S" for t in q["truth"]) + " \\\\ \\hline\n"
    s += "\\end{tabular}\\end{center}\n"
    s += "\\noindent\\textbf{Phần III.}\\par\\smallskip\n\\begin{center}\\begin{tabular}{|c|" + "c|" * len(P3) + "}\\hline\n"
    s += "\\textbf{Câu} & " + " & ".join(str(i + 1) for i in range(len(P3))) + " \\\\ \\hline\n\\textbf{Đáp án} & " + " & ".join(q["ans"] for q in P3) + " \\\\ \\hline\n\\end{tabular}\\end{center}\n\\end{minipage}\n"
    return s


def build_latex(is_sol):
    opt = "loigiai" if is_sol else "dethi"
    foot_r = (r"Trang \thepage/\pageref{LastPage} -- Hướng dẫn giải Mã đề %s" % MADE) if is_sol \
        else (r"Trang \thepage/\pageref{LastPage} -- Mã đề %s" % MADE)
    s = r"""\documentclass[11pt,a4paper]{article}
\usepackage[utf8]{vietnam}
\usepackage{amsmath,amssymb,mathrsfs}
\usepackage{graphicx}
\usepackage{tikz}
\usepackage{array}
\usepackage[top=1.0cm,bottom=1.8cm,left=1.4cm,right=1.4cm,footskip=0.8cm,headheight=16pt]{geometry}
\usepackage{fancyhdr}
\usepackage{tcolorbox}
\tcbuselibrary{skins}
\usepackage{lastpage}
\usepackage[""" + opt + r"""]{ex_test}
\graphicspath{{../../Hinh_Anh/""" + FOLDER + r"""/}}

\makeatletter
\@ifundefined{c@bt}{\newcounter{bt}}{}
\makeatother

\pagestyle{fancy}
\fancyhf{}
\renewcommand{\headrulewidth}{0pt}
\renewcommand{\footrulewidth}{0.4pt}
\lfoot{\footnotesize\textsl{""" + BRAND + r"""}}
\rfoot{\footnotesize\textsl{""" + foot_r + r"""}}
\setlength{\parskip}{1pt}
\setlength{\parindent}{0pt}

\begin{document}

""" + tex_header(is_sol)

    # Phần I
    s += tex_section(0)
    for i, q in enumerate(P1):
        s += "\n%% Câu %d\n\\begin{ex}\n\t%s\n" % (i + 1, q["stem"])
        if q.get("fig"):
            s += tex_fig(q["fig"])
        s += "\t\\choice\n"
        for k, o in enumerate(q["opts"]):
            s += "\t{%s%s}\n" % ("\\True " if k == q["ans"] else "", o.rstrip("."))
        s += "\t\\loigiai{\n%s\n\t}\n\\end{ex}\n" % tex_sol(q["sol"])
    # Phần II
    s += tex_section(1)
    for i, q in enumerate(P2):
        s += "\n%% Phần II Câu %d\n\\begin{ex}\n\t%s\n" % (i + 1, q["stem"])
        if q.get("fig"):
            s += tex_fig(q["fig"])
        s += "\t\\choiceTF[t]\n"
        for t, it in zip(q["truth"], q["items"]):
            s += "\t{%s%s}\n" % ("\\True " if t else "", it)
        lines = []
        for k, (t, so) in enumerate(zip(q["truth"], q["sols"])):
            tag = "Đúng" if t else "Sai"
            col = "green!50!black" if t else "red!80!black"
            lines.append(r"{\color{%s}\textbf{%s) %s.}} %s" % (col, "abcd"[k], tag, so))
        s += "\t\\loigiai{\n%s\n\t}\n\\end{ex}\n" % tex_sol(lines)
    # Phần III
    s += tex_section(2)
    for i, q in enumerate(P3):
        s += "\n%% Phần III Câu %d\n\\begin{ex}\n\t%s\n" % (i + 1, q["stem"])
        if q.get("fig"):
            s += tex_fig(q["fig"])
        s += "\t\\par\\shortans[oly]{%s}\n" % q["ans"]
        s += "\t\\loigiai{\n%s\n\t}\n\\end{ex}\n" % tex_sol(q["sol"])
    # Phần IV
    if P4:
        s += tex_section(3)
    for i, q in enumerate(P4):
        s += "\n%% Phần IV Câu %d\n\\begin{ex}\n\t%s\n" % (i + 1, q["stem"])
        if q["parts"]:
            s += "\t\\par " + " \\par ".join("\\textbf{%s)} %s" % ("abc"[k], p) for k, p in enumerate(q["parts"])) + "\n"
        if q.get("fig"):
            s += tex_fig(q["fig"])
        s += "\t\\loigiai{\n%s\n%s\t}\n\\end{ex}\n" % (tex_sol(q["sol"]), tex_fig(q["solfig"]) if q.get("solfig") else "")

    if is_sol:
        s += tex_answer_key()
    s += "\n\\vspace{0.4cm}\n\\begin{center}\n\t\\textbf{--------- HẾT ---------}\n\\end{center}\n\n\\end{document}\n"
    return s


def compile_latex(name, content):
    tex_path = os.path.join(LATEX_DIR, name + ".tex")
    with open(tex_path, "w", encoding="utf-8") as f:
        f.write(content)
    for run in range(2):
        res = subprocess.run(["pdflatex", "-interaction=nonstopmode", name + ".tex"], cwd=LATEX_DIR,
                             capture_output=True, text=True, encoding="utf-8", errors="ignore")
    log = open(os.path.join(LATEX_DIR, name + ".log"), encoding="utf-8", errors="ignore").read()
    errors = [l for l in log.splitlines() if l.startswith("!")]
    pages = re.search(r"Output written on .*?\((\d+) pages?", log)
    print(f"[LaTeX] {name}: {len(errors)} lỗi, {pages.group(1) if pages else '?'} trang")
    for e in errors[:10]:
        print("   ", e)
    return len(errors) == 0


# ============================================================================
# 2. SINH WORD (Pandoc -> OMML + python-docx)
# ============================================================================
def vis_len(s):
    """Độ dài hiển thị ước lượng của một phương án."""
    math = re.findall(r"\$(.*?)\$", s)
    txt = re.sub(r"\$(.*?)\$", "", s)
    mlen = sum(len(re.sub(r"\\[a-zA-Z]+|[{}]", "", m)) for m in math)
    if "frac" in s:
        mlen = int(mlen * 0.8)
    return len(txt) + mlen


def word_opts(opts):
    L = max(vis_len(o) for o in opts) + 4
    items = [r"\textbf{%s.} %s" % (LETTERS[k], o if o.endswith(".") else o + ".") for k, o in enumerate(opts)]
    if L <= 24:
        return "@@TAB@@" + "@@TAB@@".join(items) + "\n\n"
    if L <= 48:
        return "@@TAB@@" + items[0] + "@@TAB@@" + items[1] + "\n\n@@TAB@@" + items[2] + "@@TAB@@" + items[3] + "\n\n"
    return "".join("@@TAB@@" + it + "\n\n" for it in items)


def word_sol(lines):
    return "\n\n".join(lines) + "\n\n"


def word_answer_key():
    s = "\n\n@@KEY_TITLE@@\n\n\\textbf{Phần I.}\n\n\\begin{tabular}{|c|" + "c|" * len(P1) + "}\\hline\n"
    s += "Câu & " + " & ".join(str(i + 1) for i in range(len(P1))) + " \\\\ \\hline\n"
    s += "Đáp án & " + " & ".join(LETTERS[q["ans"]] for q in P1) + " \\\\ \\hline\n\\end{tabular}\n\n"
    s += "\\textbf{Phần II.}\n\n\\begin{tabular}{|c|c|c|c|c|}\\hline\nCâu & a) & b) & c) & d) \\\\ \\hline\n"
    for i, q in enumerate(P2):
        s += "%d & " % (i + 1) + " & ".join("Đ" if t else "S" for t in q["truth"]) + " \\\\ \\hline\n"
    s += "\\end{tabular}\n\n\\textbf{Phần III.}\n\n\\begin{tabular}{|c|" + "c|" * len(P3) + "}\\hline\nCâu & " + " & ".join(str(i + 1) for i in range(len(P3))) + " \\\\ \\hline\n"
    s += "Đáp án & " + " & ".join(q["ans"] for q in P3) + " \\\\ \\hline\n\\end{tabular}\n\n"
    return s


def build_pandoc_tex(is_sol):
    s = "\\documentclass{article}\n\\usepackage{amsmath,amssymb}\n\\begin{document}\n\n@@DOCUMENT_HEADER@@\n\n"
    s += "@@SECTION_1_HEADER@@\n\n"
    for i, q in enumerate(P1):
        s += "\\textbf{Câu %d.} %s\n\n" % (i + 1, q["stem"].replace("\\\\", "\n\n"))
        if q.get("fig"):
            s += "@@CENTER_IMAGE_%s@@\n\n" % q["fig"][0]
        s += word_opts(q["opts"])
        if is_sol:
            s += "\\textbf{Lời giải.} " + word_sol(q["sol"]) + "@@CHON@@ \\textbf{Chọn %s.}\n\n" % LETTERS[q["ans"]]
    s += "@@SECTION_2_HEADER@@\n\n"
    for i, q in enumerate(P2):
        stem = q["stem"].replace("\\begin{center}", "").replace("\\end{center}", "")
        s += "\\textbf{Câu %d.} %s\n\n" % (i + 1, stem)
        if q.get("fig"):
            s += "@@CENTER_IMAGE_%s@@\n\n" % q["fig"][0]
        for k, it in enumerate(q["items"]):
            s += "@@TAB1@@\\textbf{%s)} %s\n\n" % ("abcd"[k], it)
        if is_sol:
            s += "\\textbf{Lời giải.}\n\n"
            for k, (t, so) in enumerate(zip(q["truth"], q["sols"])):
                s += "@@TF_%s@@ \\textbf{%s) %s.} %s\n\n" % ("D" if t else "S", "abcd"[k], "Đúng" if t else "Sai", so)
    s += "@@SECTION_3_HEADER@@\n\n"
    for i, q in enumerate(P3):
        s += "\\textbf{Câu %d.} %s\n\n" % (i + 1, q["stem"])
        if q.get("fig"):
            s += "@@CENTER_IMAGE_%s@@\n\n" % q["fig"][0]
        if is_sol:
            s += "\\textbf{Lời giải.} " + word_sol(q["sol"]) + "@@CHON@@ \\textbf{Đáp số: %s.}\n\n" % q["ans"]
        else:
            s += "@@ANSWER_BOX@@\n\n"
    if P4:
        s += "@@SECTION_4_HEADER@@\n\n"
    for i, q in enumerate(P4):
        s += "\\textbf{Câu %d.} %s\n\n" % (i + 1, q["stem"])
        for k, p in enumerate(q["parts"]):
            s += "@@TAB1@@\\textbf{%s)} %s\n\n" % ("abc"[k], p)
        if q.get("fig"):
            s += "@@CENTER_IMAGE_%s@@\n\n" % q["fig"][0]
        if is_sol:
            s += "\\textbf{Lời giải.} " + word_sol(q["sol"])
            if q.get("solfig"):
                s += "@@CENTER_IMAGE_%s@@\n\n" % q["solfig"][0]
    if is_sol:
        s += word_answer_key()
    s += "\n\n@@HET@@\n\n\\end{document}\n"
    return s


# ---------------------- python-docx helpers ----------------------
def set_cell_borders(cell, spec):
    tcPr = cell._tc.get_or_add_tcPr()
    for b in tcPr.findall(qn('w:tcBorders')):
        tcPr.remove(b)
    xml = '<w:tcBorders %s>' % nsdecls('w')
    for side in ("top", "left", "bottom", "right"):
        v = spec.get(side, "none")
        if v == "none":
            xml += '<w:%s w:val="nil"/>' % side
        else:
            xml += '<w:%s w:val="single" w:sz="%s" w:space="0" w:color="%s"/>' % (side, v[0], v[1])
    xml += '</w:tcBorders>'
    tcPr.append(parse_xml(xml))


def add_run(p, text, bold=False, italic=False, size=11, color=None):
    r = p.add_run(text)
    r.bold = bold
    r.italic = italic
    r.font.name = "Times New Roman"
    r.font.size = Pt(size)
    if color:
        r.font.color.rgb = RGBColor(*color)
    return r


def replace_with(p, element):
    p._p.addprevious(element)
    p._p.getparent().remove(p._p)


def make_header_table(doc, is_sol):
    t = doc.add_table(rows=2, cols=3)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    r0 = t.rows[0]
    cl = r0.cells[0]
    cr = r0.cells[1].merge(r0.cells[2])
    cl.width, cr.width = Cm(7.0), Cm(10.5)
    p = cl.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(3)
    add_run(p, "LỚP TOÁN CÔ THÚY\n", True, size=12, color=(31, 73, 125))
    add_run(p, "SĐT: 0935.322.328\nĐịa chỉ: 50/2C Phạm Thị Liên", size=10.5)
    p = cr.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(3)
    if is_sol:
        add_run(p, "HƯỚNG DẪN GIẢI CHI TIẾT GIỮA KỲ I\n", True, size=12, color=(192, 0, 0))
        add_run(p, f"Môn: VẬT LÍ 10 – {SCHOOL}\n", True, size=11)
        add_run(p, f"Năm học {YEAR_W} – Mã đề: {MADE}", italic=True, size=10.5)
    else:
        add_run(p, f"KIỂM TRA GIỮA KỲ I – NĂM HỌC {YEAR_W}\n", True, size=12, color=(192, 0, 0))
        add_run(p, f"Môn: VẬT LÍ, Lớp 10 – {SCHOOL}\n", True, size=10.5)
        add_run(p, "Thời gian làm bài: 45 phút (Không kể thời gian phát đề)", italic=True, size=10)
    r1 = t.rows[1]
    if is_sol:
        c0 = r1.cells[0].merge(r1.cells[1])
        p = c0.paragraphs[0]
        add_run(p, " Giáo viên: ", True, size=10.5)
        add_run(p, "HỒ THỊ THÚY     ", size=10.5)
        add_run(p, "Môn học: ", True, size=10.5)
        add_run(p, "VẬT LÍ 10", size=10.5)
    else:
        add_run(r1.cells[0].paragraphs[0], "Họ và tên thí sinh: ..............................", size=10)
        add_run(r1.cells[1].paragraphs[0], "SBD: ......................", size=10)
    p = r1.cells[2].paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(p, f"Mã đề thi {MADE}", True, size=11)
    for row in t.rows:
        row._tr.get_or_add_trPr().append(parse_xml(r'<w:cantSplit %s/>' % nsdecls('w')))
        for c in row.cells:
            set_cell_borders(c, {k: ("6", "808080") for k in ("top", "left", "bottom", "right")})
            for pp in c.paragraphs:
                pp.paragraph_format.space_before = Pt(3)
                pp.paragraph_format.space_after = Pt(3)
    return t._tbl


def make_section_box(doc, title, note):
    t = doc.add_table(rows=1, cols=1)
    c = t.rows[0].cells[0]
    c.width = Cm(17.5)
    set_cell_borders(c, {"left": ("24", "1F497D")})
    c._tc.get_or_add_tcPr().append(parse_xml(r'<w:shd %s w:val="clear" w:color="auto" w:fill="F2F5F9"/>' % nsdecls('w')))
    p = c.paragraphs[0]
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(3)
    add_run(p, title, True, size=11.5, color=(31, 73, 125))
    add_run(p, "\n" + note, italic=True, size=10, color=(80, 80, 80))
    return t._tbl


def make_answer_box(doc):
    t = doc.add_table(rows=1, cols=5)
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    c0 = t.rows[0].cells[0]
    c0.width = Cm(1.8)
    set_cell_borders(c0, {})
    p = c0.paragraphs[0]
    p.add_run("\t")
    add_run(p, "KQ:", True)
    for k in range(1, 5):
        c = t.rows[0].cells[k]
        c.width = Cm(0.65)
        c.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_borders(c, {s: ("8", "000000") for s in ("top", "left", "bottom", "right")})
    t.rows[0]._tr.get_or_add_trPr().append(parse_xml(r'<w:trHeight %s w:val="380" w:hRule="exact"/>' % nsdecls('w')))
    return t._tbl


def add_page_field(p, size=9.5):
    def fld(instr):
        r = p.add_run()
        r.font.size = Pt(size)
        r.italic = True
        b = OxmlElement('w:fldChar'); b.set(qn('w:fldCharType'), 'begin')
        it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve'); it.text = instr
        e = OxmlElement('w:fldChar'); e.set(qn('w:fldCharType'), 'end')
        r._r.append(b); r._r.append(it); r._r.append(e)
    add_run(p, "Trang ", italic=True, size=size, color=(80, 80, 80))
    fld("PAGE")
    add_run(p, "/", italic=True, size=size, color=(80, 80, 80))
    fld("NUMPAGES")


def strip_marker(p, marker):
    for r in p.runs:
        if marker in r.text:
            r.text = r.text.replace(marker, "").lstrip()


def set_spacing(p, before, after, jc=None):
    pPr = p._p.get_or_add_pPr()
    for sp in pPr.findall(qn('w:spacing')):
        pPr.remove(sp)
    pPr.append(parse_xml(r'<w:spacing %s w:before="%d" w:after="%d" w:line="276" w:lineRule="auto"/>' % (nsdecls('w'), before, after)))
    if jc:
        for j in pPr.findall(qn('w:jc')):
            pPr.remove(j)
        pPr.append(parse_xml(r'<w:jc %s w:val="%s"/>' % (nsdecls('w'), jc)))


def apply_tabs(p):
    xml = p._p.xml
    n = xml.count("@@TAB@@")
    one = "@@TAB1@@" in xml
    xml = xml.replace("@@TAB1@@", '</w:t></w:r><w:r><w:tab/></w:r><w:r><w:t xml:space="preserve">')
    xml = xml.replace("@@TAB@@", '</w:t></w:r><w:r><w:tab/></w:r><w:r><w:t xml:space="preserve">')
    new_p = parse_xml(xml)
    p._p.getparent().replace(p._p, new_p)
    pPr = new_p.get_or_add_pPr()
    for tabs in pPr.findall(qn('w:tabs')):
        pPr.remove(tabs)
    pos = {4: [400, 2800, 5200, 7600], 2: [400, 5200]}.get(n, [400])
    if one:
        pos = [400]
    pPr.append(parse_xml('<w:tabs %s>' % nsdecls('w') + "".join('<w:tab w:val="left" w:pos="%d"/>' % x for x in pos) + '</w:tabs>'))
    for sp in pPr.findall(qn('w:spacing')):
        pPr.remove(sp)
    pPr.append(parse_xml(r'<w:spacing %s w:before="20" w:after="30" w:line="276" w:lineRule="auto"/>' % nsdecls('w')))


def style_tables(doc):
    """Bảng đáp án / bảng số liệu do Pandoc sinh: kẻ viền, căn giữa."""
    for t in doc.tables:
        if t.style is not None and t.style.name == "Table":
            pass
        tblPr = t._tbl.tblPr
        if tblPr.find(qn('w:tblBorders')) is None and len(t.columns) >= 5 and "KQ:" not in t.rows[0].cells[0].text:
            tblPr.append(parse_xml(r'<w:tblBorders %s><w:top w:val="single" w:sz="6" w:color="000000"/><w:left w:val="single" w:sz="6" w:color="000000"/><w:bottom w:val="single" w:sz="6" w:color="000000"/><w:right w:val="single" w:sz="6" w:color="000000"/><w:insideH w:val="single" w:sz="6" w:color="000000"/><w:insideV w:val="single" w:sz="6" w:color="000000"/></w:tblBorders>' % nsdecls('w')))
            for j in tblPr.findall(qn('w:jc')):
                tblPr.remove(j)
            tblPr.append(parse_xml(r'<w:jc %s w:val="center"/>' % nsdecls('w')))
            for row in t.rows:
                for c in row.cells:
                    for pp in c.paragraphs:
                        pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        for r in pp.runs:
                            r.font.name = "Times New Roman"
                            r.font.size = Pt(11)


def build_word(is_sol, out_path):
    kind = "hdg" if is_sol else "de"
    tmp_tex = os.path.join(SCRATCH_DIR, f"temp_{kind}.tex")
    tmp_docx = os.path.join(SCRATCH_DIR, f"temp_{kind}.docx")
    with open(tmp_tex, "w", encoding="utf-8") as f:
        f.write(build_pandoc_tex(is_sol))
    res = subprocess.run([PANDOC, tmp_tex, "-o", tmp_docx, "--from=latex", "--to=docx"],
                         capture_output=True, text=True, encoding="utf-8")
    if res.returncode != 0:
        print("[LỖI PANDOC]", res.stderr)
        return False
    doc = docx.Document(tmp_docx)

    for s in doc.sections:
        s.page_width, s.page_height = Cm(21.0), Cm(29.7)
        s.top_margin, s.bottom_margin = Cm(1.5), Cm(1.5)
        s.left_margin, s.right_margin = Cm(1.8), Cm(1.5)
        s.footer_distance = Cm(0.8)
        fp = s.footer.paragraphs[0]
        fp.text = ""
        fp.paragraph_format.tab_stops.add_tab_stop(Cm(17.7), alignment=2)
        add_run(fp, BRAND_W, italic=True, size=9.5, color=(100, 100, 100))
        add_run(fp, "\t", size=9.5)
        add_page_field(fp)
        add_run(fp, f" – Mã đề {MADE}", italic=True, size=9.5, color=(80, 80, 80))

    st = doc.styles['Normal']
    st.font.name = 'Times New Roman'
    st.font.size = Pt(11)
    st.element.get_or_add_rPr().append(parse_xml(r'<w:rFonts %s w:ascii="Times New Roman" w:hAnsi="Times New Roman" w:cs="Times New Roman" w:eastAsia="Times New Roman"/>' % nsdecls('w')))
    for sname in ("Body Text", "First Paragraph", "Compact"):
        try:
            sty = doc.styles[sname]
            sty.font.name = 'Times New Roman'
            sty.paragraph_format.space_before = Pt(0)
            sty.paragraph_format.space_after = Pt(2)
        except KeyError:
            pass

    style_tables(doc)

    for p in list(doc.paragraphs):
        t = p.text
        if "@@DOCUMENT_HEADER@@" in t:
            replace_with(p, make_header_table(doc, is_sol))
        elif "@@SECTION_" in t:
            k = int(re.search(r"@@SECTION_(\d)_HEADER@@", t).group(1)) - 1
            replace_with(p, make_section_box(doc, *SECTIONS[k]))
        elif "@@ANSWER_BOX@@" in t:
            replace_with(p, make_answer_box(doc))
        elif "@@CENTER_IMAGE_" in t:
            name = re.search(r"@@CENTER_IMAGE_(\w+?)@@", t).group(1)
            width = next(f[1] for q in P1 + P2 + P3 + P4 for f in (q.get("fig"), q.get("solfig")) if f and f[0] == name)
            for r in list(p.runs):
                r._r.getparent().remove(r._r)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(os.path.join(FIG_DIR, name + ".png"), width=Cm(width))
            set_spacing(p, 40, 40, "center")
        elif "@@KEY_TITLE@@" in t:
            for r in list(p.runs):
                r._r.getparent().remove(r._r)
            add_run(p, f"BẢNG ĐÁP ÁN – MÃ ĐỀ {MADE}", True, size=12.5, color=(192, 0, 0))
            set_spacing(p, 200, 80, "center")
        elif "@@HET@@" in t:
            for r in list(p.runs):
                r._r.getparent().remove(r._r)
            add_run(p, "--------- HẾT ---------", True)
            set_spacing(p, 160, 0, "center")

    for p in list(doc.paragraphs):
        t = p.text
        if "@@TAB" in p._p.xml:
            apply_tabs(p)
            continue
        if "@@CHON@@" in t:
            strip_marker(p, "@@CHON@@")
            for r in p.runs:
                r.bold = True
                r.font.color.rgb = RGBColor(192, 0, 0)
        elif "@@TF_D@@" in t or "@@TF_S@@" in t:
            ok = "@@TF_D@@" in t
            strip_marker(p, "@@TF_D@@")
            strip_marker(p, "@@TF_S@@")
            for r in p.runs:
                if r.bold and ("Đúng." in r.text or "Sai." in r.text):
                    r.font.color.rgb = RGBColor(0, 128, 0) if ok else RGBColor(192, 0, 0)
        elif t.startswith("Lời giải."):
            for r in p.runs:
                if "Lời giải." in r.text:
                    r.bold = True
                    r.font.color.rgb = RGBColor(31, 73, 125)
                    break
        if t.startswith("Câu "):
            set_spacing(p, 100, 30, "both")
            for r in p.runs[:1]:
                r.font.color.rgb = RGBColor(31, 73, 125)
        for r in p.runs:
            r.font.name = "Times New Roman"
            if r.font.size is None:
                r.font.size = Pt(11)

    # Kiểm tra sót marker + đếm OMML
    body_xml = doc.element.body.xml
    left = re.findall(r"@@\w+@@", body_xml)
    omml = body_xml.count("<m:oMath>") + body_xml.count("<m:oMath ")
    try:
        doc.save(out_path)
    except PermissionError:
        alt = out_path.replace(".docx", "_moi.docx")
        doc.save(alt)
        print(f"[CẢNH BÁO] {os.path.basename(out_path)} đang mở -> lưu tạm {os.path.basename(alt)}")
        out_path = alt
    print(f"[Word] {os.path.basename(out_path)}: {omml} công thức OMML, marker sót: {left if left else 0}")
    return True


def safe_copy(src, dst):
    try:
        shutil.copy2(src, dst)
        print(f"[COPY] {os.path.basename(dst)}")
    except PermissionError:
        print(f"[CẢNH BÁO] {os.path.basename(dst)} đang được mở, bỏ qua sao chép.")


def main():
    print("=" * 60)
    print(f" XUẤT BẢN ĐỀ VẬT LÍ 10 – {SCHOOL} – MÃ ĐỀ {MADE}")
    print("=" * 60)
    ok_de = compile_latex(NAME_DE, build_latex(False))
    ok_hdg = compile_latex(NAME_HDG, build_latex(True))
    build_word(False, os.path.join(SAN_PHAM_DIR, NAME_DE + ".docx"))
    build_word(True, os.path.join(SAN_PHAM_DIR, NAME_HDG + ".docx"))
    for n in (NAME_DE, NAME_HDG):
        src = os.path.join(LATEX_DIR, n + ".pdf")
        if os.path.exists(src):
            safe_copy(src, os.path.join(SAN_PHAM_DIR, n + ".pdf"))
    print("HOÀN TẤT!" if ok_de and ok_hdg else "CÓ LỖI LATEX – kiểm tra log.")


def run(cfg):
    """cfg: FOLDER, MADE, PREFIX, SCHOOL, YEAR, YEAR_W, P1, P2, P3, P4, [POINTS]"""
    g = globals()
    g.update(cfg)
    g["FIG_DIR"] = os.path.join(BASE_DIR, "He_Thong", "Hinh_Anh", g["FOLDER"])
    g["LATEX_DIR"] = os.path.join(BASE_DIR, "He_Thong", "LaTeX", g["FOLDER"])
    g["SCRATCH_DIR"] = os.path.join(CURRENT_DIR, "scratch_" + g["FOLDER"].lower())
    for d in (g["LATEX_DIR"], g["SCRATCH_DIR"], g["FIG_DIR"]):
        os.makedirs(d, exist_ok=True)
    sty = os.path.join(g["LATEX_DIR"], "ex_test.sty")
    if not os.path.exists(sty):
        shutil.copy2(os.path.join(BASE_DIR, "He_Thong", "Quy_Chuan", "ex_test.sty"), sty)
    g["NAME_DE"] = "%s_De_%s" % (g["PREFIX"], g["MADE"])
    g["NAME_HDG"] = "%s_HDG_%s" % (g["PREFIX"], g["MADE"])
    n1, n2, n3, n4 = len(g["P1"]), len(g["P2"]), len(g["P3"]), len(g["P4"])
    pts = cfg.get("POINTS", ("3,0", "2,0", "2,0", "3,0"))
    secs = [("PHẦN I. Câu trắc nghiệm nhiều phương án lựa chọn (%s điểm)" % pts[0],
             "Thí sinh trả lời từ câu 1 đến câu %d. Mỗi câu hỏi thí sinh chỉ chọn một phương án." % n1),
            ("PHẦN II. Câu trắc nghiệm đúng sai (%s điểm)" % pts[1],
             "Thí sinh trả lời từ câu 1 đến câu %d. Trong mỗi ý a), b), c), d) ở mỗi câu, thí sinh chọn đúng hoặc sai." % n2),
            ("PHẦN III. Câu trắc nghiệm trả lời ngắn (%s điểm)" % pts[2],
             "Thí sinh trả lời từ câu 1 đến câu %d." % n3)]
    if n4:
        secs.append(("PHẦN IV. Tự luận (%s điểm)" % pts[3],
                     "Thí sinh trình bày lời giải từ câu 1 đến câu %d." % n4))
    g["SECTIONS"] = secs
    main()

