import os
import sys
import subprocess
import shutil
import docx
from docx.shared import Pt, Cm, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn

if sys.stdout.encoding.lower() != 'utf-8':
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except Exception:
        pass

CURRENT_DIR = os.path.dirname(os.path.abspath(__file__))
BASE_DIR = os.path.abspath(os.path.join(CURRENT_DIR, "..", ".."))
HINH_ANH_DIR = os.path.join(BASE_DIR, "He_Thong", "Hinh_Anh", "Nguyen_Hue_10")
LATEX_DIR = os.path.join(BASE_DIR, "He_Thong", "LaTeX", "Nguyen_Hue_10")
SAN_PHAM_DIR = os.path.join(BASE_DIR, "San_Pham")
pandoc_exe = os.path.expandvars(r"%LOCALAPPDATA%\Pandoc\pandoc.exe")

os.makedirs(SAN_PHAM_DIR, exist_ok=True)

OUTPUT_DE_DOCX = os.path.join(SAN_PHAM_DIR, "PHY10_Nguyen_Hue_De_101.docx")
OUTPUT_HDG_DOCX = os.path.join(SAN_PHAM_DIR, "PHY10_Nguyen_Hue_HDG_101.docx")

def create_answer_box_table(doc):
    tbl = doc.add_table(rows=1, cols=5)
    tbl.alignment = WD_TABLE_ALIGNMENT.LEFT
    
    tblPr = tbl._tbl.tblPr
    for look in tblPr.findall(qn('w:tblLook')):
        tblPr.remove(look)
    tblLook = parse_xml(r'<w:tblLook %s w:val="0000" w:firstRow="0" w:lastRow="0" w:firstColumn="0" w:lastColumn="0" w:noHBand="0" w:noVBand="0"/>' % nsdecls('w'))
    for b in tblPr.findall(qn('w:tblBorders')):
        tblPr.remove(b)
    tblBorders = parse_xml(
        r'<w:tblBorders %s>'
        r'<w:top w:val="none"/>'
        r'<w:left w:val="none"/>'
        r'<w:bottom w:val="none"/>'
        r'<w:right w:val="none"/>'
        r'<w:insideH w:val="none"/>'
        r'<w:insideV w:val="none"/>'
        r'</w:tblBorders>' % nsdecls('w')
    )
    tblPr.append(tblBorders)
    tblPr.append(tblLook)
    
    cell_kq = tbl.rows[0].cells[0]
    cell_kq.width = Cm(1.8)
    cell_kq.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    tcPr0 = cell_kq._tc.get_or_add_tcPr()
    for b in tcPr0.findall(qn('w:tcBorders')):
        tcPr0.remove(b)
    tcPr0.append(parse_xml(
        r'<w:tcBorders %s>'
        r'<w:top w:val="none"/>'
        r'<w:left w:val="none"/>'
        r'<w:bottom w:val="none"/>'
        r'<w:right w:val="none"/>'
        r'</w:tcBorders>' % nsdecls('w')
    ))
    
    p_kq = cell_kq.paragraphs[0]
    p_kq.paragraph_format.space_before = Pt(0)
    p_kq.paragraph_format.space_after = Pt(0)
    p_kq.add_run("\t")
    r = p_kq.add_run("KQ:")
    r.bold = True
    r.font.name = "Times New Roman"
    r.font.size = Pt(11)
    
    for col_idx in range(1, 5):
        cell = tbl.rows[0].cells[col_idx]
        cell.width = Cm(0.65)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        tcPr = cell._tc.get_or_add_tcPr()
        for b in tcPr.findall(qn('w:tcBorders')):
            tcPr.remove(b)
        borders = parse_xml(
            r'<w:tcBorders %s>'
            r'<w:top w:val="single" w:sz="8" w:space="0" w:color="000000"/>'
            r'<w:left w:val="single" w:sz="8" w:space="0" w:color="000000"/>'
            r'<w:bottom w:val="single" w:sz="8" w:space="0" w:color="000000"/>'
            r'<w:right w:val="single" w:sz="8" w:space="0" w:color="000000"/>'
            r'</w:tcBorders>' % nsdecls('w')
        )
        tcPr.append(borders)
    
    trPr = tbl.rows[0]._tr.get_or_add_trPr()
    trHeight = parse_xml(r'<w:trHeight %s w:val="380" w:hRule="exact"/>' % nsdecls('w'))
    trPr.append(trHeight)
    return tbl

def insert_header_and_code(doc, target_p, is_solution=False):
    table = doc.add_table(rows=2, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    row0 = table.rows[0]
    cell_left = row0.cells[0]
    cell_right = row0.cells[1]
    cell_right.merge(row0.cells[2])
    
    cell_left.width = Cm(8.0)
    cell_right.width = Cm(9.5)
    
    p_l = cell_left.paragraphs[0]
    p_l.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_l.paragraph_format.line_spacing = 1.15
    p_l.paragraph_format.space_before = Pt(3)
    p_l.paragraph_format.space_after = Pt(3)
    r1 = p_l.add_run("LỚP TOÁN CÔ THÚY\n")
    r1.bold = True
    r1.font.name = "Times New Roman"
    r1.font.size = Pt(12)
    r1.font.color.rgb = RGBColor(31, 73, 125)
    
    r2 = p_l.add_run("SĐT: 0935.322.328\nĐịa chỉ: 50/2C Phạm Thị Liên")
    r2.font.name = "Times New Roman"
    r2.font.size = Pt(10.5)
    
    p_r = cell_right.paragraphs[0]
    p_r.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_r.paragraph_format.line_spacing = 1.15
    p_r.paragraph_format.space_before = Pt(3)
    p_r.paragraph_format.space_after = Pt(3)
    
    if is_solution:
        r_top = p_r.add_run("HƯỚNG DẪN GIẢI CHI TIẾT GIỮA KỲ I\n")
        r_top.bold = True
        r_top.font.name = "Times New Roman"
        r_top.font.size = Pt(12)
        r_top.font.color.rgb = RGBColor(192, 0, 0)
        
        r3 = p_r.add_run("Môn: VẬT LÍ 10 – THPT NGUYỄN HUỆ (2024 – 2025)\n")
        r3.bold = True
        r3.font.name = "Times New Roman"
        r3.font.size = Pt(11)
        r3.font.color.rgb = RGBColor(0, 0, 0)
        
        r_src = p_r.add_run("Thời gian làm bài: 45 phút – Mã đề: 101")
        r_src.italic = True
        r_src.font.name = "Times New Roman"
        r_src.font.size = Pt(10.5)
    else:
        r3 = p_r.add_run("KIỂM TRA GIỮA KỲ I – NĂM HỌC 2024 – 2025\n")
        r3.bold = True
        r3.font.name = "Times New Roman"
        r3.font.size = Pt(12)
        r3.font.color.rgb = RGBColor(192, 0, 0)
        
        r_src = p_r.add_run("Môn: VẬT LÍ, Lớp 10 – THPT NGUYỄN HUỆ\n")
        r_src.bold = True
        r_src.font.name = "Times New Roman"
        r_src.font.size = Pt(10.5)
        
        r4 = p_r.add_run("Thời gian làm bài: 45 phút (Không kể thời gian phát đề)")
        r4.italic = True
        r4.font.name = "Times New Roman"
        r4.font.size = Pt(10)
    
    row1 = table.rows[1]
    if is_solution:
        c_hoten = row1.cells[0]
        c_hoten.merge(row1.cells[1])
        c_hoten.width = Cm(13.5)
        c_made = row1.cells[2]
        c_made.width = Cm(4.0)
        
        p_info = c_hoten.paragraphs[0]
        p_info.paragraph_format.space_before = Pt(3)
        p_info.paragraph_format.space_after = Pt(3)
        r_ht = p_info.add_run(" Giáo viên: ")
        r_ht.bold = True
        r_ht.font.name = "Times New Roman"
        r_ht.font.size = Pt(10.5)
        r_ht_v = p_info.add_run("HỒ THỊ THÚY     ")
        r_ht_v.font.name = "Times New Roman"
        r_ht_v.font.size = Pt(10.5)
        
        r_mh = p_info.add_run("Môn học: ")
        r_mh.bold = True
        r_mh.font.name = "Times New Roman"
        r_mh.font.size = Pt(10.5)
        r_mh_v = p_info.add_run("VẬT LÍ 10")
        r_mh_v.font.name = "Times New Roman"
        r_mh_v.font.size = Pt(10.5)
        
        p_code = c_made.paragraphs[0]
        p_code.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_code.paragraph_format.space_before = Pt(3)
        p_code.paragraph_format.space_after = Pt(3)
        rc = p_code.add_run("Mã đề thi 101")
        rc.bold = True
        rc.font.name = "Times New Roman"
        rc.font.size = Pt(11)
    else:
        c_hoten = row1.cells[0]
        c_sbd = row1.cells[1]
        c_made = row1.cells[2]
        c_hoten.width = Cm(8.0)
        c_sbd.width = Cm(5.5)
        c_made.width = Cm(4.0)
        
        p_ht = c_hoten.paragraphs[0]
        p_ht.paragraph_format.space_before = Pt(3)
        p_ht.paragraph_format.space_after = Pt(3)
        r_ht = p_ht.add_run("Họ và tên thí sinh: ............................")
        r_ht.font.name = "Times New Roman"
        r_ht.font.size = Pt(10)
        
        p_sbd = c_sbd.paragraphs[0]
        p_sbd.paragraph_format.space_before = Pt(3)
        p_sbd.paragraph_format.space_after = Pt(3)
        r_sbd = p_sbd.add_run("SBD: ......................")
        r_sbd.font.name = "Times New Roman"
        r_sbd.font.size = Pt(10)
        
        p_code = c_made.paragraphs[0]
        p_code.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_code.paragraph_format.space_before = Pt(3)
        p_code.paragraph_format.space_after = Pt(3)
        rc = p_code.add_run("Mã đề thi 101")
        rc.bold = True
        rc.font.name = "Times New Roman"
        rc.font.size = Pt(11)
        
    for row in table.rows:
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(r'<w:cantSplit %s/>' % nsdecls('w')))
        for cell in row.cells:
            tcPr = cell._tc.get_or_add_tcPr()
            for b in tcPr.findall(qn('w:tcBorders')):
                tcPr.remove(b)
            borders = parse_xml(
                r'<w:tcBorders %s>'
                r'<w:top w:val="single" w:sz="6" w:space="0" w:color="808080"/>'
                r'<w:left w:val="single" w:sz="6" w:space="0" w:color="808080"/>'
                r'<w:bottom w:val="single" w:sz="6" w:space="0" w:color="808080"/>'
                r'<w:right w:val="single" w:sz="6" w:space="0" w:color="808080"/>'
                r'</w:tcBorders>' % nsdecls('w')
            )
            tcPr.append(borders)
            
    target_p._p.addprevious(table._tbl)
    target_p._p.getparent().remove(target_p._p)

def insert_section_header(doc, target_p, title, note=""):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    cell = table.rows[0].cells[0]
    cell.width = Cm(17.5)
    
    tcPr = cell._tc.get_or_add_tcPr()
    for b in tcPr.findall(qn('w:tcBorders')):
        tcPr.remove(b)
    borders = parse_xml(
        r'<w:tcBorders %s>'
        r'<w:top w:val="none"/>'
        r'<w:left w:val="single" w:sz="24" w:space="0" w:color="1F497D"/>'
        r'<w:bottom w:val="none"/>'
        r'<w:right w:val="none"/>'
        r'</w:tcBorders>' % nsdecls('w')
    )
    tcPr.append(borders)
    shd = parse_xml(r'<w:shd %s w:fill="F2F5F9"/>' % nsdecls('w'))
    tcPr.append(shd)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    
    r_title = p.add_run(title)
    r_title.bold = True
    r_title.font.name = "Times New Roman"
    r_title.font.size = Pt(11.5)
    r_title.font.color.rgb = RGBColor(31, 73, 125)
    
    if note:
        r_note = p.add_run(f"\n{note}")
        r_note.italic = True
        r_note.font.name = "Times New Roman"
        r_note.font.size = Pt(10)
        r_note.font.color.rgb = RGBColor(80, 80, 80)
        
    target_p._p.addprevious(table._tbl)
    target_p._p.getparent().remove(target_p._p)

def format_doc_paragraph(p):
    p_xml = p._p.xml
    text = p.text.strip()
    
    if "@@TAB@@" in p_xml:
        tab_count = p_xml.count("@@TAB@@")
        new_xml = p_xml.replace("@@TAB@@", '</w:t><w:tab/><w:t>')
        new_p = parse_xml(new_xml)
        p._p.getparent().replace(p._p, new_p)
        
        pPr = new_p.get_or_add_pPr()
        for tabs in pPr.findall(qn('w:tabs')):
            pPr.remove(tabs)
            
        tabs_xml = r'<w:tabs %s>' % nsdecls('w')
        if tab_count == 4:
            tabs_xml += r'<w:tab w:val="left" w:pos="400"/>'
            tabs_xml += r'<w:tab w:val="left" w:pos="2800"/>'
            tabs_xml += r'<w:tab w:val="left" w:pos="5200"/>'
            tabs_xml += r'<w:tab w:val="left" w:pos="7600"/>'
        elif tab_count == 2:
            tabs_xml += r'<w:tab w:val="left" w:pos="400"/>'
            tabs_xml += r'<w:tab w:val="left" w:pos="5200"/>'
        else:
            tabs_xml += r'<w:tab w:val="left" w:pos="400"/>'
        tabs_xml += r'</w:tabs>'
        pPr.append(parse_xml(tabs_xml))
        
        for sp in pPr.findall(qn('w:spacing')):
            pPr.remove(sp)
        sp_xml = r'<w:spacing %s w:before="30" w:after="40" w:line="276" w:lineRule="auto"/>' % nsdecls('w')
        pPr.append(parse_xml(sp_xml))
        return

    pPr = p._p.get_or_add_pPr()
    if text.startswith("Câu "):
        for jc in pPr.findall(qn('w:jc')):
            pPr.remove(jc)
        pPr.append(parse_xml(r'<w:jc %s w:val="both"/>' % nsdecls('w')))
        for sp in pPr.findall(qn('w:spacing')):
            pPr.remove(sp)
        sp_xml = r'<w:spacing %s w:before="60" w:after="30" w:line="276" w:lineRule="auto"/>' % nsdecls('w')
        pPr.append(parse_xml(sp_xml))
    elif "HẾT" in text:
        for jc in pPr.findall(qn('w:jc')):
            pPr.remove(jc)
        pPr.append(parse_xml(r'<w:jc %s w:val="center"/>' % nsdecls('w')))
        for sp in pPr.findall(qn('w:spacing')):
            pPr.remove(sp)
        sp_xml = r'<w:spacing %s w:before="120" w:after="120" w:line="276" w:lineRule="auto"/>' % nsdecls('w')
        pPr.append(parse_xml(sp_xml))

    for r in p.runs:
        r.font.name = 'Times New Roman'
        r.font.size = Pt(11)
        r_rPr = r._r.get_or_add_rPr()
        f = parse_xml(r'<w:rFonts %s w:ascii="Times New Roman" w:hAnsi="Times New Roman" w:cs="Times New Roman"/>' % nsdecls('w'))
        r_rPr.append(f)

tex_de = r"""\documentclass[11pt,a4paper]{article}
\usepackage[utf8]{vietnam}
\usepackage{amsmath,amssymb}
\usepackage{graphicx}

\begin{document}

@@DOCUMENT_HEADER@@

@@SECTION_1_HEADER@@

% Câu 1
\textbf{Câu 1.} Đối tượng nghiên cứu của Vật lí là

@@TAB@@\textbf{A.} sự thay đổi của các chất khi kết hợp với nhau.@@TAB@@\textbf{B.} qui luật tương tác của các dạng năng lượng.

@@TAB@@\textbf{C.} nghiên cứu về nhiệt động lực học.@@TAB@@\textbf{D.} các dạng vận động của vật chất và năng lượng.

% Câu 2
\textbf{Câu 2.} Một vật rơi tự do từ độ cao $80\text{ m}$ xuống đất. Lấy $g = 10\text{ m/s}^2$, tốc độ vật lúc vừa chạm đất là

@@TAB@@\textbf{A.} $40\text{ m/s}$.@@TAB@@\textbf{B.} $20\text{ m/s}$.@@TAB@@\textbf{C.} $10\text{ m/s}$.@@TAB@@\textbf{D.} $30\text{ m/s}$.

% Câu 3
\textbf{Câu 3.} Sự rơi tự do là

@@TAB@@\textbf{A.} chuyển động chỉ dưới tác dụng của trọng lực.@@TAB@@\textbf{B.} một dạng chuyển động thẳng đều.

@@TAB@@\textbf{C.} chuyển động khi bỏ qua mọi lực.@@TAB@@\textbf{D.} chuyển động không chịu bất cứ lực tác dụng nào.

% Câu 4
\textbf{Câu 4.} Trong chuyển động thẳng biến đổi đều thì

@@TAB@@\textbf{A.} vận tốc là đại lượng biến thiên theo thời gian theo quy luật hàm bậc hai.

@@TAB@@\textbf{B.} gia tốc là đại lượng không đổi.

@@TAB@@\textbf{C.} vận tốc là đại lượng không đổi.

@@TAB@@\textbf{D.} gia tốc là đại lượng biến thiên theo thời gian.

% Câu 5
\textbf{Câu 5.} Gia tốc là một đại lượng

@@TAB@@\textbf{A.} vô hướng, đặc trưng cho sự biến đổi nhanh hay chậm của chuyển động.

@@TAB@@\textbf{B.} vectơ, đặc trưng cho sự biến đổi nhanh hay chậm của chuyển động.

@@TAB@@\textbf{C.} vô hướng, đặc trưng cho tính không đổi của vận tốc.

@@TAB@@\textbf{D.} vectơ, đặc trưng cho sự biến đổi nhanh hay chậm của vận tốc.

% Câu 6
\textbf{Câu 6.} Canô chuyển động thẳng xuôi dòng từ P đến Q mất 2 giờ và ngược dòng từ Q về P mất 3 giờ. Khi nước yên lặng canô chuyển động có tốc độ $50\text{ km/h}$. Tốc độ của nước so với bờ là

@@TAB@@\textbf{A.} $12{,}5\text{ km/h}$.@@TAB@@\textbf{B.} $9\text{ km/h}$.@@TAB@@\textbf{C.} $10\text{ km/h}$.@@TAB@@\textbf{D.} $20\text{ km/h}$.

% Câu 7
\textbf{Câu 7.} Kí hiệu dưới đây mang ý nghĩa gì?

@@CENTER_IMAGE_c7_rac@@

@@TAB@@\textbf{A.} Không được phép bỏ vào thùng rác.@@TAB@@\textbf{B.} Tránh ánh nắng chiếu trực tiếp.

@@TAB@@\textbf{C.} Dụng cụ dễ vỡ.@@TAB@@\textbf{D.} Dụng cụ cần đặt đứng.

% Câu 8
\textbf{Câu 8.} Đồ thị độ dịch chuyển -- thời gian trong chuyển động thẳng của một chất điểm có dạng như hình vẽ. Trong thời gian nào xe chuyển động thẳng đều?

@@CENTER_IMAGE_c8_dothidt@@

@@TAB@@\textbf{A.} Không có lúc nào xe chuyển động thẳng đều.

@@TAB@@\textbf{B.} Trong khoảng thời gian từ $0$ đến $t_1$.

@@TAB@@\textbf{C.} Trong khoảng thời gian từ $t_1$ đến $t_2$.

@@TAB@@\textbf{D.} Trong khoảng thời gian từ $0$ đến $t_2$.

% Câu 9
\textbf{Câu 9.} Phát biểu nào sau đây \textbf{không đúng}?

@@TAB@@\textbf{A.} Độ dịch chuyển là vectơ nối vị trí đầu và vị trí cuối của chất điểm chuyển động.

@@TAB@@\textbf{B.} Chất điểm đi trên một đường thẳng rồi quay về vị trí ban đầu thì có độ dịch chuyển bằng không.

@@TAB@@\textbf{C.} Độ dịch chuyển bằng quãng đường đi được của chất điểm.

@@TAB@@\textbf{D.} Độ dịch chuyển là đại lượng vectơ còn quãng đường đi được là đại lượng vô hướng.

% Câu 10
\textbf{Câu 10.} Dùng thước đo có sai số dụng cụ là $1\text{ mm}$ để đo 5 lần khoảng cách giữa hai điểm M và N đều cho một giá trị như nhau là $56\text{ mm}$. Kết quả của phép đo được viết là

@@TAB@@\textbf{A.} $d = 56 \pm 0\text{ mm}$.@@TAB@@\textbf{B.} $d = 56 \pm 2\text{ mm}$.@@TAB@@\textbf{C.} $d = 56 \pm 1\text{ mm}$.@@TAB@@\textbf{D.} $d = 56 \pm 0{,}5\text{ mm}$.

% Câu 11
\textbf{Câu 11.} Trong các chuyển động sau, chuyển động nào được coi là sự rơi tự do?

@@TAB@@\textbf{A.} Chiếc lá đang rơi.@@TAB@@\textbf{B.} Hạt bụi chuyển động trong không khí.

@@TAB@@\textbf{C.} Vận động viên đang nhảy dù.@@TAB@@\textbf{D.} Quả tạ rơi trong không khí từ độ cao $50\text{ m}$.

% Câu 12
\textbf{Câu 12.} Một ô tô đang chuyển động cùng chiều dương với tốc độ $10\text{ m/s}$ trên đoạn đường thẳng thì người lái xe hãm phanh và ô tô chuyển động chậm dần đều. Cho tới khi dừng hẳn thì ô tô đã chạy thêm được $100\text{ m}$. Gia tốc của xe có giá trị bằng

@@TAB@@\textbf{A.} $-0{,}2\text{ m/s}^2$.@@TAB@@\textbf{B.} $-0{,}5\text{ m/s}^2$.@@TAB@@\textbf{C.} $0{,}2\text{ m/s}^2$.@@TAB@@\textbf{D.} $0{,}5\text{ m/s}^2$.

% Câu 13
\textbf{Câu 13.} Một vật được thả rơi tự do, tốc độ của vật khi chạm đất là $70\text{ m/s}$. Cho $g = 10\text{ m/s}^2$. Thời gian vật rơi là

@@TAB@@\textbf{A.} $3{,}5\text{ s}$.@@TAB@@\textbf{B.} $7\text{ s}$.@@TAB@@\textbf{C.} $8\text{ s}$.@@TAB@@\textbf{D.} $4\text{ s}$.

% Câu 14
\textbf{Câu 14.} Chọn câu đúng. Để đo tốc độ trung bình của vật trong phòng thí nghiệm, ta cần dùng

@@TAB@@\textbf{A.} thước đo chiều dài và đồng hồ đo thời gian.@@TAB@@\textbf{B.} thước đo và lực kế.

@@TAB@@\textbf{C.} máy bắn tốc độ.@@TAB@@\textbf{D.} lực kế và đồng hồ đo thời gian.

% Câu 15
\textbf{Câu 15.} Cặp đồ thị nào ở hình dưới đây là của chuyển động thẳng đều?

@@CENTER_IMAGE_c15_dothi4@@

@@TAB@@\textbf{A.} I và III.@@TAB@@\textbf{B.} I và IV.@@TAB@@\textbf{C.} II và IV.@@TAB@@\textbf{D.} II và III.

% Câu 16
\textbf{Câu 16.} Vật chuyển động thẳng nhanh dần thì vectơ gia tốc và vectơ vận tốc

@@TAB@@\textbf{A.} ngược hướng.@@TAB@@\textbf{B.} cùng hướng.@@TAB@@\textbf{C.} không thay đổi.@@TAB@@\textbf{D.} vuông góc.

% Câu 17
\textbf{Câu 17.} Câu trả lời nào sau đây \textbf{Sai}. Chuyển động thẳng nhanh dần đều là chuyển động có

@@TAB@@\textbf{A.} quãng đường đi được của vật luôn tỉ lệ thuận với thời gian vật đi.

@@TAB@@\textbf{B.} quỹ đạo là đường thẳng.

@@TAB@@\textbf{C.} vectơ gia tốc của vật có độ lớn là một hằng số.

@@TAB@@\textbf{D.} vận tốc có độ lớn tăng theo hàm bậc nhất đối với thời gian.

% Câu 18
\textbf{Câu 18.} Vào lúc $10\text{ h}$, người lái xe nhìn vào tốc kế và thấy tốc kế chỉ $40\text{ km/h}$. Số liệu này cho biết

@@TAB@@\textbf{A.} tốc độ trung bình của xe.@@TAB@@\textbf{B.} vận tốc trung bình của xe.

@@TAB@@\textbf{C.} tốc độ tức thời của xe.@@TAB@@\textbf{D.} vận tốc tức thời của xe.

@@SECTION_2_HEADER@@

% Phần II Câu 1
\textbf{Câu 1.} Hai bạn Đức và Triết bơi trong bể bơi có chiều dài $25\text{ m}$. Hai bạn xuất phát từ đầu bể bơi đến cuối bể bơi thì Đức dừng lại nghỉ, còn Triết quay lại và tiếp tục bơi tiếp về đầu bể bơi mới nghỉ. Thời gian bơi của Đức là $25\text{ s}$; Triết bơi đi mất $23\text{ s}$ và bơi về mất $27\text{ s}$. Biết rằng Đức và Triết bơi trên một đường thẳng song song với thành bể.

\textbf{a)} Tốc độ trung bình trong thời gian $25\text{ s}$ của Đức là $1\text{ m/s}$.

\textbf{b)} Vận tốc trung bình trong thời gian $50\text{ s}$ của Triết là $1\text{ m/s}$.

\textbf{c)} Độ dịch chuyển tổng hợp của Triết có độ lớn bằng $50\text{ m}$.

\textbf{d)} Quãng đường Đức bơi được là $25\text{ m}$.

% Phần II Câu 2
\textbf{Câu 2.} Một máy bay Boeing 747 bắt đầu xuất phát trên một đường băng thẳng dài với gia tốc $a = 2\text{ m/s}^2$. Tốc độ cần thiết để máy bay rời đường băng là $360\text{ km/h}$. Sau khi xuất phát được $30\text{ s}$ thì nhận được lệnh huỷ bay từ trạm điều khiển không lưu do thời tiết xấu nên phi công cho máy bay hãm phanh chuyển động chậm dần đều và dừng lại sau $20\text{ s}$ kể từ khi có lệnh huỷ bay.

\textbf{a)} Chuyển động của máy bay trong $30\text{ s}$ đầu tiên là chuyển động thẳng nhanh dần đều.

\textbf{b)} Tốc độ của máy bay tại thời điểm nhận được lệnh huỷ bay là $60\text{ m/s}$.

\textbf{c)} Độ lớn gia tốc của máy bay trong giai đoạn hãm phanh là $3\text{ m/s}^2$.

\textbf{d)} Tốc độ trung bình của máy bay trong khoảng thời gian từ lúc xuất phát cho đến khi dừng lại là $40\text{ m/s}$.

% Phần II Câu 3
\textbf{Câu 3.} Từ các đồ thị trong hình sau. Nhận định sau đúng hay sai?

@@CENTER_IMAGE_p2_c3_dothi3@@

\textbf{a)} Chuyển động thẳng nhanh dần đều là đồ thị hình a và b.

\textbf{b)} Đồ thị hình b có công thức vận tốc là $v = v_0 + at\ (a < 0)$.

\textbf{c)} Chuyển động thẳng chậm dần đều là đồ thị hình c.

\textbf{d)} Đồ thị hình c có công thức vận tốc là $v = v_0 + at\ (a > 0)$.

% Phần II Câu 4
\textbf{Câu 4.} Bỏ qua sức cản không khí, một vật rơi tự do tại một địa điểm có độ cao $180\text{ m}$, lấy $g = 10\text{ m/s}^2$.

\textbf{a)} Quãng đường vật rơi được sau $3\text{ s}$ là $30\text{ m}$.

\textbf{b)} Phương rơi của vật là phương thẳng đứng.

\textbf{c)} Thời gian vật rơi hết $100\text{ m}$ cuối cùng là $2\text{ s}$.

\textbf{d)} Thời gian vật rơi hết quãng đường là $9\text{ s}$.

@@SECTION_3_HEADER@@

% Phần III Câu 1
\textbf{Câu 1.} Cho đồ thị độ dịch chuyển theo thời gian của một vật chuyển động như hình dưới đây. Hỏi tại thời điểm $t$ bằng bao nhiêu giây thì vật đổi chiều chuyển động?

@@CENTER_IMAGE_p3_c1_dothi@@

@@ANSWER_BOX@@

% Phần III Câu 2
\textbf{Câu 2.} Chạy marathon là một bộ môn thể thao dưới hình thức chạy bộ được ưa chuộng nhất trên thế giới, là một môn thể thao mang đến rất nhiều lợi ích cho sức khỏe. Người tham gia có thể chạy ở các cự ly $5\text{ km}$, $10\text{ km}$ hoặc $21\text{ km}$ rồi nâng lên đường chạy dài $42\text{ km}$ trong marathon. Tùy thuộc vào thể trạng mà chúng ta lựa chọn luyện tập ở cự ly phù hợp.

Pace là một từ tiếng Anh được sử dụng để chỉ nhịp điệu trong chạy bộ. Pace là số phút để một người hoàn thành $1\text{ km}$ trong khi chạy. Bạn Hương mới tập luyện nên chọn cự ly $5\text{ km}$, trong một lần tập bạn đạt được thời gian hoàn thành cự ly $5\text{ km}$ là 42 phút. Pace trung bình trong bài luyện tập chạy bộ của bạn Hương là bao nhiêu phút/km?

@@ANSWER_BOX@@

% Phần III Câu 3
\textbf{Câu 3.} Đồ thị bên dưới mô tả sự thay đổi vận tốc theo thời gian trong chuyển động của một ô tô. Gia tốc của ô tô từ $t = 15\text{ s}$ đến $t = 30\text{ s}$ có giá trị bao nhiêu $\text{m/s}^2$?

@@CENTER_IMAGE_p3_c3_dothi@@

@@ANSWER_BOX@@

% Phần III Câu 4
\textbf{Câu 4.} Một học sinh đo chiều dài của bàn học, kết quả thu được như sau $d = 120 \pm 1\text{ cm}$. Sai số tuyệt đối của phép đo trên là bao nhiêu cm?

@@ANSWER_BOX@@

% Phần III Câu 5
\textbf{Câu 5.} Đồ thị bên dưới mô tả sự thay đổi vận tốc theo thời gian trong chuyển động của một vật. Vật thay đổi tính chất chuyển động từ chuyển động thẳng đều sang chuyển động thẳng biến đổi đều ở thời điểm nào (bao nhiêu giây)?

@@CENTER_IMAGE_p3_c5_dothi@@

@@ANSWER_BOX@@

% Phần III Câu 6
\textbf{Câu 6.} Bạn Tuấn chạy bộ qua cầu đi thẳng $40\text{ m}$ theo hướng Đông, sau đó rẽ trái chạy thẳng theo hướng Bắc $100\text{ m}$ rồi quay ngược về hướng Nam $70\text{ m}$. Độ lớn độ dịch chuyển của Tuấn là bao nhiêu m?

@@ANSWER_BOX@@

\vspace{0.4cm}
\begin{center}
\textbf{--------- HẾT ---------}
\end{center}

\end{document}
"""

tex_hdg = r"""\documentclass[11pt,a4paper]{article}
\usepackage[utf8]{vietnam}
\usepackage{amsmath,amssymb}
\usepackage{graphicx}

\begin{document}

@@DOCUMENT_HEADER@@

@@SECTION_1_HEADER@@

% Câu 1
\textbf{Câu 1.} Đối tượng nghiên cứu của Vật lí là

@@TAB@@\textbf{A.} sự thay đổi của các chất khi kết hợp với nhau.@@TAB@@\textbf{B.} qui luật tương tác của các dạng năng lượng.

@@TAB@@\textbf{C.} nghiên cứu về nhiệt động lực học.@@TAB@@\textbf{D.} các dạng vận động của vật chất và năng lượng.

\textbf{Lời giải.} Vật lí là môn khoa học tự nhiên có đối tượng nghiên cứu tập trung vào các dạng vận động cơ bản của vật chất và năng lượng, cũng như quy luật tương tác giữa chúng.

\textbf{==> Chọn D.}

% Câu 2
\textbf{Câu 2.} Một vật rơi tự do từ độ cao $80\text{ m}$ xuống đất. Lấy $g = 10\text{ m/s}^2$, tốc độ vật lúc vừa chạm đất là

@@TAB@@\textbf{A.} $40\text{ m/s}$.@@TAB@@\textbf{B.} $20\text{ m/s}$.@@TAB@@\textbf{C.} $10\text{ m/s}$.@@TAB@@\textbf{D.} $30\text{ m/s}$.

\textbf{Lời giải.} Tốc độ của vật khi vừa chạm đất sau khi rơi tự do từ độ cao $h = 80\text{ m}$ là:
\[
v = \sqrt{2gh} = \sqrt{2 \cdot 10 \cdot 80} = \sqrt{1600} = 40\text{ m/s}.
\]

\textbf{==> Chọn A.}

% Câu 3
\textbf{Câu 3.} Sự rơi tự do là

@@TAB@@\textbf{A.} chuyển động chỉ dưới tác dụng của trọng lực.@@TAB@@\textbf{B.} một dạng chuyển động thẳng đều.

@@TAB@@\textbf{C.} chuyển động khi bỏ qua mọi lực.@@TAB@@\textbf{D.} chuyển động không chịu bất cứ lực tác dụng nào.

\textbf{Lời giải.} Theo định nghĩa trong SGK Vật lí 10, sự rơi tự do là chuyển động của một vật chỉ chịu tác dụng duy nhất của trọng lực (khi bỏ qua sức cản của không khí).

\textbf{==> Chọn A.}

% Câu 4
\textbf{Câu 4.} Trong chuyển động thẳng biến đổi đều thì

@@TAB@@\textbf{A.} vận tốc là đại lượng biến thiên theo thời gian theo quy luật hàm bậc hai.

@@TAB@@\textbf{B.} gia tốc là đại lượng không đổi.

@@TAB@@\textbf{C.} vận tốc là đại lượng không đổi.

@@TAB@@\textbf{D.} gia tốc là đại lượng biến thiên theo thời gian.

\textbf{Lời giải.} Chuyển động thẳng biến đổi đều là chuyển động thẳng có gia tốc không đổi theo thời gian ($a = \text{const}$). Vận tốc biến thiên theo hàm bậc nhất của thời gian: $v = v_0 + at$.

\textbf{==> Chọn B.}

% Câu 5
\textbf{Câu 5.} Gia tốc là một đại lượng

@@TAB@@\textbf{A.} vô hướng, đặc trưng cho sự biến đổi nhanh hay chậm của chuyển động.

@@TAB@@\textbf{B.} vectơ, đặc trưng cho sự biến đổi nhanh hay chậm của chuyển động.

@@TAB@@\textbf{C.} vô hướng, đặc trưng cho tính không đổi của vận tốc.

@@TAB@@\textbf{D.} vectơ, đặc trưng cho sự biến đổi nhanh hay chậm của vận tốc.

\textbf{Lời giải.} Gia tốc là một đại lượng vectơ đặc trưng cho sự biến đổi nhanh hay chậm của vận tốc theo thời gian, được xác định bằng:
\[
\vec{a} = \frac{\Delta \vec{v}}{\Delta t}.
\]

\textbf{==> Chọn D.}

% Câu 6
\textbf{Câu 6.} Canô chuyển động thẳng xuôi dòng từ P đến Q mất 2 giờ và ngược dòng từ Q về P mất 3 giờ. Khi nước yên lặng canô chuyển động có tốc độ $50\text{ km/h}$. Tốc độ của nước so với bờ là

@@TAB@@\textbf{A.} $12{,}5\text{ km/h}$.@@TAB@@\textbf{B.} $9\text{ km/h}$.@@TAB@@\textbf{C.} $10\text{ km/h}$.@@TAB@@\textbf{D.} $20\text{ km/h}$.

\textbf{Lời giải.} Gọi $s$ là khoảng cách giữa hai bến P và Q ($s > 0$). Tốc độ của canô đối với nước là $v_c = 50\text{ km/h}$, tốc độ dòng nước so với bờ là $v_n$ ($0 < v_n < 50$).
- Khi xuôi dòng: $v_{\text{xuôi}} = v_c + v_n = 50 + v_n \implies s = 2(50 + v_n)$.
- Khi ngược dòng: $v_{\text{ngược}} = v_c - v_n = 50 - v_n \implies s = 3(50 - v_n)$.
Vì quãng đường không đổi:
\[
2(50 + v_n) = 3(50 - v_n) \iff 100 + 2v_n = 150 - 3v_n \iff 5v_n = 50 \iff v_n = 10\text{ km/h}.
\]

\textbf{==> Chọn C.}

% Câu 7
\textbf{Câu 7.} Kí hiệu dưới đây mang ý nghĩa gì?

@@CENTER_IMAGE_c7_rac@@

@@TAB@@\textbf{A.} Không được phép bỏ vào thùng rác.@@TAB@@\textbf{B.} Tránh ánh nắng chiếu trực tiếp.

@@TAB@@\textbf{C.} Dụng cụ dễ vỡ.@@TAB@@\textbf{D.} Dụng cụ cần đặt đứng.

\textbf{Lời giải.} Biểu tượng thùng rác có bánh xe bị gạch chéo (WEEE Symbol) là biển cảnh báo đối với các thiết bị điện và điện tử, chỉ dẫn rằng không được vứt bỏ thiết bị cùng với rác thải sinh hoạt thông thường mà phải được thu gom, phân loại và xử lý riêng biệt.

\textbf{==> Chọn A.}

% Câu 8
\textbf{Câu 8.} Đồ thị độ dịch chuyển -- thời gian trong chuyển động thẳng của một chất điểm có dạng như hình vẽ. Trong thời gian nào xe chuyển động thẳng đều?

@@CENTER_IMAGE_c8_dothidt@@

@@TAB@@\textbf{A.} Không có lúc nào xe chuyển động thẳng đều.

@@TAB@@\textbf{B.} Trong khoảng thời gian từ $0$ đến $t_1$.

@@TAB@@\textbf{C.} Trong khoảng thời gian từ $t_1$ đến $t_2$.

@@TAB@@\textbf{D.} Trong khoảng thời gian từ $0$ đến $t_2$.

\textbf{Lời giải.} Trên đồ thị độ dịch chuyển -- thời gian ($d - t$):
- Trong khoảng thời gian từ $0$ đến $t_1$, đồ thị là đường thẳng xiên dốc lên xuất phát từ gốc tọa độ, hệ số góc không đổi $v = \dfrac{\Delta d}{\Delta t} > 0$, biểu diễn chuyển động thẳng đều.
- Trong khoảng thời gian từ $t_1$ đến $t_2$, đồ thị là đường nằm ngang ($d = \text{const}$), biểu diễn chất điểm đứng yên.

\textbf{==> Chọn B.}

% Câu 9
\textbf{Câu 9.} Phát biểu nào sau đây \textbf{không đúng}?

@@TAB@@\textbf{A.} Độ dịch chuyển là vectơ nối vị trí đầu và vị trí cuối của chất điểm chuyển động.

@@TAB@@\textbf{B.} Chất điểm đi trên một đường thẳng rồi quay về vị trí ban đầu thì có độ dịch chuyển bằng không.

@@TAB@@\textbf{C.} Độ dịch chuyển bằng quãng đường đi được của chất điểm.

@@TAB@@\textbf{D.} Độ dịch chuyển là đại lượng vectơ còn quãng đường đi được là đại lượng vô hướng.

\textbf{Lời giải.} Độ lớn của độ dịch chuyển chỉ bằng quãng đường đi được ($d = s$) khi chất điểm chuyển động thẳng và không đổi chiều. Trong chuyển động có đổi chiều hoặc chuyển động cong thì $d < s$. Do đó phát biểu "Độ dịch chuyển bằng quãng đường đi được của chất điểm" là không đúng.

\textbf{==> Chọn C.}

% Câu 10
\textbf{Câu 10.} Dùng thước đo có sai số dụng cụ là $1\text{ mm}$ để đo 5 lần khoảng cách giữa hai điểm M và N đều cho một giá trị như nhau là $56\text{ mm}$. Kết quả của phép đo được viết là

@@TAB@@\textbf{A.} $d = 56 \pm 0\text{ mm}$.@@TAB@@\textbf{B.} $d = 56 \pm 2\text{ mm}$.@@TAB@@\textbf{C.} $d = 56 \pm 1\text{ mm}$.@@TAB@@\textbf{D.} $d = 56 \pm 0{,}5\text{ mm}$.

\textbf{Lời giải.} Giá trị trung bình của phép đo là $\overline{d} = 56\text{ mm}$. Vì cả 5 lần đo đều ra cùng giá trị $56\text{ mm}$ nên sai số ngẫu nhiên $\overline{\Delta d} = 0$. Sai số tuyệt đối của phép đo bằng sai số dụng cụ: $\Delta d = \Delta d_{dc} = 1\text{ mm}$.
Kết quả phép đo được viết là: $d = \overline{d} \pm \Delta d = 56 \pm 1\text{ mm}$.

\textbf{==> Chọn C.}

% Câu 11
\textbf{Câu 11.} Trong các chuyển động sau, chuyển động nào được coi là sự rơi tự do?

@@TAB@@\textbf{A.} Chiếc lá đang rơi.@@TAB@@\textbf{B.} Hạt bụi chuyển động trong không khí.

@@TAB@@\textbf{C.} Vận động viên đang nhảy dù.@@TAB@@\textbf{D.} Quả tạ rơi trong không khí từ độ cao $50\text{ m}$.

\textbf{Lời giải.} Khi một vật rơi trong không khí mà lực cản của không khí rất nhỏ không đáng kể so với trọng lượng của vật thì sự rơi đó được coi gần đúng là sự rơi tự do. Quả tạ đặc bằng kim loại có trọng lượng rất lớn so với lực cản không khí nên chuyển động rơi của nó được coi là rơi tự do.

\textbf{==> Chọn D.}

% Câu 12
\textbf{Câu 12.} Một ô tô đang chuyển động cùng chiều dương với tốc độ $10\text{ m/s}$ trên đoạn đường thẳng thì người lái xe hãm phanh và ô tô chuyển động chậm dần đều. Cho tới khi dừng hẳn thì ô tô đã chạy thêm được $100\text{ m}$. Gia tốc của xe có giá trị bằng

@@TAB@@\textbf{A.} $-0{,}2\text{ m/s}^2$.@@TAB@@\textbf{B.} $-0{,}5\text{ m/s}^2$.@@TAB@@\textbf{C.} $0{,}2\text{ m/s}^2$.@@TAB@@\textbf{D.} $0{,}5\text{ m/s}^2$.

\textbf{Lời giải.} Áp dụng công thức độc lập với thời gian: $v^2 - v_0^2 = 2as$.
Với $v_0 = 10\text{ m/s}$, $v = 0$ (khi dừng hẳn), quãng đường $s = 100\text{ m}$:
\[
0^2 - 10^2 = 2 \cdot a \cdot 100 \iff -100 = 200a \implies a = -0{,}5\text{ m/s}^2.
\]

\textbf{==> Chọn B.}

% Câu 13
\textbf{Câu 13.} Một vật được thả rơi tự do, tốc độ của vật khi chạm đất là $70\text{ m/s}$. Cho $g = 10\text{ m/s}^2$. Thời gian vật rơi là

@@TAB@@\textbf{A.} $3{,}5\text{ s}$.@@TAB@@\textbf{B.} $7\text{ s}$.@@TAB@@\textbf{C.} $8\text{ s}$.@@TAB@@\textbf{D.} $4\text{ s}$.

\textbf{Lời giải.} Thời gian vật rơi tự do được tính từ công thức vận tốc:
\[
v = gt \implies t = \frac{v}{g} = \frac{70}{10} = 7\text{ s}.
\]

\textbf{==> Chọn B.}

% Câu 14
\textbf{Câu 14.} Chọn câu đúng. Để đo tốc độ trung bình của vật trong phòng thí nghiệm, ta cần dùng

@@TAB@@\textbf{A.} thước đo chiều dài và đồng hồ đo thời gian.@@TAB@@\textbf{B.} thước đo và lực kế.

@@TAB@@\textbf{C.} máy bắn tốc độ.@@TAB@@\textbf{D.} lực kế và đồng hồ đo thời gian.

\textbf{Lời giải.} Tốc độ trung bình được xác định theo công thức $v_{tb} = \dfrac{s}{t}$. Do đó, để đo tốc độ trung bình trong phòng thí nghiệm ta cần dùng thước đo chiều dài (đo quãng đường $s$) và đồng hồ bấm giây/đồng hồ hiện số (đo thời gian $t$).

\textbf{==> Chọn A.}

% Câu 15
\textbf{Câu 15.} Cặp đồ thị nào ở hình dưới đây là của chuyển động thẳng đều?

@@CENTER_IMAGE_c15_dothi4@@

@@TAB@@\textbf{A.} I và III.@@TAB@@\textbf{B.} I và IV.@@TAB@@\textbf{C.} II và IV.@@TAB@@\textbf{D.} II và III.

\textbf{Lời giải.} Trong chuyển động thẳng đều:
- Vận tốc không đổi theo thời gian: $v(t) = \text{const}$, đồ thị $v - t$ là đường thẳng nằm ngang song song trục thời gian (hình IV).
- Độ dịch chuyển tỉ lệ thuận với thời gian: $d(t) = v \cdot t$, đồ thị $d - t$ là đường thẳng đi qua gốc tọa độ (hình I).
Vậy cặp đồ thị đúng là I và IV.

\textbf{==> Chọn B.}

% Câu 16
\textbf{Câu 16.} Vật chuyển động thẳng nhanh dần thì vectơ gia tốc và vectơ vận tốc

@@TAB@@\textbf{A.} ngược hướng.@@TAB@@\textbf{B.} cùng hướng.@@TAB@@\textbf{C.} không thay đổi.@@TAB@@\textbf{D.} vuông góc.

\textbf{Lời giải.} Trong chuyển động thẳng nhanh dần, tốc độ của vật tăng theo thời gian nên tích $a \cdot v > 0$. Điều này có nghĩa là vectơ gia tốc $\vec{a}$ và vectơ vận tốc $\vec{v}$ luôn cùng hướng với nhau.

\textbf{==> Chọn B.}

% Câu 17
\textbf{Câu 17.} Câu trả lời nào sau đây \textbf{Sai}. Chuyển động thẳng nhanh dần đều là chuyển động có

@@TAB@@\textbf{A.} quãng đường đi được của vật luôn tỉ lệ thuận với thời gian vật đi.

@@TAB@@\textbf{B.} quỹ đạo là đường thẳng.

@@TAB@@\textbf{C.} vectơ gia tốc của vật có độ lớn là một hằng số.

@@TAB@@\textbf{D.} vận tốc có độ lớn tăng theo hàm bậc nhất đối với thời gian.

\textbf{Lời giải.} Phương trình quãng đường trong chuyển động thẳng nhanh dần đều có dạng $s = v_0 t + \dfrac{1}{2}at^2$, là hàm bậc hai theo thời gian, nên quãng đường không tỉ lệ thuận với thời gian. Do đó phát biểu A là sai.

\textbf{==> Chọn A.}

% Câu 18
\textbf{Câu 18.} Vào lúc $10\text{ h}$, người lái xe nhìn vào tốc kế và thấy tốc kế chỉ $40\text{ km/h}$. Số liệu này cho biết

@@TAB@@\textbf{A.} tốc độ trung bình của xe.@@TAB@@\textbf{B.} vận tốc trung bình của xe.

@@TAB@@\textbf{C.} tốc độ tức thời của xe.@@TAB@@\textbf{D.} vận tốc tức thời của xe.

\textbf{Lời giải.} Tốc kế (speedometer) gắn trên xe hiển thị độ lớn vận tốc của xe tại thời điểm quan sát, tức là tốc độ tức thời của xe.

\textbf{==> Chọn C.}

@@SECTION_2_HEADER@@

% Phần II Câu 1
\textbf{Câu 1.} Hai bạn Đức và Triết bơi trong bể bơi có chiều dài $25\text{ m}$. Hai bạn xuất phát từ đầu bể bơi đến cuối bể bơi thì Đức dừng lại nghỉ, còn Triết quay lại và tiếp tục bơi tiếp về đầu bể bơi mới nghỉ. Thời gian bơi của Đức là $25\text{ s}$; Triết bơi đi mất $23\text{ s}$ và bơi về mất $27\text{ s}$. Biết rằng Đức và Triết bơi trên một đường thẳng song song với thành bể.

\textbf{Lời giải.} Chiều dài bể bơi là $L = 25\text{ m}$.

- \textbf{a) Đúng.} Quãng đường Đức bơi được là $s_{\text{Đ}} = 25\text{ m}$, thời gian bơi $t_{\text{Đ}} = 25\text{ s}$. Tốc độ trung bình của Đức là:
\[
v_{tb(\text{Đ})} = \frac{s_{\text{Đ}}}{t_{\text{Đ}}} = \frac{25}{25} = 1\text{ m/s}.
\]

- \textbf{b) Sai.} Tổng thời gian bơi của Triết là $t_T = 23 + 27 = 50\text{ s}$. Vì Triết bơi đi rồi bơi về lại vị trí xuất phát nên độ dịch chuyển tổng hợp $d_T = 0\text{ m}$. Vận tốc trung bình của Triết là:
\[
v_{tb} = \frac{d_T}{t_T} = \frac{0}{50} = 0\text{ m/s} \ne 1\text{ m/s}.
\]

- \textbf{c) Sai.} Vì vị trí xuất phát và vị trí kết thúc của Triết trùng nhau (đầu bể bơi) nên độ dịch chuyển tổng hợp của Triết có độ lớn bằng $0\text{ m}$, không phải $50\text{ m}$.

- \textbf{d) Đúng.} Đức bơi từ đầu bể đến cuối bể nên quãng đường Đức bơi được bằng đúng chiều dài bể bơi là $25\text{ m}$.

% Phần II Câu 2
\textbf{Câu 2.} Một máy bay Boeing 747 bắt đầu xuất phát trên một đường băng thẳng dài với gia tốc $a = 2\text{ m/s}^2$. Tốc độ cần thiết để máy bay rời đường băng là $360\text{ km/h}$. Sau khi xuất phát được $30\text{ s}$ thì nhận được lệnh huỷ bay từ trạm điều khiển không lưu do thời tiết xấu nên phi công cho máy bay hãm phanh chuyển động chậm dần đều và dừng lại sau $20\text{ s}$ kể từ khi có lệnh huỷ bay.

\textbf{Lời giải.} 

- \textbf{a) Đúng.} Trong $30\text{ s}$ đầu tiên, máy bay xuất phát từ trạng thái nghỉ ($v_0 = 0$) với gia tốc không đổi $a_1 = 2\text{ m/s}^2 > 0$. Vận tốc tăng dần đều theo thời gian nên đây là chuyển động thẳng nhanh dần đều.

- \textbf{b) Đúng.} Vận tốc của máy bay tại thời điểm nhận được lệnh huỷ bay ($t_1 = 30\text{ s}$) là:
\[
v_1 = v_0 + a_1 t_1 = 0 + 2 \cdot 30 = 60\text{ m/s}.
\]

- \textbf{c) Đúng.} Giai đoạn hãm phanh từ vận tốc $v_1 = 60\text{ m/s}$ đến dừng hẳn ($v_2 = 0$) trong thời gian $t_2 = 20\text{ s}$. Gia tốc của máy bay trong giai đoạn này:
\[
a_2 = \frac{v_2 - v_1}{t_2} = \frac{0 - 60}{20} = -3\text{ m/s}^2.
\]
Độ lớn gia tốc hãm phanh là $|a_2| = 3\text{ m/s}^2$.

- \textbf{d) Sai.} Quãng đường chạy trong $30\text{ s}$ đầu: $s_1 = \dfrac{1}{2} a_1 t_1^2 = \dfrac{1}{2} \cdot 2 \cdot 30^2 = 900\text{ m}$.
Quãng đường hãm phanh: $s_2 = \dfrac{v_1 + v_2}{2} \cdot t_2 = \dfrac{60 + 0}{2} \cdot 20 = 600\text{ m}$.
Tổng quãng đường: $s = s_1 + s_2 = 900 + 600 = 1500\text{ m}$.
Tổng thời gian: $t = t_1 + t_2 = 30 + 20 = 50\text{ s}$.
Tốc độ trung bình trong toàn bộ quá trình:
\[
v_{tb} = \frac{s}{t} = \frac{1500}{50} = 30\text{ m/s} \ne 40\text{ m/s}.
\]

% Phần II Câu 3
\textbf{Câu 3.} Từ các đồ thị trong hình sau. Nhận định sau đúng hay sai?

@@CENTER_IMAGE_p2_c3_dothi3@@

\textbf{Lời giải.} 

- \textbf{a) Đúng.} Đồ thị hình a và hình b đều là đường thẳng dốc lên từ trái sang phải, vận tốc có độ lớn tăng đều theo thời gian, biểu diễn chuyển động thẳng nhanh dần đều.

- \textbf{b) Sai.} Ở hình b, đồ thị dốc lên nên hệ số góc $a = \dfrac{\Delta v}{\Delta t} > 0$. Công thức vận tốc có dạng $v = v_0 + at$ với $a > 0$ (chứ không phải $a < 0$).

- \textbf{c) Đúng.} Ở hình c, đồ thị là đường thẳng dốc xuống từ $v_0$ về 0, độ lớn vận tốc giảm đều theo thời gian, biểu diễn chuyển động thẳng chậm dần đều.

- \textbf{d) Sai.} Ở hình c, đồ thị dốc xuống nên hệ số góc $a = \dfrac{\Delta v}{\Delta t} < 0$. Công thức vận tốc có dạng $v = v_0 + at$ với $a < 0$ (chứ không phải $a > 0$).

% Phần II Câu 4
\textbf{Câu 4.} Bỏ qua sức cản không khí, một vật rơi tự do tại một địa điểm có độ cao $180\text{ m}$, lấy $g = 10\text{ m/s}^2$.

\textbf{Lời giải.} 

- \textbf{a) Sai.} Quãng đường vật rơi được sau $3\text{ s}$ là:
\[
s_3 = \frac{1}{2}gt_3^2 = \frac{1}{2} \cdot 10 \cdot 3^2 = 45\text{ m} \ne 30\text{ m}.
\]

- \textbf{b) Đúng.} Chuyển động rơi tự do luôn có phương thẳng đứng, chiều hướng từ trên xuống dưới.

- \textbf{c) Đúng.} Để rơi hết $100\text{ m}$ cuối cùng, trước đó vật rơi quãng đường là $s_1 = 180 - 100 = 80\text{ m}$.
Thời gian rơi $80\text{ m}$ đầu tiên:
\[
t_1 = \sqrt{\frac{2s_1}{g}} = \sqrt{\frac{2 \cdot 80}{10}} = \sqrt{16} = 4\text{ s}.
\]
Tổng thời gian rơi hết độ cao $180\text{ m}$:
\[
t = \sqrt{\frac{2h}{g}} = \sqrt{\frac{2 \cdot 180}{10}} = \sqrt{36} = 6\text{ s}.
\]
Thời gian rơi hết $100\text{ m}$ cuối cùng là: $\Delta t = t - t_1 = 6 - 4 = 2\text{ s}$.

- \textbf{d) Sai.} Thời gian rơi hết toàn bộ quãng đường là $t = 6\text{ s}$, không phải $9\text{ s}$.

@@SECTION_3_HEADER@@

% Phần III Câu 1
\textbf{Câu 1.} Cho đồ thị độ dịch chuyển theo thời gian của một vật chuyển động như hình dưới đây. Hỏi tại thời điểm $t$ bằng bao nhiêu giây thì vật đổi chiều chuyển động?

@@CENTER_IMAGE_p3_c1_dothi@@

\textbf{Lời giải.} Từ đồ thị $d - t$:
- Trong khoảng $t \in [0; 2\text{ s}]$, độ dịch chuyển tăng từ $0$ lên $30\text{ m}$, vận tốc $v = \dfrac{30 - 0}{2 - 0} = 15\text{ m/s} > 0$ (vật chuyển động theo chiều dương).
- Trong khoảng $t \in [2\text{ s}; 4\text{ s}]$, độ dịch chuyển giảm từ $30\text{ m}$ về $20\text{ m}$, vận tốc $v' = \dfrac{20 - 30}{4 - 2} = -5\text{ m/s} < 0$ (vật chuyển động theo chiều âm).
Như vậy, tại thời điểm $t = 2\text{ s}$, độ dịch chuyển đạt cực đại và vận tốc đổi dấu từ dương sang âm, tức là vật đổi chiều chuyển động.

\textbf{==> Đáp án: 2.}

% Phần III Câu 2
\textbf{Câu 2.} Chạy marathon là một bộ môn thể thao dưới hình thức chạy bộ được ưa chuộng nhất trên thế giới, là một môn thể thao mang đến rất nhiều lợi ích cho sức khỏe. Người tham gia có thể chạy ở các cự ly $5\text{ km}$, $10\text{ km}$ hoặc $21\text{ km}$ rồi nâng lên đường chạy dài $42\text{ km}$ trong marathon. Tùy thuộc vào thể trạng mà chúng ta lựa chọn luyện tập ở cự ly phù hợp.

Pace là một từ tiếng Anh được sử dụng để chỉ nhịp điệu trong chạy bộ. Pace là số phút để một người hoàn thành $1\text{ km}$ trong khi chạy. Bạn Hương mới tập luyện nên chọn cự ly $5\text{ km}$, trong một lần tập bạn đạt được thời gian hoàn thành cự ly $5\text{ km}$ là 42 phút. Pace trung bình trong bài luyện tập chạy bộ của bạn Hương là bao nhiêu phút/km?

\textbf{Lời giải.} Pace trung bình được tính bằng tỉ số giữa thời gian chạy (tính bằng phút) và quãng đường (tính bằng km):
\[
\text{Pace} = \frac{t}{s} = \frac{42\text{ phút}}{5\text{ km}} = 8{,}4\text{ phút/km}.
\]

\textbf{==> Đáp án: 8,4.}

% Phần III Câu 3
\textbf{Câu 3.} Đồ thị bên dưới mô tả sự thay đổi vận tốc theo thời gian trong chuyển động của một ô tô. Gia tốc của ô tô từ $t = 15\text{ s}$ đến $t = 30\text{ s}$ có giá trị bao nhiêu $\text{m/s}^2$?

@@CENTER_IMAGE_p3_c3_dothi@@

\textbf{Lời giải.} Từ đồ thị vận tốc -- thời gian ($v - t$):
- Tại thời điểm $t_1 = 15\text{ s}$, vận tốc của ô tô là $v_1 = 30\text{ m/s}$.
- Tại thời điểm $t_2 = 30\text{ s}$, vận tốc của ô tô là $v_2 = 0\text{ m/s}$.
Gia tốc của ô tô trong khoảng thời gian này là:
\[
a = \frac{v_2 - v_1}{t_2 - t_1} = \frac{0 - 30}{30 - 15} = \frac{-30}{15} = -2\text{ m/s}^2.
\]

\textbf{==> Đáp án: -2.}

% Phần III Câu 4
\textbf{Câu 4.} Một học sinh đo chiều dài của bàn học, kết quả thu được như sau $d = 120 \pm 1\text{ cm}$. Sai số tuyệt đối của phép đo trên là bao nhiêu cm?

\textbf{Lời giải.} Kết quả đo được viết dưới dạng $d = \overline{d} \pm \Delta d$. Với kết quả $d = 120 \pm 1\text{ cm}$, giá trị trung bình là $\overline{d} = 120\text{ cm}$ và sai số tuyệt đối của phép đo là $\Delta d = 1\text{ cm}$.

\textbf{==> Đáp án: 1.}

% Phần III Câu 5
\textbf{Câu 5.} Đồ thị bên dưới mô tả sự thay đổi vận tốc theo thời gian trong chuyển động của một vật. Vật thay đổi tính chất chuyển động từ chuyển động thẳng đều sang chuyển động thẳng biến đổi đều ở thời điểm nào (bao nhiêu giây)?

@@CENTER_IMAGE_p3_c5_dothi@@

\textbf{Lời giải.} Phân tích tính chất chuyển động trên từng đoạn của đồ thị:
- Đoạn $AB$ ($0 \le t \le 40\text{ s}$): Đồ thị dốc lên, vận tốc tăng đều từ $40\text{ cm/s}$ đến $120\text{ cm/s}$ $\implies$ Chuyển động thẳng nhanh dần đều.
- Đoạn $BD$ ($40\text{ s} \le t \le 80\text{ s}$): Đồ thị là đoạn nằm ngang, vận tốc không đổi $v = 120\text{ cm/s}$ $\implies$ Chuyển động thẳng đều.
- Đoạn $DF$ ($80\text{ s} \le t \le 160\text{ s}$): Đồ thị dốc xuống, vận tốc giảm đều từ $120\text{ cm/s}$ về $0$ $\implies$ Chuyển động thẳng chậm dần đều (là chuyển động thẳng biến đổi đều).
Vậy tại điểm $D$ ứng với thời điểm $t = 80\text{ s}$, vật chuyển tính chất chuyển động từ thẳng đều sang biến đổi đều.

\textbf{==> Đáp án: 80.}

% Phần III Câu 6
\textbf{Câu 6.} Bạn Tuấn chạy bộ qua cầu đi thẳng $40\text{ m}$ theo hướng Đông, sau đó rẽ trái chạy thẳng theo hướng Bắc $100\text{ m}$ rồi quay ngược về hướng Nam $70\text{ m}$. Độ lớn độ dịch chuyển của Tuấn là bao nhiêu m?

\textbf{Lời giải.} Chọn hệ trục tọa độ $Oxy$ gắn với mặt phẳng ngang: gốc $O$ tại vị trí xuất phát, trục $Ox$ hướng Đông, trục $Oy$ hướng Bắc.
- Đi thẳng $40\text{ m}$ hướng Đông: $\vec{d}_1 = (40; 0)\text{ m}$.
- Chạy thẳng $100\text{ m}$ hướng Bắc: $\vec{d}_2 = (0; 100)\text{ m}$.
- Quay ngược về hướng Nam $70\text{ m}$: $\vec{d}_3 = (0; -70)\text{ m}$.
Vectơ độ dịch chuyển tổng hợp của Tuấn:
\[
\vec{d} = \vec{d}_1 + \vec{d}_2 + \vec{d}_3 = (40 + 0 + 0;\; 0 + 100 - 70) = (40; 30)\text{ m}.
\]
Độ lớn độ dịch chuyển:
\[
d = |\vec{d}| = \sqrt{40^2 + 30^2} = \sqrt{1600 + 900} = \sqrt{2500} = 50\text{ m}.
\]

\textbf{==> Đáp án: 50.}

\vspace{0.4cm}
\begin{center}
\textbf{--------- HẾT ---------}
\end{center}

\end{document}
"""

def build_nguyen_hue_doc(tex_content, output_path, is_solution=False):
    scratch_dir = os.path.join(CURRENT_DIR, "scratch_nguyen_hue_10")
    os.makedirs(scratch_dir, exist_ok=True)
    
    file_type = "hdg" if is_solution else "de"
    temp_tex = os.path.join(scratch_dir, f"temp_{file_type}_101.tex")
    temp_docx = os.path.join(scratch_dir, f"temp_{file_type}_101.docx")
    
    with open(temp_tex, "w", encoding="utf-8") as f:
        f.write(tex_content)
        
    cmd = [pandoc_exe, temp_tex, "-o", temp_docx, "--from=latex", "--to=docx"]
    print(f"[*] Đang xuất bản Pandoc: {os.path.basename(output_path)}...")
    res = subprocess.run(cmd, capture_output=True, text=True, encoding="utf-8")
    if res.returncode != 0:
        print("[LỖI PANDOC]:", res.stderr)
        return
        
    doc = docx.Document(temp_docx)
    
    # Thiết lập lề trang
    for s in doc.sections:
        s.top_margin = Cm(1.6)
        s.bottom_margin = Cm(1.6)
        s.left_margin = Cm(2.0)
        s.right_margin = Cm(1.5)
        s.page_width = Cm(21.0)
        s.page_height = Cm(29.7)
        
        # Footer
        footer = s.footer
        f_p = footer.paragraphs[0]
        f_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        f_p.text = ""
        r_l = f_p.add_run("Lớp Toán Cô Thúy  •  SĐT: 0935.322.328  •  50/2C Phạm Thị Liên                                   ")
        r_l.font.name = "Times New Roman"
        r_l.font.size = Pt(9.5)
        r_l.italic = True
        r_l.font.color.rgb = RGBColor(100, 100, 100)
        
        r_r = f_p.add_run("Trang Mã đề 101")
        r_r.font.name = "Times New Roman"
        r_r.font.size = Pt(9.5)
        r_r.italic = True
        r_r.font.color.rgb = RGBColor(80, 80, 80)
        
    style = doc.styles['Normal']
    style.font.name = 'Times New Roman'
    style.font.size = Pt(11)
    style.font.color.rgb = RGBColor(0, 0, 0)
    rPr = style.element.get_or_add_rPr()
    rFonts = parse_xml(r'<w:rFonts %s w:ascii="Times New Roman" w:hAnsi="Times New Roman" w:cs="Times New Roman"/>' % nsdecls('w'))
    rPr.append(rFonts)
    
    # Duyệt và thay thế Header sections
    for p in list(doc.paragraphs):
        if "@@DOCUMENT_HEADER@@" in p.text:
            insert_header_and_code(doc, p, is_solution)
        elif "@@SECTION_1_HEADER@@" in p.text:
            insert_section_header(doc, p, "PHẦN I. Câu trắc nghiệm nhiều phương án lựa chọn.", "Thí sinh trả lời từ câu 1 đến câu 18. Mỗi câu hỏi thí sinh chỉ chọn một phương án.")
        elif "@@SECTION_2_HEADER@@" in p.text:
            insert_section_header(doc, p, "PHẦN II. Câu trắc nghiệm đúng sai.", "Thí sinh trả lời từ câu 1 đến câu 4. Trong mỗi ý a), b), c), d) ở mỗi câu, thí sinh chọn đúng hoặc sai.")
        elif "@@SECTION_3_HEADER@@" in p.text:
            insert_section_header(doc, p, "PHẦN III. Câu trắc nghiệm yêu cầu trả lời ngắn.", "Thí sinh trả lời từ câu 1 đến câu 6.")

    # Chèn các ảnh căn giữa
    img_map = {
        "c7_rac": (os.path.join(HINH_ANH_DIR, "c7_rac.png"), Cm(1.3)),
        "c8_dothidt": (os.path.join(HINH_ANH_DIR, "c8_dothidt.png"), Cm(4.8)),
        "c15_dothi4": (os.path.join(HINH_ANH_DIR, "c15_dothi4.png"), Cm(13.5)),
        "p2_c3_dothi3": (os.path.join(HINH_ANH_DIR, "p2_c3_dothi3.png"), Cm(12.5)),
        "p3_c1_dothi": (os.path.join(HINH_ANH_DIR, "p3_c1_dothi.png"), Cm(7.0)),
        "p3_c3_dothi": (os.path.join(HINH_ANH_DIR, "p3_c3_dothi.png"), Cm(7.6)),
        "p3_c5_dothi": (os.path.join(HINH_ANH_DIR, "p3_c5_dothi.png"), Cm(8.2)),
    }
    
    for p in list(doc.paragraphs):
        for img_tag, (img_file, img_w) in img_map.items():
            token = f"@@CENTER_IMAGE_{img_tag}@@"
            if token in p.text:
                p.text = ""
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.space_after = Pt(4)
                run = p.add_run()
                if os.path.exists(img_file):
                    run.add_picture(img_file, width=img_w)
                break
                
        if "@@ANSWER_BOX@@" in p.text:
            tbl = create_answer_box_table(doc)
            p._p.addprevious(tbl._tbl)
            p._p.getparent().remove(p._p)

    # Định dạng các đoạn văn bản (tab, font, màu sắc)
    for p in doc.paragraphs:
        p_text = p.text.strip()
        
        # Format lời giải
        if is_solution:
            if p_text.startswith("Lời giải."):
                for r in p.runs:
                    if "Lời giải." in r.text:
                        r.bold = True
                        r.font.color.rgb = RGBColor(31, 73, 125)
            elif "==> Chọn" in p_text or "==> Đáp án:" in p_text:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                for r in p.runs:
                    r.bold = True
                    r.font.size = Pt(11)
                    r.font.color.rgb = RGBColor(192, 0, 0)
            elif p_text.startswith("- a) Đúng.") or p_text.startswith("- b) Đúng.") or p_text.startswith("- c) Đúng.") or p_text.startswith("- d) Đúng.") or \
                 p_text.startswith("- a) Sai.") or p_text.startswith("- b) Sai.") or p_text.startswith("- c) Sai.") or p_text.startswith("- d) Sai."):
                for r in p.runs:
                    if "Đúng." in r.text:
                        r.bold = True
                        r.font.color.rgb = RGBColor(0, 128, 0)
                    elif "Sai." in r.text:
                        r.bold = True
                        r.font.color.rgb = RGBColor(192, 0, 0)

        format_doc_paragraph(p)
        
    doc.save(output_path)
    print(f"[THÀNH CÔNG] Đã lưu tài liệu Word: {os.path.basename(output_path)}")

def main():
    print("==================================================")
    print(" BẮT ĐẦU XUẤT BẢN WORD CHO ĐỀ NGUYỄN HUỆ (MÃ 101) ")
    print("==================================================")
    
    # 1. Build Đề docx
    build_nguyen_hue_doc(tex_de, OUTPUT_DE_DOCX, is_solution=False)
    
    # 2. Build HDG docx
    build_nguyen_hue_doc(tex_hdg, OUTPUT_HDG_DOCX, is_solution=True)
    
    # 3. Copy compiled PDF to San_Pham
    pdf_de_src = os.path.join(LATEX_DIR, "PHY10_Nguyen_Hue_De_101.pdf")
    pdf_hdg_src = os.path.join(LATEX_DIR, "PHY10_Nguyen_Hue_HDG_101.pdf")
    pdf_de_dst = os.path.join(SAN_PHAM_DIR, "PHY10_Nguyen_Hue_De_101.pdf")
    pdf_hdg_dst = os.path.join(SAN_PHAM_DIR, "PHY10_Nguyen_Hue_HDG_101.pdf")
    
    if os.path.exists(pdf_de_src):
        shutil.copy2(pdf_de_src, pdf_de_dst)
        print(f"[THÀNH CÔNG] Đã sao chép PDF Đề: {os.path.basename(pdf_de_dst)}")
    if os.path.exists(pdf_hdg_src):
        shutil.copy2(pdf_hdg_src, pdf_hdg_dst)
        print(f"[THÀNH CÔNG] Đã sao chép PDF HDG: {os.path.basename(pdf_hdg_dst)}")
        
    print("\nHOÀN TẤT TẤT CẢ SẢN PHẨM MÃ ĐỀ 101!")

if __name__ == "__main__":
    main()
