import os
import sys
import subprocess
import docx
from docx.shared import Pt, Cm, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn

if sys.stdout.encoding.lower() != 'utf-8':
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except Exception:
        pass

CURRENT_DIR = os.path.dirname(os.path.abspath(__file__))
BASE_DIR = os.path.abspath(os.path.join(CURRENT_DIR, "..", ".."))
HINH_ANH_DIR = os.path.join(BASE_DIR, "He_Thong", "Hinh_Anh")
SAN_PHAM_DIR = os.path.join(BASE_DIR, "San_Pham")
OUTPUT_DOCX = os.path.join(SAN_PHAM_DIR, "De_Thi_2009.docx")

os.makedirs(SAN_PHAM_DIR, exist_ok=True)

# 1. SOẠN THẢO MÃ NGUỒN CHO PANDOC VỚI CÁC TOKEN ĐÁNH DẤU CHUẨN FORM GIÁO VIÊN
# Chuẩn: Dấu @@TAB@@ ở đầu mỗi dòng phương án và giữa các phương án
tex_content = r"""\documentclass[12pt,a4paper]{article}
\usepackage[utf8]{vietnam}
\usepackage{amsmath,amssymb}
\usepackage{graphicx}
\usepackage[top=1.6cm,bottom=1.6cm,left=2.0cm,right=1.5cm]{geometry}

\begin{document}

@@DOCUMENT_HEADER@@

@@SECTION_1_HEADER@@

% CÂU 1 (4 dòng riêng biệt cho A, B, C, D có TAB đầu dòng)
\textbf{Câu 1.} Cho hàm số $y = f(x)$ có đạo hàm trên $\mathbb{R}$ thỏa $f'(x) < 0$, $\forall x \in (1; 2)$ và $f'(x) > 0$, $\forall x \in (2; 3)$. Phát biểu nào sau đây là đúng?

@@TAB@@\textbf{A.} Hàm số $y = f(x)$ đồng biến trên cả hai khoảng $(1; 2)$ và $(2; 3)$.

@@TAB@@\textbf{B.} Hàm số $y = f(x)$ nghịch biến trên cả hai khoảng $(1; 2)$ và $(2; 3)$.

@@TAB@@\textbf{C.} Hàm số $y = f(x)$ đồng biến trên khoảng $(1; 2)$ và nghịch biến trên khoảng $(2; 3)$.

@@TAB@@\textbf{D.} Hàm số $y = f(x)$ nghịch biến trên khoảng $(1; 2)$ và đồng biến trên khoảng $(2; 3)$.

% CÂU 2 (1 dòng 4 phương án dùng TAB chuẩn xác)
\textbf{Câu 2.} Giá trị cực đại của hàm số $f(x) = 2x^3 - 9x^2 - 24x + 1$ là

@@TAB@@\textbf{A.} $-1$.@@TAB@@\textbf{B.} $14$.@@TAB@@\textbf{C.} $4$.@@TAB@@\textbf{D.} $-111$.

% CÂU 3 (2 dòng x 2 phương án dùng TAB chuẩn xác)
\textbf{Câu 3.} Cho lăng trụ $ABC.A'B'C'$. Khẳng định nào sau đây đúng?

@@TAB@@\textbf{A.} $\overrightarrow{BA} + \overrightarrow{A'C'} = \overrightarrow{BC}$.@@TAB@@\textbf{B.} $\overrightarrow{BA} + \overrightarrow{A'C'} = \overrightarrow{BC'}$.

@@TAB@@\textbf{C.} $\overrightarrow{BA} + \overrightarrow{A'C'} = \overrightarrow{C'B}$.@@TAB@@\textbf{D.} $\overrightarrow{BA} + \overrightarrow{A'C'} = \overrightarrow{B'C}$.

% CÂU 4 (Bên trái: đề + 2 dòng phương án TAB; Bên phải: hình ảnh đồ thị)
@@START_SIDE_BY_SIDE_CAU4@@

\textbf{Câu 4.} Hàm số $y = f(x)$ xác định trên đoạn $[-1; 6]$ và có đồ thị như hình vẽ. Hàm số đã cho nghịch biến trên khoảng nào sau đây?

@@TAB@@\textbf{A.} $(-1; 2)$.@@TAB@@\textbf{B.} $(0; 2)$.

@@TAB@@\textbf{C.} $(2; 6)$.@@TAB@@\textbf{D.} $(-2; 0)$.

@@END_SIDE_BY_SIDE_CAU4@@

% CÂU 5 (2 dòng x 2 phương án dùng TAB chuẩn xác)
\textbf{Câu 5.} Cho tứ diện $ABCD$. Lấy $G$ là trọng tâm của tam giác $ABC$. Phát biểu nào sau đây là sai?

@@TAB@@\textbf{A.} $\overrightarrow{GA} + \overrightarrow{GB} + \overrightarrow{GC} = \overrightarrow{0}$.@@TAB@@\textbf{B.} $\overrightarrow{GA} + \overrightarrow{GB} + \overrightarrow{GC} + \overrightarrow{GD} = \overrightarrow{0}$.

@@TAB@@\textbf{C.} $\overrightarrow{GD} - \overrightarrow{GA} = \overrightarrow{AD}$.@@TAB@@\textbf{D.} $\overrightarrow{DA} + \overrightarrow{DB} + \overrightarrow{DC} = 3\overrightarrow{DG}$.

% CÂU 6 (1 dòng 4 phương án dùng TAB)
\textbf{Câu 6.} Một chiếc hộp hình lập phương $ABCD.A'B'C'D'$ có cạnh bằng $12\text{ cm}$, mặt trên $A'B'C'D'$ không nắp. Có một con kiến ở đỉnh $A$ bên ngoài hộp và một miếng mồi của kiến tại điểm $O$ là tâm đáy $ABCD$ ở bên trong hộp. Quãng đường ngắn nhất mà con kiến tìm đến miếng mồi (làm tròn đến hai chữ số thập phân) là

@@TAB@@\textbf{A.} $32{,}49\text{ (cm)}$.@@TAB@@\textbf{B.} $36{,}29\text{ (cm)}$.@@TAB@@\textbf{C.} $12\text{ (cm)}$.@@TAB@@\textbf{D.} $30{,}59\text{ (cm)}$.

% CÂU 7 (1 dòng 4 phương án dùng TAB)
\textbf{Câu 7.} Giá trị nhỏ nhất của hàm số $f(x) = x^4 - 8x^2 + a$, ($a \in \mathbb{R}$) trên đoạn $[-1; 3]$ bằng

@@TAB@@\textbf{A.} $-6$.@@TAB@@\textbf{B.} $a$.@@TAB@@\textbf{C.} $-16 + a$.@@TAB@@\textbf{D.} $9 + a$.

% CÂU 8 (Bên trái: đề + 2 dòng phương án TAB; Bên phải: hình ảnh đồ thị)
@@START_SIDE_BY_SIDE_CAU8@@

\textbf{Câu 8.} Hàm số $y = f(x)$ xác định trên đoạn $[-1; 5]$ và có đồ thị như hình vẽ. Tập giá trị của hàm số $y = f(x)$ trên đoạn $[-1; 5]$ là

@@TAB@@\textbf{A.} $[-1; 5]$.@@TAB@@\textbf{B.} $[1; 3]$.

@@TAB@@\textbf{C.} $[-1; 3]$.@@TAB@@\textbf{D.} $[1; 5]$.

@@END_SIDE_BY_SIDE_CAU8@@

\newpage
% TRANG 2
% CÂU 9
\textbf{Câu 9.} Mỗi ngày bác Bình đều đi bộ để rèn luyện sức khoẻ. Quãng đường đi bộ mỗi ngày (đơn vị: $\text{km}$) của bác Bình trong 20 ngày được thống kê lại ở bảng sau:

\begin{center}
\begin{tabular}{|l|c|c|c|c|c|}
\hline
\textbf{Quãng đường} & $[2{,}7; 3{,}0)$ & $[3{,}0; 3{,}3)$ & $[3{,}3; 3{,}6)$ & $[3{,}6; 3{,}9)$ & $[3{,}9; 4{,}2)$ \\
\hline
\textbf{Số ngày} & $3$ & $6$ & $5$ & $4$ & $2$ \\
\hline
\end{tabular}
\end{center}

Khoảng biến thiên của mẫu số liệu ghép nhóm trên bằng

@@TAB@@\textbf{A.} $1{,}2$.@@TAB@@\textbf{B.} $0{,}362$.@@TAB@@\textbf{C.} $3{,}39$.@@TAB@@\textbf{D.} $1{,}5$.

% CÂU 10
\textbf{Câu 10.} Trong không gian $Oxyz$, cho hai điểm $A(1; 1; 2)$ và $B(3; 1; 0)$. Trung điểm của đoạn thẳng $AB$ có toạ độ là

@@TAB@@\textbf{A.} $(2; 1; 1)$.@@TAB@@\textbf{B.} $(4; 2; 2)$.@@TAB@@\textbf{C.} $(2; 0; -2)$.@@TAB@@\textbf{D.} $(1; 0; -1)$.

% CÂU 11
\textbf{Câu 11.} Đường tiệm cận xiên của đồ thị hàm số $y = \dfrac{x^2 + 2x - 2}{x - 2}$ là

@@TAB@@\textbf{A.} $y = -x + 3$.@@TAB@@\textbf{B.} $y = x + 3$.@@TAB@@\textbf{C.} $y = x - 3$.@@TAB@@\textbf{D.} $y = x + 4$.

% CÂU 12
\textbf{Câu 12.} Trong không gian $Oxyz$, cho điểm $A(-3; 1; -4)$, $B(1; -5; 2)$. Đường thẳng $AB$ cắt mặt phẳng $(Oxy)$ tại điểm

@@TAB@@\textbf{A.} $M\left(-\dfrac{1}{3}; -3; 0\right)$.@@TAB@@\textbf{B.} $N\left(\dfrac{1}{3}; 3; 0\right)$.

@@TAB@@\textbf{C.} $P(0; 3; 1)$.@@TAB@@\textbf{D.} $Q(-3; 1; 0)$.

@@SECTION_2_HEADER@@

% CÂU 1 PHẦN 2 (Bảng biến thiên crop chuẩn gốc)
\textbf{Câu 1.} Hàm số $y = f(x)$ liên tục trên $\mathbb{R}$ và có bảng biến thiên như sau:

@@CENTER_IMAGE_BBT@@

@@TAB@@\textbf{a)} Đồ thị hàm số đã cho có hai đường tiệm cận ngang.

@@TAB@@\textbf{b)} Giá trị nhỏ nhất của hàm số trên $(-\infty; +\infty)$ bằng $8$.

@@TAB@@\textbf{c)} Hàm số đồng biến trên $(8; 38)$.

@@TAB@@\textbf{d)} Giá trị lớn nhất của hàm số trên $\mathbb{R}$ bằng $142$.

\newpage
% TRANG 3
% CÂU 2 PHẦN 2 (Tam giác)
@@START_SIDE_BY_SIDE_TAMGIAC@@

\textbf{Câu 2.} Xét tam giác $ABC$ có $AC = 2AB$ và $BC = 10\text{ cm}$. Trên cạnh $AC$ lấy điểm $D$ sao cho $AD = \dfrac{1}{4}AC$, trên cạnh $AB$ lấy điểm $E$ sao cho $AE = \dfrac{1}{4}AB$, trên cạnh $AD$ lấy điểm $F$ sao cho $AF = \dfrac{1}{4}AD$ và tiếp tục lấy các điểm $G, H, I, J\dots$ (vô hạn lần) theo quy luật đó. Xét tính đúng sai các mệnh đề sau:

@@TAB@@\textbf{a)} $\dfrac{AB}{AC} = \dfrac{AD}{AB}$.

@@TAB@@\textbf{b)} Tam giác $ABD$ đồng dạng với tam giác $ABC$.

@@TAB@@\textbf{c)} $BD = 5\text{ cm}$; $DE = 3\text{ cm}$.

@@TAB@@\textbf{d)} Độ dài đường gấp khúc $CBDEFGH\dots$ bằng $20\text{ cm}$.

@@END_SIDE_BY_SIDE_TAMGIAC@@

% CÂU 3 PHẦN 2 (Máy bay)
@@START_SIDE_BY_SIDE_MAYBAY@@

\textbf{Câu 3.} Hình vẽ sau mô tả vị trí của máy bay vào thời điểm 9h30 phút. Biết các đơn vị trên hình tính theo đơn vị $\text{km}$. Trong các khẳng định sau đây, khẳng định nào đúng, khẳng định nào sai?

@@TAB@@\textbf{a)} Máy bay đang ở độ cao $9\text{ km}$.

@@TAB@@\textbf{b)} Tọa độ của máy bay $(300; 150; 9)$.

@@TAB@@\textbf{c)} Phi công để máy bay ở chế độ tự động với vận tốc theo hướng đông là $750\text{ km/h}$, độ cao không đổi. Biết rằng gió thổi theo hướng đông với vận tốc $10\text{ m/s}$. Giả sử vận tốc và hướng gió không đổi thì lúc 10h30 phút máy bay ở tọa độ $(150; 1086; 9)$.

@@TAB@@\textbf{d)} Sau khi bay đến vị trí lúc 10h30 thì máy bay bay theo hướng ngược lại với vận tốc $800\text{ km/h}$ với độ cao không đổi, biết lúc đó trời lặng gió thì lúc 11h máy bay ở tọa độ $(686; 150; 9)$.

@@END_SIDE_BY_SIDE_MAYBAY@@

% CÂU 4 PHẦN 2 (Bảng chiều cao)
\textbf{Câu 4.} Khảo sát chiều cao của 20 học sinh nam lớp 12A của một trường THPT X, người ta được kết quả thống kê trong bảng sau:

\begin{center}
\begin{tabular}{|l|c|c|c|c|c|}
\hline
\textbf{Chiều cao (cm)} & $[160; 165)$ & $[165; 170)$ & $[170; 175)$ & $[175; 180)$ & $[180; 185)$ \\
\hline
\textbf{Số học sinh} & $3$ & $5$ & $7$ & $4$ & $1$ \\
\hline
\end{tabular}
\end{center}

@@TAB@@\textbf{a)} Gọi $x_1; x_2; \dots; x_{20}$ là mẫu số liệu gốc gồm chiều cao của 20 học sinh trên được xếp theo thứ tự không giảm. Khi đó, $x_3 \in [165; 170)$ và $x_9 \in [170; 175)$.

@@TAB@@\textbf{b)} Tứ phân vị thứ ba của mẫu số liệu ghép nhóm đã cho bằng $175$.

@@TAB@@\textbf{c)} Khoảng tứ phân vị của mẫu số liệu ghép nhóm đã cho là $\Delta Q = Q_3 - Q_1 = 8{,}5$.

@@TAB@@\textbf{d)} Chọn ngẫu nhiên một học sinh trong nhóm khảo sát nói trên, xác suất chọn được học sinh có chiều cao từ $175\text{ cm}$ trở lên bằng $0{,}25$.

\newpage
% TRANG 4
@@SECTION_3_HEADER@@

% CÂU 1 PHẦN 3
\textbf{Câu 1.} Một doanh nghiệp dự kiến sản xuất không quá $2000$ sản phẩm cùng loại. Giả sử rằng doanh thu của doanh nghiệp khi sản xuất $x$ sản phẩm ($x \in \mathbb{N}$; $0 \leqslant x \leqslant 2000$) là $R(x) = 16x - 0{,}005x^2$ (triệu đồng) và doanh nghiệp phải nộp một khoản thuế là $5\%$ của doanh thu $R(x)$. Chi phí để sản xuất mỗi sản phẩm là bốn triệu đồng. Lợi nhuận của doanh nghiệp (khi sản xuất $x$ sản phẩm) được tính theo công thức: Lợi nhuận bằng doanh thu trừ đi thuế và chi phí. Doanh nghiệp phải sản xuất bao nhiêu sản phẩm để thu được lợi nhuận lớn nhất?

@@ANSWER_BOX_1@@

% CÂU 2 PHẦN 3 (Shipper)
@@START_SIDE_BY_SIDE_SHIPPER@@

\textbf{Câu 2.} Một bác Shipper giao hàng xuất phát từ kho $A$ để lấy hàng và đi giao tất cả các con đường sau đó lại trở về kho $A$ để trả lại những hàng hóa mà khách hàng chưa nhận. Con đường có sơ đồ và thời gian giao hàng (phút) trên mỗi con đường được mô tả trong hình bên. Thời gian ngắn nhất để bác Shipper hoàn thành công việc trên là bao nhiêu phút?

@@ANSWER_BOX_2@@

@@END_SIDE_BY_SIDE_SHIPPER@@

% CÂU 3 PHẦN 3
\textbf{Câu 3.} Trong không gian $Oxyz$, cho tam giác $ABC$ có $A(-4; -1; 2)$, $B(3; 5; -6)$ và $C(a; b; c)$. Biết trung điểm cạnh $AC$ thuộc trục tung, trung điểm cạnh $BC$ thuộc mặt phẳng $(Oxz)$. Tính $T = 2a + b - c$.

@@ANSWER_BOX_3@@

% CÂU 4 PHẦN 3
\textbf{Câu 4.} Cho tập hợp $X = \{1; 2; 3; 4; 5; 6; 7; 8\}$. Gọi $S$ là tập hợp tất cả các số tự nhiên có 4 chữ số được lập từ các chữ số thuộc tập $X$. Chọn ngẫu nhiên một số từ tập hợp $S$. Xác suất để chọn được một số chia hết cho 3 bằng $\dfrac{a}{b}$ (với $a, b \in \mathbb{N}^*$, $\dfrac{a}{b}$ là phân số tối giản). Tính $T = a + b$.

@@ANSWER_BOX_4@@

% CÂU 5 PHẦN 3
\textbf{Câu 5.} Một quần thể vi khuẩn được nuôi cấy trong phòng thí nghiệm. Nồng độ dinh dưỡng $S$ (đơn vị: $\text{mg/}\ell$) thay đổi theo thời gian $t$ giờ ($t \geqslant 0$) được mô hình hóa bởi hàm số: $S(t) = \dfrac{10t + 5}{t + 1}$. Biết tốc độ sinh trưởng $V$ của vi khuẩn phụ thuộc vào nồng độ dinh dưỡng theo hàm số $V(S) = \dfrac{5S}{S + 2}$. Khi thời gian $t$ kéo dài, tốc độ sinh trưởng $V$ tăng dần và ổn định quanh một ngưỡng $K$ nhất định. Hỏi sau bao nhiêu phút thì tốc độ sinh trưởng của vi khuẩn đạt $90\%$ ngưỡng $K$?

@@ANSWER_BOX_5@@

% CÂU 6 PHẦN 3
\textbf{Câu 6.} Trong không gian với hệ trục tọa độ $Oxyz$, mặt phẳng $(P)\colon bcx + acy + abz - abc = 0$ qua điểm $M(2; 4; 8)$ và cắt các tia $Ox, Oy, Oz$ lần lượt tại $A, B, C$ sao cho $OA = 2OB = 4OC$. Tính $T = a + b + c$.

@@ANSWER_BOX_6@@

\vspace{0.4cm}
\begin{center}
\textbf{------ HẾT ------}
\end{center}

\end{document}
"""

temp_tex = os.path.join(CURRENT_DIR, "temp_build.tex")
temp_docx = os.path.join(CURRENT_DIR, "temp_build.docx")

with open(temp_tex, "w", encoding="utf-8") as f:
    f.write(tex_content)

pandoc_exe = os.path.expandvars(r"%LOCALAPPDATA%\Pandoc\pandoc.exe")
print("[1/5] Đang chuyển đổi TeX sang Word qua Pandoc...")
subprocess.run([pandoc_exe, temp_tex, "-o", temp_docx], check=True)

# 2. KHỞI TẠO VÀ XỬ LÝ VỚI PYTHON-DOCX
print("[2/5] Đang cấu hình trang A4 và thông số phông chữ...")
doc = docx.Document(temp_docx)

# Cấu hình lề trang và Footer chuẩn
for s in doc.sections:
    s.top_margin = Cm(1.6)
    s.bottom_margin = Cm(1.6)
    s.left_margin = Cm(2.0)
    s.right_margin = Cm(1.5)
    s.page_width = Cm(21.0)
    s.page_height = Cm(29.7)
    
    footer = s.footer
    f_p = footer.paragraphs[0]
    f_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    f_p.text = ""
    r_l = f_p.add_run("Lớp Toán Cô Thúy  •  SĐT: 0935.322.328  •  50/2C Phạm Thị Liên                                   ")
    r_l.font.name = "Times New Roman"
    r_l.font.size = Pt(9.5)
    r_l.italic = True
    r_l.font.color.rgb = RGBColor(100, 100, 100)
    
    r_r = f_p.add_run("Trang Mã đề 2009")
    r_r.font.name = "Times New Roman"
    r_r.font.size = Pt(9.5)
    r_r.italic = True
    r_r.font.color.rgb = RGBColor(80, 80, 80)

style = doc.styles['Normal']
style.font.name = 'Times New Roman'
style.font.size = Pt(12)
style.font.color.rgb = RGBColor(0, 0, 0)
rPr = style.element.get_or_add_rPr()
rFonts = parse_xml(r'<w:rFonts %s w:ascii="Times New Roman" w:hAnsi="Times New Roman" w:cs="Times New Roman"/>' % nsdecls('w'))
rPr.append(rFonts)

# Helper function tạo bảng Header đúng mẫu giáo viên gửi (2 hàng, bảng có viền chuẩn)
def insert_header_and_code(target_p):
    table = doc.add_table(rows=2, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    # Hàng 1: Merge ô 1 và 2 thành ô phải
    row0 = table.rows[0]
    cell_left = row0.cells[0]
    cell_right = row0.cells[1]
    cell_right.merge(row0.cells[2])
    
    cell_left.width = Cm(9.5)
    cell_right.width = Cm(8.0)
    
    # Cột Trái: Thông tin lớp
    p_l = cell_left.paragraphs[0]
    p_l.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_l.paragraph_format.line_spacing = 1.15
    p_l.paragraph_format.space_before = Pt(2)
    p_l.paragraph_format.space_after = Pt(2)
    r1 = p_l.add_run("LỚP TOÁN CÔ THÚY\n")
    r1.bold = True
    r1.font.name = "Times New Roman"
    r1.font.size = Pt(12)
    r1.font.color.rgb = RGBColor(31, 73, 125)
    
    r2 = p_l.add_run("SĐT: 0935.322.328  •  50/2C Phạm Thị Liên\n--------------------")
    r2.font.name = "Times New Roman"
    r2.font.size = Pt(10)
    
    # Cột Phải: Thông tin môn thi
    p_r = cell_right.paragraphs[0]
    p_r.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_r.paragraph_format.line_spacing = 1.15
    p_r.paragraph_format.space_before = Pt(2)
    p_r.paragraph_format.space_after = Pt(2)
    r3 = p_r.add_run("ĐỀ ÔN TẬP TOÁN -- KHỐI LỚP: 12\n")
    r3.bold = True
    r3.font.name = "Times New Roman"
    r3.font.size = Pt(12)
    r3.font.color.rgb = RGBColor(192, 0, 0)
    
    r4 = p_r.add_run("Thời gian làm bài: 90 phút (không kể phát đề)")
    r4.italic = True
    r4.font.name = "Times New Roman"
    r4.font.size = Pt(10.5)
    
    # Hàng 2: Họ tên, Số báo danh, Mã đề (chuẩn như mẫu người dùng gửi)
    row1 = table.rows[1]
    c_hoten = row1.cells[0]
    c_sbd = row1.cells[1]
    c_made = row1.cells[2]
    
    c_hoten.width = Cm(10.5)
    c_sbd.width = Cm(4.0)
    c_made.width = Cm(3.0)
    
    p_hoten = c_hoten.paragraphs[0]
    p_hoten.paragraph_format.space_before = Pt(2)
    p_hoten.paragraph_format.space_after = Pt(2)
    r_ht = p_hoten.add_run("Họ và tên: ............................................................................")
    r_ht.font.name = "Times New Roman"
    r_ht.font.size = Pt(11)
    
    p_sbd = c_sbd.paragraphs[0]
    p_sbd.paragraph_format.space_before = Pt(2)
    p_sbd.paragraph_format.space_after = Pt(2)
    r_sb = p_sbd.add_run("Số báo danh: .......")
    r_sb.font.name = "Times New Roman"
    r_sb.font.size = Pt(11)
    
    p_md = c_made.paragraphs[0]
    p_md.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_md.paragraph_format.space_before = Pt(2)
    p_md.paragraph_format.space_after = Pt(2)
    r_md = p_md.add_run("Mã đề 2009")
    r_md.bold = True
    r_md.font.name = "Times New Roman"
    r_md.font.size = Pt(11)
    
    # Viền bảng
    tblPr = table._tbl.tblPr
    tblBorders = parse_xml(
        r'<w:tblBorders %s>'
        r'<w:top w:val="single" w:sz="8" w:space="0" w:color="000000"/>'
        r'<w:left w:val="single" w:sz="8" w:space="0" w:color="000000"/>'
        r'<w:bottom w:val="single" w:sz="8" w:space="0" w:color="000000"/>'
        r'<w:right w:val="single" w:sz="8" w:space="0" w:color="000000"/>'
        r'<w:insideH w:val="single" w:sz="6" w:space="0" w:color="000000"/>'
        r'<w:insideV w:val="single" w:sz="6" w:space="0" w:color="000000"/>'
        r'</w:tblBorders>' % nsdecls('w')
    )
    tblPr.append(tblBorders)
    
    target_p._p.addprevious(table._tbl)
    target_p._p.getparent().remove(target_p._p)

# Helper function tạo Section header chuẩn mẫu
def insert_section_header(target_p, sec_title, sec_desc):
    p_sec = doc.add_paragraph()
    p_sec.paragraph_format.space_before = Pt(8)
    p_sec.paragraph_format.space_after = Pt(2)
    r_title = p_sec.add_run(sec_title)
    r_title.bold = True
    r_title.font.name = "Times New Roman"
    r_title.font.size = Pt(12)
    
    if sec_desc:
        r_desc = p_sec.add_run(" " + sec_desc)
        r_desc.font.name = "Times New Roman"
        r_desc.font.size = Pt(11)
        r_desc.italic = True
        
    p_sec.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    target_p._p.addprevious(p_sec._p)
    target_p._p.getparent().remove(target_p._p)

# Helper function tạo ô trả lời ngắn 4 ô vuông chuẩn THPT 2025
def create_answer_box_table():
    tbl = doc.add_table(rows=1, cols=5)
    tbl.alignment = WD_TABLE_ALIGNMENT.LEFT
    
    cell_kq = tbl.rows[0].cells[0]
    cell_kq.width = Cm(1.8) # Đủ rộng để có thụt lề tab
    p_kq = cell_kq.paragraphs[0]
    p_kq.paragraph_format.space_before = Pt(0)
    p_kq.paragraph_format.space_after = Pt(0)
    p_kq.add_run("\t")
    r = p_kq.add_run("KQ:")
    r.bold = True
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)
    
    for col_idx in range(1, 5):
        cell = tbl.rows[0].cells[col_idx]
        cell.width = Cm(0.65)
        tcPr = cell._tc.get_or_add_tcPr()
        borders = parse_xml(
            r'<w:tcBorders %s>'
            r'<w:top w:val="single" w:sz="8" w:space="0" w:color="000000"/>'
            r'<w:left w:val="single" w:sz="8" w:space="0" w:color="000000"/>'
            r'<w:bottom w:val="single" w:sz="8" w:space="0" w:color="000000"/>'
            r'<w:right w:val="single" w:sz="8" w:space="0" w:color="000000"/>'
            r'</w:tcBorders>' % nsdecls('w')
        )
        tcPr.append(borders)
    
    trPr = tbl.rows[0]._tr.get_or_add_trPr()
    trHeight = parse_xml(r'<w:trHeight %s w:val="380" w:hRule="exact"/>' % nsdecls('w'))
    trPr.append(trHeight)
    
    tblPr = tbl._tbl.tblPr
    tblBorders = parse_xml(
        r'<w:tblBorders %s>'
        r'<w:top w:val="none"/>'
        r'<w:left w:val="none"/>'
        r'<w:bottom w:val="none"/>'
        r'<w:right w:val="none"/>'
        r'<w:insideH w:val="none"/>'
        r'<w:insideV w:val="none"/>'
        r'</w:tblBorders>' % nsdecls('w')
    )
    tblPr.append(tblBorders)
    return tbl

# 3. THAY THẾ HEADER VÀ CÁC SECTION BANNERS
print("[3/5] Đang thiết kế Tiêu đề bài thi và Khung nhận diện các phần...")
for p in list(doc.paragraphs):
    if "@@DOCUMENT_HEADER@@" in p.text:
        insert_header_and_code(p)
    elif "@@SECTION_1_HEADER@@" in p.text:
        insert_section_header(p, "PHẦN I. Câu trắc nghiệm nhiều phương án lựa chọn.", "Thí sinh trả lời từ câu 1 đến câu 12. Mỗi câu hỏi thí sinh chỉ chọn một phương án.")
    elif "@@SECTION_2_HEADER@@" in p.text:
        insert_section_header(p, "PHẦN II. Câu trắc nghiệm đúng sai.", "Thí sinh trả lời từ câu 1 đến câu 4. Trong mỗi ý a), b), c), d) ở mỗi câu, thí sinh chọn đúng hoặc sai.")
    elif "@@SECTION_3_HEADER@@" in p.text:
        insert_section_header(p, "PHẦN III. Câu trắc nghiệm yêu cầu trả lời ngắn.", "Thí sinh trả lời từ câu 1 đến câu 6.")

# 4. XỬ LÝ CÁC KHỐI SIDE-BY-SIDE (ĐỒ THỊ/HÌNH VẼ BÊN PHẢI)
print("[4/5] Đang định dạng các khối đồ thị & hình vẽ bên phải...")
side_by_side_configs = [
    ("CAU4", "fig_cau4.png", 4.6),
    ("CAU8", "fig_cau8.png", 4.6),
    ("TAMGIAC", "fig_tamgiac.png", 4.4),
    ("MAYBAY", "fig_maybay.png", 5.0),
    ("SHIPPER", "fig_shipper.png", 5.0),
]

for tag, img_name, img_width_cm in side_by_side_configs:
    start_tag = f"@@START_SIDE_BY_SIDE_{tag}@@"
    end_tag = f"@@END_SIDE_BY_SIDE_{tag}@@"
    
    start_p = None
    end_p = None
    inner_ps = []
    found = False
    
    for p in doc.paragraphs:
        if start_tag in p.text:
            start_p = p
            found = True
            continue
        if end_tag in p.text:
            end_p = p
            found = False
            continue
        if found:
            inner_ps.append(p)
            
    if start_p and end_p:
        tbl = doc.add_table(rows=1, cols=2)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        
        tblPr = tbl._tbl.tblPr
        tblBorders = parse_xml(
            r'<w:tblBorders %s>'
            r'<w:top w:val="none"/>'
            r'<w:left w:val="none"/>'
            r'<w:bottom w:val="none"/>'
            r'<w:right w:val="none"/>'
            r'<w:insideH w:val="none"/>'
            r'<w:insideV w:val="none"/>'
            r'</w:tblBorders>' % nsdecls('w')
        )
        tblPr.append(tblBorders)
        
        cell_0 = tbl.rows[0].cells[0]
        cell_1 = tbl.rows[0].cells[1]
        cell_0.width = Cm(11.8)
        cell_1.width = Cm(5.7)
        cell_0.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        cell_1.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        
        start_p._p.addprevious(tbl._tbl)
        cell_0._tc.remove(cell_0.paragraphs[0]._p)
        
        for p in inner_ps:
            if "@@ANSWER_BOX_2@@" in p.text:
                box_tbl = create_answer_box_table()
                cell_0._tc.append(box_tbl._tbl)
                p._p.getparent().remove(p._p)
            else:
                cell_0._tc.append(p._p)
                
        img_path = os.path.join(HINH_ANH_DIR, img_name)
        cell_1_p = cell_1.paragraphs[0]
        cell_1_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cell_1_p.paragraph_format.space_before = Pt(0)
        cell_1_p.paragraph_format.space_after = Pt(0)
        if os.path.exists(img_path):
            cell_1_p.add_run().add_picture(img_path, width=Cm(img_width_cm))
            
        start_p._p.getparent().remove(start_p._p)
        end_p._p.getparent().remove(end_p._p)

# Chèn bảng biến thiên và các ô điền đáp án ngắn còn lại
for p in list(doc.paragraphs):
    if "@@CENTER_IMAGE_BBT@@" in p.text:
        img_path = os.path.join(HINH_ANH_DIR, "fig_bbt.png")
        p.text = ""
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(6)
        if os.path.exists(img_path):
            p.add_run().add_picture(img_path, width=Cm(11.0))
            
    for num in [1, 3, 4, 5, 6]:
        tag = f"@@ANSWER_BOX_{num}@@"
        if tag in p.text:
            box_tbl = create_answer_box_table()
            p._p.addprevious(box_tbl._tbl)
            p._p.getparent().remove(p._p)
            break

# 5. CHUẨN HOÁ TOÀN DIỆN CĂN LỀ, CANH TAB VÀ JUSTIFIED
print("[5/5] Đang chuẩn hoá Canh lề Justified, Canh Tab Ruler và phông chữ...")

def format_doc_paragraph(p, inside_table_cell=False):
    p_xml = p._p.xml
    text = p.text.strip()
    
    # 5.1. XỬ LÝ CANH TAB CHO CÁC PHƯƠNG ÁN TRẮC NGHIỆM VÀ MỆNH ĐỀ
    if "@@TAB@@" in p_xml:
        tab_count = p_xml.count("@@TAB@@")
        new_xml = p_xml.replace("@@TAB@@", '</w:t><w:tab/><w:t>')
        new_p = parse_xml(new_xml)
        p._p.getparent().replace(p._p, new_p)
        
        pPr = new_p.get_or_add_pPr()
        
        # Cấu hình Tab stops và alignment
        for tabs in pPr.findall(qn('w:tabs')):
            pPr.remove(tabs)
            
        if tab_count >= 4:
            # 4 phương án trên 1 dòng với Tab đầu dòng (Câu 2, 6, 7, 9, 10, 11)
            # Mốc: Tab1=0.75cm (A), Tab2=4.75cm (B), Tab3=9.0cm (C), Tab4=13.25cm (D)
            tabs_xml = (
                r'<w:tabs %s>'
                r'<w:tab w:val="left" w:pos="425"/>'
                r'<w:tab w:val="left" w:pos="2693"/>'
                r'<w:tab w:val="left" w:pos="5103"/>'
                r'<w:tab w:val="left" w:pos="7512"/>'
                r'</w:tabs>' % nsdecls('w')
            )
            pPr.append(parse_xml(tabs_xml))
            for jc in pPr.findall(qn('w:jc')):
                pPr.remove(jc)
            pPr.append(parse_xml(r'<w:jc %s w:val="left"/>' % nsdecls('w')))
            
        elif tab_count == 2:
            # 2 phương án trên 1 dòng với Tab đầu dòng
            if inside_table_cell:
                # Trong ô 11.8cm (Câu 4, 8): Tab1=0.5cm (A/C), Tab2=6.0cm (B/D)
                tabs_xml = (
                    r'<w:tabs %s>'
                    r'<w:tab w:val="left" w:pos="283"/>'
                    r'<w:tab w:val="left" w:pos="3402"/>'
                    r'</w:tabs>' % nsdecls('w')
                )
            else:
                # Toàn trang (Câu 3, 5, 12): Tab1=0.75cm (A/C), Tab2=9.0cm (B/D)
                tabs_xml = (
                    r'<w:tabs %s>'
                    r'<w:tab w:val="left" w:pos="425"/>'
                    r'<w:tab w:val="left" w:pos="5103"/>'
                    r'</w:tabs>' % nsdecls('w')
                )
            pPr.append(parse_xml(tabs_xml))
            for jc in pPr.findall(qn('w:jc')):
                pPr.remove(jc)
            pPr.append(parse_xml(r'<w:jc %s w:val="left"/>' % nsdecls('w')))
            
        elif tab_count == 1:
            # 1 phương án / 1 dòng có Tab đầu dòng (Câu 1 và các ý a, b, c, d Phần II)
            # Thụt dòng Tab=0.75cm, dòng tiếp theo thụt bằng lề (hanging indent)
            tabs_xml = r'<w:tabs %s><w:tab w:val="left" w:pos="425"/></w:tabs>' % nsdecls('w')
            pPr.append(parse_xml(tabs_xml))
            
            # Căn đều 2 bên (Justified) cho câu dài
            for jc in pPr.findall(qn('w:jc')):
                pPr.remove(jc)
            pPr.append(parse_xml(r'<w:jc %s w:val="both"/>' % nsdecls('w')))
            
            # Hanging indent 0.75 cm
            for ind in pPr.findall(qn('w:ind')):
                pPr.remove(ind)
            ind_xml = r'<w:ind %s w:left="425" w:hanging="425"/>' % nsdecls('w')
            pPr.append(parse_xml(ind_xml))
            
        # Spacing
        for sp in pPr.findall(qn('w:spacing')):
            pPr.remove(sp)
        sp_xml = r'<w:spacing %s w:before="0" w:after="50" w:line="276" w:lineRule="auto"/>' % nsdecls('w')
        pPr.append(parse_xml(sp_xml))
        return

    # 5.2. CĂN ĐỀU 2 BÊN (JUSTIFY) CHO CÂU HỎI
    pPr = p._p.get_or_add_pPr()
    if text.startswith("Câu "):
        for jc in pPr.findall(qn('w:jc')):
            pPr.remove(jc)
        pPr.append(parse_xml(r'<w:jc %s w:val="both"/>' % nsdecls('w')))
        for sp in pPr.findall(qn('w:spacing')):
            pPr.remove(sp)
        sp_xml = r'<w:spacing %s w:before="60" w:after="40" w:line="276" w:lineRule="auto"/>' % nsdecls('w')
        pPr.append(parse_xml(sp_xml))
    elif "HẾT" in text:
        for jc in pPr.findall(qn('w:jc')):
            pPr.remove(jc)
        pPr.append(parse_xml(r'<w:jc %s w:val="center"/>' % nsdecls('w')))
        for sp in pPr.findall(qn('w:spacing')):
            pPr.remove(sp)
        sp_xml = r'<w:spacing %s w:before="140" w:after="140" w:line="276" w:lineRule="auto"/>' % nsdecls('w')
        pPr.append(parse_xml(sp_xml))

    # Đảm bảo phông chữ Times New Roman 12pt
    for r in p.runs:
        r.font.name = 'Times New Roman'
        r.font.size = Pt(12)
        r_rPr = r._r.get_or_add_rPr()
        f = parse_xml(r'<w:rFonts %s w:ascii="Times New Roman" w:hAnsi="Times New Roman" w:cs="Times New Roman"/>' % nsdecls('w'))
        r_rPr.append(f)

# Duyệt các đoạn văn ở Body
for p in list(doc.paragraphs):
    format_doc_paragraph(p, inside_table_cell=False)

# Duyệt và định dạng các bảng
for table in doc.tables:
    is_data_table = False
    for row in table.rows:
        for cell in row.cells:
            text = cell.text.strip()
            if 'Quãng đường' in text or 'Số ngày' in text or 'Chiều cao' in text or 'Số học sinh' in text:
                is_data_table = True
                
    if is_data_table:
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        tblPr = table._tbl.tblPr
        tblBorders = parse_xml(
            r'<w:tblBorders %s>'
            r'<w:top w:val="single" w:sz="8" w:space="0" w:color="000000"/>'
            r'<w:left w:val="single" w:sz="8" w:space="0" w:color="000000"/>'
            r'<w:bottom w:val="single" w:sz="8" w:space="0" w:color="000000"/>'
            r'<w:right w:val="single" w:sz="8" w:space="0" w:color="000000"/>'
            r'<w:insideH w:val="single" w:sz="8" w:space="0" w:color="000000"/>'
            r'<w:insideV w:val="single" w:sz="8" w:space="0" w:color="000000"/>'
            r'</w:tblBorders>' % nsdecls('w')
        )
        tblPr.append(tblBorders)
        for row in table.rows:
            for cell in row.cells:
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                for p in cell.paragraphs:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p.paragraph_format.line_spacing = 1.0
                    p.paragraph_format.space_before = Pt(2)
                    p.paragraph_format.space_after = Pt(2)

    for row in table.rows:
        for cell in row.cells:
            for p in list(cell.paragraphs):
                format_doc_paragraph(p, inside_table_cell=True)

# Tắt Compatibility Mode để kích hoạt engine Word 2013-2024
settings = doc.settings.element
compat = parse_xml(
    r'<w:compat %s>'
    r'<w:compatSetting w:name="compatibilityMode" w:uri="http://schemas.microsoft.com/office/word" w:val="15"/>'
    r'</w:compat>' % nsdecls('w')
)
settings.append(compat)

doc.save(OUTPUT_DOCX)
print(f"[THÀNH CÔNG RỰC RỠ] Đã xuất bản file Word đạt chuẩn mẫu giáo viên: {OUTPUT_DOCX}")

if os.path.exists(temp_tex):
    os.remove(temp_tex)
if os.path.exists(temp_docx):
    os.remove(temp_docx)
