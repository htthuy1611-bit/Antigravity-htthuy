# -*- coding: utf-8 -*-
"""
BUILD ĐỀ KIỂM TRA GIỮA KỲ I - VẬT LÍ 10 - THPT HAI BÀ TRƯNG (HUẾ) - MÃ ĐỀ 209
Bản quyền: LỚP TOÁN CÔ THÚY - GV: HỒ THỊ THÚY - SĐT: 0935.322.328

Một nguồn dữ liệu duy nhất (QUESTIONS) -> sinh ra:
  1. LaTeX Đề  (ex_test [dethi])   -> pdflatex -> PDF
  2. LaTeX HDG (ex_test [loigiai]) -> pdflatex -> PDF
  3. Word Đề / HDG (Pandoc -> OMML 100% + python-docx hậu xử lý chuẩn BTPro)
"""
import os
import re
import sys
import shutil
import subprocess
import docx
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

if sys.stdout.encoding.lower() != 'utf-8':
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except Exception:
        pass

CURRENT_DIR = os.path.dirname(os.path.abspath(__file__))
BASE_DIR = os.path.abspath(os.path.join(CURRENT_DIR, "..", ".."))
FIG_DIR = os.path.join(BASE_DIR, "He_Thong", "Hinh_Anh", "HBT_10")
LATEX_DIR = os.path.join(BASE_DIR, "He_Thong", "LaTeX", "HBT_10")
SCRATCH_DIR = os.path.join(CURRENT_DIR, "scratch_hbt_10")
SAN_PHAM_DIR = os.path.join(BASE_DIR, "San_Pham")
PANDOC = os.path.expandvars(r"%LOCALAPPDATA%\Pandoc\pandoc.exe")
for d in (LATEX_DIR, SCRATCH_DIR, SAN_PHAM_DIR):
    os.makedirs(d, exist_ok=True)

MADE = "209"
NAME_DE = f"VL10_HBT_De_{MADE}"
NAME_HDG = f"VL10_HBT_HDG_{MADE}"
SCHOOL = "THPT HAI BÀ TRƯNG"
YEAR = "2025 -- 2026"
YEAR_W = "2025 – 2026"
BRAND = "Lớp Toán Cô Thúy -- SĐT: 0935.322.328 -- 50/2C Phạm Thị Liên"
BRAND_W = "Lớp Toán Cô Thúy  •  SĐT: 0935.322.328  •  50/2C Phạm Thị Liên"
LETTERS = "ABCD"

# ============================================================================
# DỮ LIỆU ĐỀ THI
# ============================================================================
P1 = [
    dict(stem=r"Mỗi ngày bố của bạn An chở bạn ấy từ nhà đến trường mất $20\text{ phút}$. Hôm nay là ngày thi nên bố bạn ấy chạy xe nhanh hơn và đến trường sớm hơn thường ngày $5\text{ phút}$. Biết quãng đường từ nhà bạn An đến trường là $10\text{ km}$. Tốc độ trung bình của xe hôm nay là",
         opts=[r"$40\text{ km/h}$", r"$50\text{ km/h}$", r"$30\text{ km/h}$", r"$60\text{ km/h}$"], ans=0,
         sol=[r"Thời gian chuyển động hôm nay: $t' = 20 - 5 = 15\text{ phút} = 0{,}25\text{ h}$.",
              r"Tốc độ trung bình của xe hôm nay:",
              r"\[v_{\text{tb}} = \frac{s}{t'} = \frac{10}{0{,}25} = 40\text{ km/h}.\]"]),
    dict(stem=r"Hình dưới mô tả đồ thị độ dịch chuyển -- thời gian của hai xe chuyển động thẳng trong cùng một khoảng thời gian. Chọn câu đúng.",
         fig=("fig_cau2", 4.8),
         opts=[r"Tốc độ vật 1 lớn hơn tốc độ vật 2.", r"Vận tốc vật 1 nhỏ hơn vận tốc vật 2.",
               r"Tốc độ vật 2 lớn hơn tốc độ vật 1.", r"Vận tốc vật 1 có thể lớn hơn, nhỏ hơn hoặc bằng vận tốc vật 2."], ans=0,
         sol=[r"Trong đồ thị $d - t$ của chuyển động thẳng đều, độ dốc của đường biểu diễn cho biết vận tốc: $v = \dfrac{\Delta d}{\Delta t}$.",
              r"Đường (1) dốc hơn đường (2): trong cùng thời gian $t_1$ thì $d_1 > d_2$, suy ra $v_1 > v_2$. Vậy tốc độ vật 1 lớn hơn tốc độ vật 2."]),
    dict(stem=r"Hãy sắp xếp các phương pháp nghiên cứu vật lí sau theo đúng tiến trình lịch sử phát triển của vật lí học:\\ (1) Sử dụng phương pháp thực nghiệm để tìm hiểu thế giới tự nhiên.\\ (2) Dựa trên quan sát và suy luận chủ quan để tìm hiểu thế giới tự nhiên.\\ (3) Sử dụng kết hợp các mô hình lí thuyết tìm hiểu thế giới vi mô và thí nghiệm để kiểm chứng.",
         opts=[r"(1) -- (2) -- (3).", r"(2) -- (1) -- (3).", r"(2) -- (3) -- (1).", r"(1) -- (3) -- (2)."], ans=1,
         sol=[r"Tiến trình phát triển của vật lí học gồm ba giai đoạn:",
              r"-- Tiền vật lí: tìm hiểu tự nhiên dựa trên quan sát và suy luận chủ quan (2).",
              r"-- Vật lí cổ điển: phương pháp thực nghiệm do Galilei khởi xướng (1).",
              r"-- Vật lí hiện đại: kết hợp mô hình lí thuyết về thế giới vi mô với thí nghiệm kiểm chứng (3).",
              r"Thứ tự đúng: (2) -- (1) -- (3)."]),
    dict(stem=r"Hai vật có khối lượng $m_1 < m_2$ được thả rơi tự do tại cùng một độ cao trong chân không, vận tốc tương ứng khi chạm đất là $v_1$ và $v_2$. Kết luận nào sau đây đúng?",
         opts=[r"$v_1 \ge v_2$ hoặc $v_1 < v_2$.", r"$v_1 = v_2$.", r"$v_1 > v_2$.", r"$v_1 < v_2$."], ans=1,
         sol=[r"Vận tốc chạm đất của vật rơi tự do $v = \sqrt{2gh}$ không phụ thuộc khối lượng, chỉ phụ thuộc độ cao $h$. Hai vật rơi từ cùng độ cao nên $v_1 = v_2$."]),
    dict(stem=r"Hai vật được thả rơi tự do đồng thời từ hai độ cao khác nhau $h_1$ và $h_2$. Thời gian rơi của vật thứ nhất lớn gấp ba lần thời gian rơi của vật thứ hai. Bỏ qua lực cản của không khí. Tỉ số các độ cao tương ứng là",
         opts=[r"$\dfrac{h_1}{h_2} = 9$.", r"$\dfrac{h_1}{h_2} = 3$.", r"$\dfrac{h_1}{h_2} = \dfrac{1}{9}$.", r"$\dfrac{h_1}{h_2} = \dfrac{1}{3}$."], ans=0,
         sol=[r"Độ cao rơi tự do: $h = \dfrac{1}{2}gt^2$, suy ra",
              r"\[\frac{h_1}{h_2} = \left(\frac{t_1}{t_2}\right)^2 = 3^2 = 9.\]"]),
    dict(stem=r"Trong các trường hợp dưới đây, trị số vận tốc nào là tốc độ trung bình?",
         opts=[r"Xe lửa chạy với tốc độ $40\text{ km/h}$ trên hành trình từ Huế ra Quảng Trị.",
               r"Tốc độ của bạn Long đo được lúc đúng 7 giờ 15 phút là $5\text{ km/h}$.",
               r"Viên đạn bay ra khỏi nòng súng với tốc độ $600\text{ m/s}$.",
               r"Tốc độ của đầu búa máy ngay tại thời điểm va chạm là $8\text{ m/s}$."], ans=0,
         sol=[r"Tốc độ trung bình đặc trưng cho độ nhanh chậm trên cả một quãng đường (hành trình Huế -- Quảng Trị).",
              r"Các trị số đo tại một thời điểm hay một vị trí xác định (lúc 7 giờ 15 phút, khi ra khỏi nòng súng, lúc va chạm) là tốc độ tức thời."]),
    dict(stem=r"Gia tốc rơi tự do được xác định qua phép đo độ cao $h$ và thời gian rơi $t$ theo công thức $g = \dfrac{2h}{t^2}$. Biểu thức xác định sai số tuyệt đối $\Delta g$ của phép đo là",
         opts=[r"$\Delta g = \overline{g}\left(2\dfrac{\Delta h}{\overline{h}} + \dfrac{\Delta t}{\overline{t}}\right)$.",
               r"$\Delta g = \overline{g}\left(\dfrac{\Delta h}{\overline{h}} + \dfrac{\Delta t}{\overline{t}}\right)$.",
               r"$\Delta g = \overline{g}\left(\dfrac{\Delta h}{\overline{h}} + 2\dfrac{\Delta t}{\overline{t}}\right)$.",
               r"$\Delta g = \dfrac{\Delta h}{\Delta t^2}$."], ans=2,
         sol=[r"Với $g = 2h\,t^{-2}$, sai số tỉ đối bằng tổng sai số tỉ đối của các thừa số (nhân với số mũ tương ứng):",
              r"\[\frac{\Delta g}{\overline{g}} = \frac{\Delta h}{\overline{h}} + 2\frac{\Delta t}{\overline{t}} \Rightarrow \Delta g = \overline{g}\left(\frac{\Delta h}{\overline{h}} + 2\frac{\Delta t}{\overline{t}}\right).\]"]),
    dict(stem=r"Một vật chuyển động thẳng nhanh dần đều không vận tốc đầu với phương trình độ dịch chuyển $d = 3t^2\text{ (m)}$. Gọi $v$ là vận tốc của vật tại thời điểm $t$, $S$ là quãng đường vật đi được sau thời gian $t$. Hệ thức đúng là",
         opts=[r"$v^2 = 12S$.", r"$v^2 = 6S$.", r"$v^2 = 4S$.", r"$v^2 = 8S$."], ans=0,
         sol=[r"So sánh $d = v_0t + \dfrac{1}{2}at^2 = 3t^2$ ta được $v_0 = 0$ và $a = 6\text{ m/s}^2$.",
              r"Vật chuyển động thẳng không đổi chiều nên $S = d$. Áp dụng hệ thức độc lập thời gian:",
              r"\[v^2 - v_0^2 = 2aS \Rightarrow v^2 = 2 \cdot 6 \cdot S = 12S.\]"]),
    dict(stem=r"Một vật chuyển động thẳng biến đổi đều dọc theo trục $Ox$. Tại thời điểm $t_0$ vận tốc của vật là $v_0$, tại thời điểm $t$ vận tốc của vật là $v$. Gia tốc của vật được xác định bằng công thức",
         opts=[r"$a = \dfrac{v^2 + v_0^2}{t - t_0}$.", r"$a = \dfrac{v - v_0}{t - t_0}$.", r"$a = \dfrac{v^2 - v_0^2}{t - t_0}$.", r"$a = \dfrac{v + v_0}{t + t_0}$."], ans=1,
         sol=[r"Gia tốc là độ biến thiên vận tốc trong một đơn vị thời gian: $a = \dfrac{\Delta v}{\Delta t} = \dfrac{v - v_0}{t - t_0}$."]),
    dict(stem=r"Khi lắp ráp và sử dụng các thiết bị điện trong phòng thực hành Vật lí, thao tác nào sau đây là yêu cầu bắt buộc và quan trọng nhất?",
         opts=[r"Không cần quan tâm sử dụng đúng chức năng của thiết bị.",
               r"Chỉ cần quan sát sơ bộ hình dáng bên ngoài rồi bật công tắc ngay.",
               r"Quan sát kĩ các kí hiệu cảnh báo và thông số kĩ thuật trên thiết bị để sử dụng đúng chức năng, đúng điện áp.",
               r"Bật nguồn điện trước rồi mới cắm dây tiến hành thí nghiệm."], ans=2,
         sol=[r"Để đảm bảo an toàn cho người và thiết bị, trước khi sử dụng phải quan sát kĩ các kí hiệu cảnh báo, nhãn thông số kĩ thuật để dùng đúng chức năng và đúng điện áp."]),
    dict(stem=r"Cho các phép đo: (1) dùng thước đo chiều cao; (2) dùng đồng hồ bấm giây đo thời gian; (3) đo gia tốc rơi tự do thông qua công thức; (4) đo tốc độ chuyển động thông qua quãng đường và thời gian. Các phép đo gián tiếp là",
         opts=[r"(1), (2), (4).", r"(1), (2).", r"(2), (3), (4).", r"(3), (4)."], ans=3,
         sol=[r"Phép đo trực tiếp: đọc kết quả ngay trên dụng cụ đo -- (1), (2).",
              r"Phép đo gián tiếp: xác định qua công thức liên hệ với các đại lượng đo trực tiếp -- (3) $g = \dfrac{2h}{t^2}$ và (4) $v = \dfrac{s}{t}$."]),
    dict(stem=r"Đồ thị vận tốc -- thời gian $(v - t)$ của một vật chuyển động thẳng dọc theo trục $Ox$ đi qua ba điểm $A$, $B$, $C$ như hình vẽ. Quãng đường vật đi được trong khoảng thời gian từ $0\text{ s}$ đến $5\text{ s}$ bằng",
         fig=("fig_cau12", 5.2),
         opts=[r"$6{,}25\text{ m}$.", r"$7{,}25\text{ m}$.", r"$6\text{ m}$.", r"$12{,}5\text{ m}$."], ans=0,
         sol=[r"Quãng đường bằng tổng diện tích các phần giới hạn bởi đồ thị $v - t$ và trục $Ot$ (lấy giá trị dương):",
              r"-- Từ $0$ đến $2{,}5\text{ s}$ (tam giác $OAB$): $s_1 = \dfrac{1}{2} \cdot 2{,}5 \cdot 2{,}5 = 3{,}125\text{ m}$.",
              r"-- Từ $2{,}5\text{ s}$ đến $5\text{ s}$ (vật đổi chiều, tam giác dưới trục): $s_2 = \dfrac{1}{2} \cdot 2{,}5 \cdot 2{,}5 = 3{,}125\text{ m}$.",
              r"\[s = s_1 + s_2 = 3{,}125 + 3{,}125 = 6{,}25\text{ m}.\]"]),
]

P2 = [
    dict(stem=r"Một quả cầu bắt đầu lăn không vận tốc đầu từ đỉnh $A$ của một dốc thẳng dài $100\text{ m}$, sau $10\text{ s}$ thì tới chân dốc $B$. Từ $B$, quả cầu tiếp tục lăn chậm dần đều trên mặt phẳng ngang thêm $50\text{ m}$ thì dừng lại tại $C$. Chọn chiều dương là chiều chuyển động của quả cầu.",
         fig=("fig_cau13", 10.5),
         items=[r"Gia tốc của quả cầu khi lăn xuống dốc $AB$ bằng $2{,}5\text{ m/s}^2$.",
                r"Vận tốc của quả cầu tại chân dốc $B$ bằng $20\text{ m/s}$.",
                r"Độ lớn gia tốc của quả cầu khi lăn chậm dần đều trên đoạn $BC$ bằng $4\text{ m/s}^2$.",
                r"Tổng thời gian chuyển động của quả cầu từ $A$ đến khi dừng lại tại $C$ bằng $15\text{ s}$."],
         truth=[False, True, True, True],
         sols=[r"Trên dốc $AB$: $s_1 = \dfrac{1}{2}a_1t_1^2 \Rightarrow 100 = \dfrac{1}{2}a_1 \cdot 10^2 \Rightarrow a_1 = 2\text{ m/s}^2 \ne 2{,}5\text{ m/s}^2$.",
               r"$v_B = a_1t_1 = 2 \cdot 10 = 20\text{ m/s}$.",
               r"Trên $BC$: $v_C^2 - v_B^2 = 2a_2s_2 \Rightarrow 0 - 20^2 = 2a_2 \cdot 50 \Rightarrow a_2 = -4\text{ m/s}^2$, độ lớn $4\text{ m/s}^2$.",
               r"Thời gian trên $BC$: $t_2 = \dfrac{v_C - v_B}{a_2} = \dfrac{0 - 20}{-4} = 5\text{ s}$. Tổng thời gian: $t = t_1 + t_2 = 10 + 5 = 15\text{ s}$."]),
    dict(stem=r"Một nhóm học sinh làm thí nghiệm đo tốc độ trung bình của một xe chuyển động thẳng từ $A$ đến $B$. Dùng thước đo được quãng đường $s = 0{,}500 \pm 0{,}002\text{ (m)}$. Dùng đồng hồ hiện số có sai số dụng cụ $0{,}001\text{ s}$ đo thời gian chuyển động, kết quả 4 lần đo ghi ở bảng sau:"
              "\n\n\\begin{center}\\begin{tabular}{|c|c|c|c|c|}\\hline \\textbf{Lần đo} & \\textbf{1} & \\textbf{2} & \\textbf{3} & \\textbf{4} \\\\ \\hline \\textbf{$t$ (s)} & 0,776 & 0,778 & 0,772 & 0,774 \\\\ \\hline\\end{tabular}\\end{center}",
         items=[r"Phép đo tốc độ trung bình trong thí nghiệm này là phép đo trực tiếp.",
                r"Kết quả đo thời gian là $t = 0{,}775 \pm 0{,}003\text{ (s)}$.",
                r"Kết quả đo tốc độ trung bình là $v = 0{,}645 \pm 0{,}005\text{ (m/s)}$.",
                r"Sai số tỉ đối của phép đo tốc độ trung bình là $\delta v = 0{,}5\%$."],
         truth=[False, True, True, False],
         sols=[r"Tốc độ trung bình được tính qua công thức $v = \dfrac{s}{t}$ từ hai đại lượng đo trực tiếp nên là phép đo gián tiếp.",
               r"$\overline{t} = \dfrac{0{,}776 + 0{,}778 + 0{,}772 + 0{,}774}{4} = 0{,}775\text{ s}$; $\overline{\Delta t} = \dfrac{0{,}001 + 0{,}003 + 0{,}003 + 0{,}001}{4} = 0{,}002\text{ s}$. Vậy $\Delta t = 0{,}002 + 0{,}001 = 0{,}003\text{ s}$ và $t = 0{,}775 \pm 0{,}003\text{ s}$.",
               r"$\overline{v} = \dfrac{0{,}500}{0{,}775} \approx 0{,}645\text{ m/s}$; $\delta v = \dfrac{0{,}002}{0{,}500} + \dfrac{0{,}003}{0{,}775} \approx 0{,}0079$; $\Delta v = \overline{v}\cdot\delta v \approx 0{,}005\text{ m/s}$. Vậy $v = 0{,}645 \pm 0{,}005\text{ m/s}$.",
               r"$\delta v \approx 0{,}79\% \approx 0{,}8\% \ne 0{,}5\%$."]),
]

P3 = [
    dict(stem=r"Bạn Khánh Linh chạy bộ trên một đường thẳng trong $10\text{ phút}$. Trong $4\text{ phút}$ đầu, bạn chạy đều với tốc độ $4\text{ m/s}$; thời gian còn lại bạn chạy đều với tốc độ $3\text{ m/s}$. Tốc độ trung bình của bạn Khánh Linh trong $10\text{ phút}$ đó là bao nhiêu $\text{m/s}$?",
         ans="3,4",
         sol=[r"$t_1 = 4\text{ phút} = 240\text{ s}$; $t_2 = 6\text{ phút} = 360\text{ s}$; tổng thời gian $t = 600\text{ s}$.",
              r"Quãng đường: $s = v_1t_1 + v_2t_2 = 4 \cdot 240 + 3 \cdot 360 = 2040\text{ m}$.",
              r"\[v_{\text{tb}} = \frac{s}{t} = \frac{2040}{600} = 3{,}4\text{ m/s}.\]"]),
    dict(stem=r"Trong đợt lũ lụt tại các tỉnh miền Bắc đầu tháng 10 năm 2025, dòng nước lũ chảy với tốc độ khoảng $4\text{ m/s}$ so với bờ. Một ca nô cứu hộ chạy ngược dòng với tốc độ $8\text{ m/s}$ so với dòng nước để tiếp cận một mái nhà có người mắc kẹt cách trạm xuất phát $1920\text{ m}$. Thời gian để ca nô tiếp cận chỗ người bị nạn là bao nhiêu phút?",
         ans="8",
         sol=[r"Vận tốc ca nô so với bờ khi chạy ngược dòng: $v = v_{\text{cn/n}} - v_{\text{n/b}} = 8 - 4 = 4\text{ m/s}$.",
              r"\[t = \frac{s}{v} = \frac{1920}{4} = 480\text{ s} = 8\text{ phút}.\]"]),
    dict(stem=r"Hình vẽ là đồ thị độ dịch chuyển -- thời gian $(d - t)$ của một vật chuyển động thẳng. Quãng đường vật đi được từ thời điểm $t = 0$ đến thời điểm $t = 4\text{ s}$ là bao nhiêu mét?",
         fig=("fig_cau17", 6.8), ans="120",
         sol=[r"-- Từ $0$ đến $2\text{ s}$: độ dịch chuyển tăng từ $0$ lên $60\text{ m}$, vật đi theo chiều dương được $s_1 = 60\text{ m}$.",
              r"-- Từ $2\text{ s}$ đến $4\text{ s}$: độ dịch chuyển giảm từ $60\text{ m}$ về $0$, vật đổi chiều và đi được $s_2 = 60\text{ m}$.",
              r"\[s = s_1 + s_2 = 60 + 60 = 120\text{ m}.\]"]),
    dict(stem=r"Hai giọt nước mưa rơi tự do từ cùng một độ cao $h$, giọt thứ hai rơi sau giọt thứ nhất $1\text{ s}$. Khi giọt thứ nhất vừa chạm đất thì giọt thứ hai còn cách mặt đất $45\text{ m}$. Bỏ qua lực cản không khí, lấy $g = 10\text{ m/s}^2$. Độ cao $h$ bằng bao nhiêu mét?",
         ans="125",
         sol=[r"Gọi $t$ là thời gian rơi của giọt thứ nhất: $h = \dfrac{1}{2}gt^2 = 5t^2$.",
              r"Khi đó giọt thứ hai đã rơi $(t - 1)$ giây: $s_2 = 5(t - 1)^2$.",
              r"\[h - s_2 = 45 \Rightarrow 5t^2 - 5(t - 1)^2 = 45 \Rightarrow 5(2t - 1) = 45 \Rightarrow t = 5\text{ s}.\]",
              r"\[h = 5 \cdot 5^2 = 125\text{ m}.\]"]),
]

P4 = [
    dict(stem=r"Một ô tô chuyển động thẳng đều với tốc độ $80\text{ km/h}$ từ $A$ về hướng Đông. Sau khi đi được $32\text{ km}$, ô tô rẽ trái chuyển động thẳng đều về hướng Bắc với tốc độ $40\text{ km/h}$ trong $36\text{ phút}$ thì đến $B$. Tính:",
         parts=[r"Quãng đường ô tô đi từ $A$ đến $B$.",
                r"Thời gian ô tô đi từ $A$ đến $B$.",
                r"Tốc độ trung bình và độ lớn vận tốc trung bình của ô tô trên cả hành trình."],
         sol=[r"\textbf{a)} $t_2 = 36\text{ phút} = 0{,}6\text{ h}$, $s_2 = v_2t_2 = 40 \cdot 0{,}6 = 24\text{ km}$. Quãng đường: $s = s_1 + s_2 = 32 + 24 = 56\text{ km}$.",
              r"\textbf{b)} $t_1 = \dfrac{s_1}{v_1} = \dfrac{32}{80} = 0{,}4\text{ h}$. Thời gian: $t = t_1 + t_2 = 0{,}4 + 0{,}6 = 1\text{ h}$.",
              r"\textbf{c)} Tốc độ trung bình: $v_{\text{tb}} = \dfrac{s}{t} = \dfrac{56}{1} = 56\text{ km/h}$.",
              r"Hai chặng vuông góc nên độ dịch chuyển: $d = \sqrt{s_1^2 + s_2^2} = \sqrt{32^2 + 24^2} = 40\text{ km}$.",
              r"Độ lớn vận tốc trung bình: $v = \dfrac{d}{t} = \dfrac{40}{1} = 40\text{ km/h}$."]),
    dict(stem=r"Một ô tô đang chạy thẳng đều với tốc độ $72\text{ km/h}$ thì hãm phanh, chuyển động thẳng chậm dần đều và đi thêm $250\text{ m}$ nữa thì dừng hẳn. Chọn chiều dương là chiều chuyển động. Tính:",
         parts=[r"Gia tốc $a$ của ô tô.",
                r"Quãng đường ô tô đi được trong $10\text{ s}$ cuối cùng trước khi dừng hẳn."],
         sol=[r"\textbf{a)} $v_0 = 72\text{ km/h} = 20\text{ m/s}$, $v = 0$. Áp dụng $v^2 - v_0^2 = 2as$:",
              r"\[0 - 20^2 = 2 \cdot a \cdot 250 \Rightarrow a = -0{,}8\text{ m/s}^2.\]",
              r"\textbf{b)} Thời gian hãm phanh: $t = \dfrac{0 - 20}{-0{,}8} = 25\text{ s} > 10\text{ s}$.",
              r"Xét ngược thời gian, $10\text{ s}$ cuối của chuyển động chậm dần đều tương đương $10\text{ s}$ đầu của chuyển động nhanh dần đều không vận tốc đầu với gia tốc $0{,}8\text{ m/s}^2$:",
              r"\[s' = \frac{1}{2}|a|t'^2 = \frac{1}{2} \cdot 0{,}8 \cdot 10^2 = 40\text{ m}.\]"]),
    dict(stem=r"Hai xe xuất phát cùng lúc từ trạng thái nghỉ, chuyển động thẳng nhanh dần đều cùng chiều dương. Hình vẽ là đồ thị quãng đường -- thời gian $(s - t)$ của xe (1) và xe (2). Dựa vào đồ thị, hãy tính giá trị $S_1$.",
         fig=("fig_cau21", 7.0), parts=[],
         sol=[r"Với chuyển động nhanh dần đều không vận tốc đầu: $s = \dfrac{1}{2}at^2$.",
              r"-- Xe (2) đi được $10\text{ m}$ sau $5\text{ s}$: $a_2 = \dfrac{2 \cdot 10}{5^2} = 0{,}8\text{ m/s}^2$ (kiểm tra: tại $t = 8\text{ s}$, $s_2 = \dfrac{1}{2} \cdot 0{,}8 \cdot 8^2 = 25{,}6\text{ m}$, đúng như đồ thị).",
              r"-- Xe (1) đi được $10\text{ m}$ sau $10\text{ s}$: $a_1 = \dfrac{2 \cdot 10}{10^2} = 0{,}2\text{ m/s}^2$.",
              r"Tại $t = 8\text{ s}$, quãng đường xe (1) đi được:",
              r"\[S_1 = \frac{1}{2}a_1t^2 = \frac{1}{2} \cdot 0{,}2 \cdot 8^2 = 6{,}4\text{ m}.\]"]),
]

SECTIONS = [
    ("PHẦN I. Câu trắc nghiệm nhiều phương án lựa chọn (3,0 điểm)",
     "Thí sinh trả lời từ câu 1 đến câu 12. Mỗi câu hỏi thí sinh chỉ chọn một phương án."),
    ("PHẦN II. Câu trắc nghiệm đúng sai (2,0 điểm)",
     "Thí sinh trả lời từ câu 1 đến câu 2. Trong mỗi ý a), b), c), d) ở mỗi câu, thí sinh chọn đúng hoặc sai."),
    ("PHẦN III. Câu trắc nghiệm trả lời ngắn (2,0 điểm)",
     "Thí sinh trả lời từ câu 1 đến câu 4."),
    ("PHẦN IV. Tự luận (3,0 điểm)",
     "Thí sinh trình bày lời giải từ câu 1 đến câu 3."),
]


# ============================================================================
# 1. SINH LATEX (ex_test)
# ============================================================================
def tex_fig(fig):
    name, w = fig
    return "\t\\begin{center}\n\t\t\\includegraphics[width=%.1fcm]{%s.pdf}\n\t\\end{center}\n" % (w, name)


def tex_sol(lines):
    out = []
    for ln in lines:
        if ln.startswith(r"\["):
            out.append("\t\t" + ln)
        else:
            out.append("\t\t" + ln + "\\par")
    return "\n".join(out)


def tex_section(idx):
    title, note = SECTIONS[idx]
    return ("\n\\vspace{0.15cm}\n\\begin{flushleft}\n"
            "\t{\\color{blue!50!green}\\fbox{\\fontfamily{qag}\\bfseries\\selectfont %s}}\n"
            "\t\\textit{\\small (%s)}\n\\end{flushleft}\n\\setcounter{ex}{0}\n\\setcounter{bt}{0}\n" % (title, note))


def tex_header(is_sol):
    if is_sol:
        top = r"{\bfseries\color{red!80!black} HƯỚNG DẪN GIẢI CHI TIẾT GIỮA KỲ I}\\[3pt]" + "\n" + \
              r"\textbf{Môn: VẬT LÍ 10 -- %s}\\[2pt]" % SCHOOL + "\n" + \
              r"\textit{Năm học %s -- Mã đề: %s}" % (YEAR, MADE)
        row2 = (r"\rule{0pt}{14pt}\textbf{Giáo viên:} HỒ THỊ THÚY &" + "\n" +
                r"\rule{0pt}{14pt}\textbf{Môn học:} VẬT LÍ 10 &" + "\n" +
                r"\rule{0pt}{14pt}\textbf{Mã đề thi %s}\rule[-4pt]{0pt}{4pt} \\" % MADE)
    else:
        top = r"{\bfseries\color{red!80!black} KIỂM TRA GIỮA KỲ I -- NĂM HỌC %s}\\[3pt]" % YEAR + "\n" + \
              r"\textbf{Môn: VẬT LÍ, Lớp 10 -- %s}\\[2pt]" % SCHOOL + "\n" + \
              r"\textit{Thời gian làm bài: 45 phút (Không kể thời gian phát đề)}"
        row2 = (r"\rule{0pt}{13pt}Họ và tên thí sinh: \dotfill\rule[-3pt]{0pt}{3pt} &" + "\n" +
                r"\rule{0pt}{13pt}Số báo danh: \dotfill\rule[-3pt]{0pt}{3pt} &" + "\n" +
                r"\rule{0pt}{13pt}\textbf{Mã đề thi %s}\rule[-3pt]{0pt}{3pt} \\" % MADE)
    return r"""\noindent
\begin{tabular}{|p{6.6cm}|p{6.6cm}|>{\centering\arraybackslash}p{3.2cm}|}
\hline
\begin{minipage}{6.6cm}
\vspace{3pt}
\centering
{\large\bfseries\color{blue!80!black} LỚP TOÁN CÔ THÚY}\\[3pt]
\textbf{SĐT:} 0935.322.328\\[2pt]
\textbf{Địa chỉ:} 50/2C Phạm Thị Liên
\vspace{3pt}
\end{minipage}
&
\multicolumn{2}{p{10.2cm}|}{
\begin{minipage}{10.2cm}
\vspace{3pt}
\centering
""" + top + r"""
\vspace{3pt}
\end{minipage}
} \\
\hline
""" + row2 + r"""
\hline
\end{tabular}
\vspace{0.1cm}
"""


def tex_answer_key():
    s = "\n\\vspace{0.3cm}\n\\noindent\\begin{minipage}{\\linewidth}\n\\begin{center}{\\large\\bfseries\\color{red!80!black} BẢNG ĐÁP ÁN -- MÃ ĐỀ %s}\\end{center}\n" % MADE
    s += "\\noindent\\textbf{Phần I.}\\par\\smallskip\n\\begin{center}\\begin{tabular}{|c|" + "c|" * 12 + "}\\hline\n"
    s += "\\textbf{Câu} & " + " & ".join(str(i + 1) for i in range(12)) + " \\\\ \\hline\n"
    s += "\\textbf{Đáp án} & " + " & ".join("\\textbf{%s}" % LETTERS[q["ans"]] for q in P1) + " \\\\ \\hline\n\\end{tabular}\\end{center}\n"
    s += "\\noindent\\textbf{Phần II.}\\par\\smallskip\n\\begin{center}\\begin{tabular}{|c|c|c|c|c|}\\hline\n"
    s += "\\textbf{Câu} & \\textbf{a)} & \\textbf{b)} & \\textbf{c)} & \\textbf{d)} \\\\ \\hline\n"
    for i, q in enumerate(P2):
        s += "%d & " % (i + 1) + " & ".join("Đ" if t else "S" for t in q["truth"]) + " \\\\ \\hline\n"
    s += "\\end{tabular}\\end{center}\n"
    s += "\\noindent\\textbf{Phần III.}\\par\\smallskip\n\\begin{center}\\begin{tabular}{|c|c|c|c|c|}\\hline\n"
    s += "\\textbf{Câu} & 1 & 2 & 3 & 4 \\\\ \\hline\n\\textbf{Đáp án} & " + " & ".join(q["ans"] for q in P3) + " \\\\ \\hline\n\\end{tabular}\\end{center}\n\\end{minipage}\n"
    return s


def build_latex(is_sol):
    opt = "loigiai" if is_sol else "dethi"
    foot_r = (r"Trang \thepage/\pageref{LastPage} -- Hướng dẫn giải Mã đề %s" % MADE) if is_sol \
        else (r"Trang \thepage/\pageref{LastPage} -- Mã đề %s" % MADE)
    s = r"""\documentclass[11pt,a4paper]{article}
\usepackage[utf8]{vietnam}
\usepackage{amsmath,amssymb,mathrsfs}
\usepackage{graphicx}
\usepackage{tikz}
\usepackage{array}
\usepackage[top=1.0cm,bottom=1.8cm,left=1.4cm,right=1.4cm,footskip=0.8cm,headheight=16pt]{geometry}
\usepackage{fancyhdr}
\usepackage{tcolorbox}
\tcbuselibrary{skins}
\usepackage{lastpage}
\usepackage[""" + opt + r"""]{ex_test}
\graphicspath{{../../Hinh_Anh/HBT_10/}}

\makeatletter
\@ifundefined{c@bt}{\newcounter{bt}}{}
\makeatother

\pagestyle{fancy}
\fancyhf{}
\renewcommand{\headrulewidth}{0pt}
\renewcommand{\footrulewidth}{0.4pt}
\lfoot{\footnotesize\textsl{""" + BRAND + r"""}}
\rfoot{\footnotesize\textsl{""" + foot_r + r"""}}
\setlength{\parskip}{1pt}
\setlength{\parindent}{0pt}

\begin{document}

""" + tex_header(is_sol)

    # Phần I
    s += tex_section(0)
    for i, q in enumerate(P1):
        s += "\n%% Câu %d\n\\begin{ex}\n\t%s\n" % (i + 1, q["stem"])
        if q.get("fig"):
            s += tex_fig(q["fig"])
        s += "\t\\choice\n"
        for k, o in enumerate(q["opts"]):
            s += "\t{%s%s}\n" % ("\\True " if k == q["ans"] else "", o.rstrip("."))
        s += "\t\\loigiai{\n%s\n\t}\n\\end{ex}\n" % tex_sol(q["sol"])
    # Phần II
    s += tex_section(1)
    for i, q in enumerate(P2):
        s += "\n%% Phần II Câu %d\n\\begin{ex}\n\t%s\n" % (i + 1, q["stem"])
        if q.get("fig"):
            s += tex_fig(q["fig"])
        s += "\t\\choiceTF[t]\n"
        for t, it in zip(q["truth"], q["items"]):
            s += "\t{%s%s}\n" % ("\\True " if t else "", it)
        lines = []
        for k, (t, so) in enumerate(zip(q["truth"], q["sols"])):
            tag = "Đúng" if t else "Sai"
            col = "green!50!black" if t else "red!80!black"
            lines.append(r"{\color{%s}\textbf{%s) %s.}} %s" % (col, "abcd"[k], tag, so))
        s += "\t\\loigiai{\n%s\n\t}\n\\end{ex}\n" % tex_sol(lines)
    # Phần III
    s += tex_section(2)
    for i, q in enumerate(P3):
        s += "\n%% Phần III Câu %d\n\\begin{ex}\n\t%s\n" % (i + 1, q["stem"])
        if q.get("fig"):
            s += tex_fig(q["fig"])
        s += "\t\\par\\shortans[oly]{%s}\n" % q["ans"]
        s += "\t\\loigiai{\n%s\n\t}\n\\end{ex}\n" % tex_sol(q["sol"])
    # Phần IV
    s += tex_section(3)
    for i, q in enumerate(P4):
        s += "\n%% Phần IV Câu %d\n\\begin{ex}\n\t%s\n" % (i + 1, q["stem"])
        if q["parts"]:
            s += "\t\\par " + " \\par ".join("\\textbf{%s)} %s" % ("abc"[k], p) for k, p in enumerate(q["parts"])) + "\n"
        if q.get("fig"):
            s += tex_fig(q["fig"])
        s += "\t\\loigiai{\n%s\n\t}\n\\end{ex}\n" % tex_sol(q["sol"])

    if is_sol:
        s += tex_answer_key()
    s += "\n\\vspace{0.4cm}\n\\begin{center}\n\t\\textbf{--------- HẾT ---------}\n\\end{center}\n\n\\end{document}\n"
    return s


def compile_latex(name, content):
    tex_path = os.path.join(LATEX_DIR, name + ".tex")
    with open(tex_path, "w", encoding="utf-8") as f:
        f.write(content)
    for run in range(2):
        res = subprocess.run(["pdflatex", "-interaction=nonstopmode", name + ".tex"], cwd=LATEX_DIR,
                             capture_output=True, text=True, encoding="utf-8", errors="ignore")
    log = open(os.path.join(LATEX_DIR, name + ".log"), encoding="utf-8", errors="ignore").read()
    errors = [l for l in log.splitlines() if l.startswith("!")]
    pages = re.search(r"Output written on .*?\((\d+) pages?", log)
    print(f"[LaTeX] {name}: {len(errors)} lỗi, {pages.group(1) if pages else '?'} trang")
    for e in errors[:10]:
        print("   ", e)
    return len(errors) == 0


# ============================================================================
# 2. SINH WORD (Pandoc -> OMML + python-docx)
# ============================================================================
def vis_len(s):
    """Độ dài hiển thị ước lượng của một phương án."""
    math = re.findall(r"\$(.*?)\$", s)
    txt = re.sub(r"\$(.*?)\$", "", s)
    mlen = sum(len(re.sub(r"\\[a-zA-Z]+|[{}]", "", m)) for m in math)
    if "frac" in s:
        mlen = int(mlen * 0.8)
    return len(txt) + mlen


def word_opts(opts):
    L = max(vis_len(o) for o in opts) + 4
    items = [r"\textbf{%s.} %s" % (LETTERS[k], o if o.endswith(".") else o + ".") for k, o in enumerate(opts)]
    if L <= 24:
        return "@@TAB@@" + "@@TAB@@".join(items) + "\n\n"
    if L <= 48:
        return "@@TAB@@" + items[0] + "@@TAB@@" + items[1] + "\n\n@@TAB@@" + items[2] + "@@TAB@@" + items[3] + "\n\n"
    return "".join("@@TAB@@" + it + "\n\n" for it in items)


def word_sol(lines):
    return "\n\n".join(lines) + "\n\n"


def word_answer_key():
    s = "\n\n@@KEY_TITLE@@\n\n\\textbf{Phần I.}\n\n\\begin{tabular}{|c|" + "c|" * 12 + "}\\hline\n"
    s += "Câu & " + " & ".join(str(i + 1) for i in range(12)) + " \\\\ \\hline\n"
    s += "Đáp án & " + " & ".join(LETTERS[q["ans"]] for q in P1) + " \\\\ \\hline\n\\end{tabular}\n\n"
    s += "\\textbf{Phần II.}\n\n\\begin{tabular}{|c|c|c|c|c|}\\hline\nCâu & a) & b) & c) & d) \\\\ \\hline\n"
    for i, q in enumerate(P2):
        s += "%d & " % (i + 1) + " & ".join("Đ" if t else "S" for t in q["truth"]) + " \\\\ \\hline\n"
    s += "\\end{tabular}\n\n\\textbf{Phần III.}\n\n\\begin{tabular}{|c|c|c|c|c|}\\hline\nCâu & 1 & 2 & 3 & 4 \\\\ \\hline\n"
    s += "Đáp án & " + " & ".join(q["ans"] for q in P3) + " \\\\ \\hline\n\\end{tabular}\n\n"
    return s


def build_pandoc_tex(is_sol):
    s = "\\documentclass{article}\n\\usepackage{amsmath,amssymb}\n\\begin{document}\n\n@@DOCUMENT_HEADER@@\n\n"
    s += "@@SECTION_1_HEADER@@\n\n"
    for i, q in enumerate(P1):
        s += "\\textbf{Câu %d.} %s\n\n" % (i + 1, q["stem"].replace("\\\\", "\n\n"))
        if q.get("fig"):
            s += "@@CENTER_IMAGE_%s@@\n\n" % q["fig"][0]
        s += word_opts(q["opts"])
        if is_sol:
            s += "\\textbf{Lời giải.} " + word_sol(q["sol"]) + "@@CHON@@ \\textbf{Chọn %s.}\n\n" % LETTERS[q["ans"]]
    s += "@@SECTION_2_HEADER@@\n\n"
    for i, q in enumerate(P2):
        stem = q["stem"].replace("\\begin{center}", "").replace("\\end{center}", "")
        s += "\\textbf{Câu %d.} %s\n\n" % (i + 1, stem)
        if q.get("fig"):
            s += "@@CENTER_IMAGE_%s@@\n\n" % q["fig"][0]
        for k, it in enumerate(q["items"]):
            s += "@@TAB1@@\\textbf{%s)} %s\n\n" % ("abcd"[k], it)
        if is_sol:
            s += "\\textbf{Lời giải.}\n\n"
            for k, (t, so) in enumerate(zip(q["truth"], q["sols"])):
                s += "@@TF_%s@@ \\textbf{%s) %s.} %s\n\n" % ("D" if t else "S", "abcd"[k], "Đúng" if t else "Sai", so)
    s += "@@SECTION_3_HEADER@@\n\n"
    for i, q in enumerate(P3):
        s += "\\textbf{Câu %d.} %s\n\n" % (i + 1, q["stem"])
        if q.get("fig"):
            s += "@@CENTER_IMAGE_%s@@\n\n" % q["fig"][0]
        if is_sol:
            s += "\\textbf{Lời giải.} " + word_sol(q["sol"]) + "@@CHON@@ \\textbf{Đáp số: %s.}\n\n" % q["ans"]
        else:
            s += "@@ANSWER_BOX@@\n\n"
    s += "@@SECTION_4_HEADER@@\n\n"
    for i, q in enumerate(P4):
        s += "\\textbf{Câu %d.} %s\n\n" % (i + 1, q["stem"])
        for k, p in enumerate(q["parts"]):
            s += "@@TAB1@@\\textbf{%s)} %s\n\n" % ("abc"[k], p)
        if q.get("fig"):
            s += "@@CENTER_IMAGE_%s@@\n\n" % q["fig"][0]
        if is_sol:
            s += "\\textbf{Lời giải.} " + word_sol(q["sol"])
    if is_sol:
        s += word_answer_key()
    s += "\n\n@@HET@@\n\n\\end{document}\n"
    return s


# ---------------------- python-docx helpers ----------------------
def set_cell_borders(cell, spec):
    tcPr = cell._tc.get_or_add_tcPr()
    for b in tcPr.findall(qn('w:tcBorders')):
        tcPr.remove(b)
    xml = '<w:tcBorders %s>' % nsdecls('w')
    for side in ("top", "left", "bottom", "right"):
        v = spec.get(side, "none")
        if v == "none":
            xml += '<w:%s w:val="nil"/>' % side
        else:
            xml += '<w:%s w:val="single" w:sz="%s" w:space="0" w:color="%s"/>' % (side, v[0], v[1])
    xml += '</w:tcBorders>'
    tcPr.append(parse_xml(xml))


def add_run(p, text, bold=False, italic=False, size=11, color=None):
    r = p.add_run(text)
    r.bold = bold
    r.italic = italic
    r.font.name = "Times New Roman"
    r.font.size = Pt(size)
    if color:
        r.font.color.rgb = RGBColor(*color)
    return r


def replace_with(p, element):
    p._p.addprevious(element)
    p._p.getparent().remove(p._p)


def make_header_table(doc, is_sol):
    t = doc.add_table(rows=2, cols=3)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    r0 = t.rows[0]
    cl = r0.cells[0]
    cr = r0.cells[1].merge(r0.cells[2])
    cl.width, cr.width = Cm(7.0), Cm(10.5)
    p = cl.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(3)
    add_run(p, "LỚP TOÁN CÔ THÚY\n", True, size=12, color=(31, 73, 125))
    add_run(p, "SĐT: 0935.322.328\nĐịa chỉ: 50/2C Phạm Thị Liên", size=10.5)
    p = cr.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(3)
    if is_sol:
        add_run(p, "HƯỚNG DẪN GIẢI CHI TIẾT GIỮA KỲ I\n", True, size=12, color=(192, 0, 0))
        add_run(p, f"Môn: VẬT LÍ 10 – {SCHOOL}\n", True, size=11)
        add_run(p, f"Năm học {YEAR_W} – Mã đề: {MADE}", italic=True, size=10.5)
    else:
        add_run(p, f"KIỂM TRA GIỮA KỲ I – NĂM HỌC {YEAR_W}\n", True, size=12, color=(192, 0, 0))
        add_run(p, f"Môn: VẬT LÍ, Lớp 10 – {SCHOOL}\n", True, size=10.5)
        add_run(p, "Thời gian làm bài: 45 phút (Không kể thời gian phát đề)", italic=True, size=10)
    r1 = t.rows[1]
    if is_sol:
        c0 = r1.cells[0].merge(r1.cells[1])
        p = c0.paragraphs[0]
        add_run(p, " Giáo viên: ", True, size=10.5)
        add_run(p, "HỒ THỊ THÚY     ", size=10.5)
        add_run(p, "Môn học: ", True, size=10.5)
        add_run(p, "VẬT LÍ 10", size=10.5)
    else:
        add_run(r1.cells[0].paragraphs[0], "Họ và tên thí sinh: ..............................", size=10)
        add_run(r1.cells[1].paragraphs[0], "SBD: ......................", size=10)
    p = r1.cells[2].paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(p, f"Mã đề thi {MADE}", True, size=11)
    for row in t.rows:
        row._tr.get_or_add_trPr().append(parse_xml(r'<w:cantSplit %s/>' % nsdecls('w')))
        for c in row.cells:
            set_cell_borders(c, {k: ("6", "808080") for k in ("top", "left", "bottom", "right")})
            for pp in c.paragraphs:
                pp.paragraph_format.space_before = Pt(3)
                pp.paragraph_format.space_after = Pt(3)
    return t._tbl


def make_section_box(doc, title, note):
    t = doc.add_table(rows=1, cols=1)
    c = t.rows[0].cells[0]
    c.width = Cm(17.5)
    set_cell_borders(c, {"left": ("24", "1F497D")})
    c._tc.get_or_add_tcPr().append(parse_xml(r'<w:shd %s w:val="clear" w:color="auto" w:fill="F2F5F9"/>' % nsdecls('w')))
    p = c.paragraphs[0]
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(3)
    add_run(p, title, True, size=11.5, color=(31, 73, 125))
    add_run(p, "\n" + note, italic=True, size=10, color=(80, 80, 80))
    return t._tbl


def make_answer_box(doc):
    t = doc.add_table(rows=1, cols=5)
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    c0 = t.rows[0].cells[0]
    c0.width = Cm(1.8)
    set_cell_borders(c0, {})
    p = c0.paragraphs[0]
    p.add_run("\t")
    add_run(p, "KQ:", True)
    for k in range(1, 5):
        c = t.rows[0].cells[k]
        c.width = Cm(0.65)
        c.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_borders(c, {s: ("8", "000000") for s in ("top", "left", "bottom", "right")})
    t.rows[0]._tr.get_or_add_trPr().append(parse_xml(r'<w:trHeight %s w:val="380" w:hRule="exact"/>' % nsdecls('w')))
    return t._tbl


def add_page_field(p, size=9.5):
    def fld(instr):
        r = p.add_run()
        r.font.size = Pt(size)
        r.italic = True
        b = OxmlElement('w:fldChar'); b.set(qn('w:fldCharType'), 'begin')
        it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve'); it.text = instr
        e = OxmlElement('w:fldChar'); e.set(qn('w:fldCharType'), 'end')
        r._r.append(b); r._r.append(it); r._r.append(e)
    add_run(p, "Trang ", italic=True, size=size, color=(80, 80, 80))
    fld("PAGE")
    add_run(p, "/", italic=True, size=size, color=(80, 80, 80))
    fld("NUMPAGES")


def strip_marker(p, marker):
    for r in p.runs:
        if marker in r.text:
            r.text = r.text.replace(marker, "").lstrip()


def set_spacing(p, before, after, jc=None):
    pPr = p._p.get_or_add_pPr()
    for sp in pPr.findall(qn('w:spacing')):
        pPr.remove(sp)
    pPr.append(parse_xml(r'<w:spacing %s w:before="%d" w:after="%d" w:line="276" w:lineRule="auto"/>' % (nsdecls('w'), before, after)))
    if jc:
        for j in pPr.findall(qn('w:jc')):
            pPr.remove(j)
        pPr.append(parse_xml(r'<w:jc %s w:val="%s"/>' % (nsdecls('w'), jc)))


def apply_tabs(p):
    xml = p._p.xml
    n = xml.count("@@TAB@@")
    one = "@@TAB1@@" in xml
    xml = xml.replace("@@TAB1@@", '</w:t></w:r><w:r><w:tab/></w:r><w:r><w:t xml:space="preserve">')
    xml = xml.replace("@@TAB@@", '</w:t></w:r><w:r><w:tab/></w:r><w:r><w:t xml:space="preserve">')
    new_p = parse_xml(xml)
    p._p.getparent().replace(p._p, new_p)
    pPr = new_p.get_or_add_pPr()
    for tabs in pPr.findall(qn('w:tabs')):
        pPr.remove(tabs)
    pos = {4: [400, 2800, 5200, 7600], 2: [400, 5200]}.get(n, [400])
    if one:
        pos = [400]
    pPr.append(parse_xml('<w:tabs %s>' % nsdecls('w') + "".join('<w:tab w:val="left" w:pos="%d"/>' % x for x in pos) + '</w:tabs>'))
    for sp in pPr.findall(qn('w:spacing')):
        pPr.remove(sp)
    pPr.append(parse_xml(r'<w:spacing %s w:before="20" w:after="30" w:line="276" w:lineRule="auto"/>' % nsdecls('w')))


def style_tables(doc):
    """Bảng đáp án / bảng số liệu do Pandoc sinh: kẻ viền, căn giữa."""
    for t in doc.tables:
        if t.style is not None and t.style.name == "Table":
            pass
        tblPr = t._tbl.tblPr
        if tblPr.find(qn('w:tblBorders')) is None and len(t.columns) >= 5 and "KQ:" not in t.rows[0].cells[0].text:
            tblPr.append(parse_xml(r'<w:tblBorders %s><w:top w:val="single" w:sz="6" w:color="000000"/><w:left w:val="single" w:sz="6" w:color="000000"/><w:bottom w:val="single" w:sz="6" w:color="000000"/><w:right w:val="single" w:sz="6" w:color="000000"/><w:insideH w:val="single" w:sz="6" w:color="000000"/><w:insideV w:val="single" w:sz="6" w:color="000000"/></w:tblBorders>' % nsdecls('w')))
            for j in tblPr.findall(qn('w:jc')):
                tblPr.remove(j)
            tblPr.append(parse_xml(r'<w:jc %s w:val="center"/>' % nsdecls('w')))
            for row in t.rows:
                for c in row.cells:
                    for pp in c.paragraphs:
                        pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        for r in pp.runs:
                            r.font.name = "Times New Roman"
                            r.font.size = Pt(11)


def build_word(is_sol, out_path):
    kind = "hdg" if is_sol else "de"
    tmp_tex = os.path.join(SCRATCH_DIR, f"temp_{kind}.tex")
    tmp_docx = os.path.join(SCRATCH_DIR, f"temp_{kind}.docx")
    with open(tmp_tex, "w", encoding="utf-8") as f:
        f.write(build_pandoc_tex(is_sol))
    res = subprocess.run([PANDOC, tmp_tex, "-o", tmp_docx, "--from=latex", "--to=docx"],
                         capture_output=True, text=True, encoding="utf-8")
    if res.returncode != 0:
        print("[LỖI PANDOC]", res.stderr)
        return False
    doc = docx.Document(tmp_docx)

    for s in doc.sections:
        s.page_width, s.page_height = Cm(21.0), Cm(29.7)
        s.top_margin, s.bottom_margin = Cm(1.5), Cm(1.5)
        s.left_margin, s.right_margin = Cm(1.8), Cm(1.5)
        s.footer_distance = Cm(0.8)
        fp = s.footer.paragraphs[0]
        fp.text = ""
        fp.paragraph_format.tab_stops.add_tab_stop(Cm(17.7), alignment=2)
        add_run(fp, BRAND_W, italic=True, size=9.5, color=(100, 100, 100))
        add_run(fp, "\t", size=9.5)
        add_page_field(fp)
        add_run(fp, f" – Mã đề {MADE}", italic=True, size=9.5, color=(80, 80, 80))

    st = doc.styles['Normal']
    st.font.name = 'Times New Roman'
    st.font.size = Pt(11)
    st.element.get_or_add_rPr().append(parse_xml(r'<w:rFonts %s w:ascii="Times New Roman" w:hAnsi="Times New Roman" w:cs="Times New Roman" w:eastAsia="Times New Roman"/>' % nsdecls('w')))
    for sname in ("Body Text", "First Paragraph", "Compact"):
        try:
            sty = doc.styles[sname]
            sty.font.name = 'Times New Roman'
            sty.paragraph_format.space_before = Pt(0)
            sty.paragraph_format.space_after = Pt(2)
        except KeyError:
            pass

    style_tables(doc)

    for p in list(doc.paragraphs):
        t = p.text
        if "@@DOCUMENT_HEADER@@" in t:
            replace_with(p, make_header_table(doc, is_sol))
        elif "@@SECTION_" in t:
            k = int(re.search(r"@@SECTION_(\d)_HEADER@@", t).group(1)) - 1
            replace_with(p, make_section_box(doc, *SECTIONS[k]))
        elif "@@ANSWER_BOX@@" in t:
            replace_with(p, make_answer_box(doc))
        elif "@@CENTER_IMAGE_" in t:
            name = re.search(r"@@CENTER_IMAGE_(\w+?)@@", t).group(1)
            width = next(q["fig"][1] for q in P1 + P2 + P3 + P4 if q.get("fig") and q["fig"][0] == name)
            for r in list(p.runs):
                r._r.getparent().remove(r._r)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(os.path.join(FIG_DIR, name + ".png"), width=Cm(width))
            set_spacing(p, 40, 40, "center")
        elif "@@KEY_TITLE@@" in t:
            for r in list(p.runs):
                r._r.getparent().remove(r._r)
            add_run(p, f"BẢNG ĐÁP ÁN – MÃ ĐỀ {MADE}", True, size=12.5, color=(192, 0, 0))
            set_spacing(p, 200, 80, "center")
        elif "@@HET@@" in t:
            for r in list(p.runs):
                r._r.getparent().remove(r._r)
            add_run(p, "--------- HẾT ---------", True)
            set_spacing(p, 160, 0, "center")

    for p in list(doc.paragraphs):
        t = p.text
        if "@@TAB" in p._p.xml:
            apply_tabs(p)
            continue
        if "@@CHON@@" in t:
            strip_marker(p, "@@CHON@@")
            for r in p.runs:
                r.bold = True
                r.font.color.rgb = RGBColor(192, 0, 0)
        elif "@@TF_D@@" in t or "@@TF_S@@" in t:
            ok = "@@TF_D@@" in t
            strip_marker(p, "@@TF_D@@")
            strip_marker(p, "@@TF_S@@")
            for r in p.runs:
                if r.bold and ("Đúng." in r.text or "Sai." in r.text):
                    r.font.color.rgb = RGBColor(0, 128, 0) if ok else RGBColor(192, 0, 0)
        elif t.startswith("Lời giải."):
            for r in p.runs:
                if "Lời giải." in r.text:
                    r.bold = True
                    r.font.color.rgb = RGBColor(31, 73, 125)
                    break
        if t.startswith("Câu "):
            set_spacing(p, 100, 30, "both")
            for r in p.runs[:1]:
                r.font.color.rgb = RGBColor(31, 73, 125)
        for r in p.runs:
            r.font.name = "Times New Roman"
            if r.font.size is None:
                r.font.size = Pt(11)

    # Kiểm tra sót marker + đếm OMML
    body_xml = doc.element.body.xml
    left = re.findall(r"@@\w+@@", body_xml)
    omml = body_xml.count("<m:oMath>") + body_xml.count("<m:oMath ")
    try:
        doc.save(out_path)
    except PermissionError:
        alt = out_path.replace(".docx", "_moi.docx")
        doc.save(alt)
        print(f"[CẢNH BÁO] {os.path.basename(out_path)} đang mở -> lưu tạm {os.path.basename(alt)}")
        out_path = alt
    print(f"[Word] {os.path.basename(out_path)}: {omml} công thức OMML, marker sót: {left if left else 0}")
    return True


def safe_copy(src, dst):
    try:
        shutil.copy2(src, dst)
        print(f"[COPY] {os.path.basename(dst)}")
    except PermissionError:
        print(f"[CẢNH BÁO] {os.path.basename(dst)} đang được mở, bỏ qua sao chép.")


def main():
    print("=" * 60)
    print(f" XUẤT BẢN ĐỀ VẬT LÍ 10 – {SCHOOL} – MÃ ĐỀ {MADE}")
    print("=" * 60)
    ok_de = compile_latex(NAME_DE, build_latex(False))
    ok_hdg = compile_latex(NAME_HDG, build_latex(True))
    build_word(False, os.path.join(SAN_PHAM_DIR, NAME_DE + ".docx"))
    build_word(True, os.path.join(SAN_PHAM_DIR, NAME_HDG + ".docx"))
    for n in (NAME_DE, NAME_HDG):
        src = os.path.join(LATEX_DIR, n + ".pdf")
        if os.path.exists(src):
            safe_copy(src, os.path.join(SAN_PHAM_DIR, n + ".pdf"))
    print("HOÀN TẤT!" if ok_de and ok_hdg else "CÓ LỖI LATEX – kiểm tra log.")


if __name__ == "__main__":
    main()
