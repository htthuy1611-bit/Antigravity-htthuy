import os
import sys
import subprocess
import docx
from docx.shared import Pt, Cm, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn

CURRENT_DIR = os.path.dirname(os.path.abspath(__file__))
BASE_DIR = os.path.abspath(os.path.join(CURRENT_DIR, "..", ".."))
HINH_ANH_DIR = os.path.join(BASE_DIR, "He_Thong", "Hinh_Anh", "Bao_Thang_3")
SAN_PHAM_DIR = os.path.join(BASE_DIR, "San_Pham")
pandoc_exe = os.path.expandvars(r"%LOCALAPPDATA%\Pandoc\pandoc.exe")

os.makedirs(SAN_PHAM_DIR, exist_ok=True)

def create_answer_box_table(doc):
    tbl = doc.add_table(rows=1, cols=5)
    tbl.alignment = WD_TABLE_ALIGNMENT.LEFT
    
    tblPr = tbl._tbl.tblPr
    for look in tblPr.findall(qn('w:tblLook')):
        tblPr.remove(look)
    tblLook = parse_xml(r'<w:tblLook %s w:val="0000" w:firstRow="0" w:lastRow="0" w:firstColumn="0" w:lastColumn="0" w:noHBand="0" w:noVBand="0"/>' % nsdecls('w'))
    for b in tblPr.findall(qn('w:tblBorders')):
        tblPr.remove(b)
    tblBorders = parse_xml(
        r'<w:tblBorders %s>'
        r'<w:top w:val="none"/>'
        r'<w:left w:val="none"/>'
        r'<w:bottom w:val="none"/>'
        r'<w:right w:val="none"/>'
        r'<w:insideH w:val="none"/>'
        r'<w:insideV w:val="none"/>'
        r'</w:tblBorders>' % nsdecls('w')
    )
    tblPr.append(tblBorders)
    tblPr.append(tblLook)
    
    cell_kq = tbl.rows[0].cells[0]
    cell_kq.width = Cm(1.8)
    cell_kq.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    tcPr0 = cell_kq._tc.get_or_add_tcPr()
    for b in tcPr0.findall(qn('w:tcBorders')):
        tcPr0.remove(b)
    tcPr0.append(parse_xml(
        r'<w:tcBorders %s>'
        r'<w:top w:val="none"/>'
        r'<w:left w:val="none"/>'
        r'<w:bottom w:val="none"/>'
        r'<w:right w:val="none"/>'
        r'</w:tcBorders>' % nsdecls('w')
    ))
    
    p_kq = cell_kq.paragraphs[0]
    p_kq.paragraph_format.space_before = Pt(0)
    p_kq.paragraph_format.space_after = Pt(0)
    p_kq.add_run("\t")
    r = p_kq.add_run("KQ:")
    r.bold = True
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)
    
    for col_idx in range(1, 5):
        cell = tbl.rows[0].cells[col_idx]
        cell.width = Cm(0.65)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        tcPr = cell._tc.get_or_add_tcPr()
        for b in tcPr.findall(qn('w:tcBorders')):
            tcPr.remove(b)
        borders = parse_xml(
            r'<w:tcBorders %s>'
            r'<w:top w:val="single" w:sz="8" w:space="0" w:color="000000"/>'
            r'<w:left w:val="single" w:sz="8" w:space="0" w:color="000000"/>'
            r'<w:bottom w:val="single" w:sz="8" w:space="0" w:color="000000"/>'
            r'<w:right w:val="single" w:sz="8" w:space="0" w:color="000000"/>'
            r'</w:tcBorders>' % nsdecls('w')
        )
        tcPr.append(borders)
    
    trPr = tbl.rows[0]._tr.get_or_add_trPr()
    trHeight = parse_xml(r'<w:trHeight %s w:val="380" w:hRule="exact"/>' % nsdecls('w'))
    trPr.append(trHeight)
    return tbl

def format_doc_paragraph(p, inside_table_cell=False):
    p_xml = p._p.xml
    text = p.text.strip()
    
    if "@@TAB@@" in p_xml:
        tab_count = p_xml.count("@@TAB@@")
        new_xml = p_xml.replace("@@TAB@@", '</w:t><w:tab/><w:t>')
        new_p = parse_xml(new_xml)
        p._p.getparent().replace(p._p, new_p)
        
        pPr = new_p.get_or_add_pPr()
        for tabs in pPr.findall(qn('w:tabs')):
            pPr.remove(tabs)
            
        if tab_count >= 4:
            tabs_xml = (
                r'<w:tabs %s>'
                r'<w:tab w:val="left" w:pos="425"/>'
                r'<w:tab w:val="left" w:pos="2693"/>'
                r'<w:tab w:val="left" w:pos="5103"/>'
                r'<w:tab w:val="left" w:pos="7512"/>'
                r'</w:tabs>' % nsdecls('w')
            )
            pPr.append(parse_xml(tabs_xml))
            for jc in pPr.findall(qn('w:jc')):
                pPr.remove(jc)
            pPr.append(parse_xml(r'<w:jc %s w:val="left"/>' % nsdecls('w')))
            
        elif tab_count == 2:
            if inside_table_cell:
                tabs_xml = (
                    r'<w:tabs %s>'
                    r'<w:tab w:val="left" w:pos="283"/>'
                    r'<w:tab w:val="left" w:pos="3402"/>'
                    r'</w:tabs>' % nsdecls('w')
                )
            else:
                tabs_xml = (
                    r'<w:tabs %s>'
                    r'<w:tab w:val="left" w:pos="425"/>'
                    r'<w:tab w:val="left" w:pos="5103"/>'
                    r'</w:tabs>' % nsdecls('w')
                )
            pPr.append(parse_xml(tabs_xml))
            for jc in pPr.findall(qn('w:jc')):
                pPr.remove(jc)
            pPr.append(parse_xml(r'<w:jc %s w:val="left"/>' % nsdecls('w')))
            
        elif tab_count == 1:
            tabs_xml = r'<w:tabs %s><w:tab w:val="left" w:pos="425"/></w:tabs>' % nsdecls('w')
            pPr.append(parse_xml(tabs_xml))
            for jc in pPr.findall(qn('w:jc')):
                pPr.remove(jc)
            pPr.append(parse_xml(r'<w:jc %s w:val="both"/>' % nsdecls('w')))
            for ind in pPr.findall(qn('w:ind')):
                pPr.remove(ind)
            ind_xml = r'<w:ind %s w:left="425" w:hanging="425"/>' % nsdecls('w')
            pPr.append(parse_xml(ind_xml))
            
        for sp in pPr.findall(qn('w:spacing')):
            pPr.remove(sp)
        sp_xml = r'<w:spacing %s w:before="0" w:after="50" w:line="276" w:lineRule="auto"/>' % nsdecls('w')
        pPr.append(parse_xml(sp_xml))
        return

    pPr = p._p.get_or_add_pPr()
    if text.startswith("Câu "):
        for jc in pPr.findall(qn('w:jc')):
            pPr.remove(jc)
        pPr.append(parse_xml(r'<w:jc %s w:val="both"/>' % nsdecls('w')))
        for sp in pPr.findall(qn('w:spacing')):
            pPr.remove(sp)
        sp_xml = r'<w:spacing %s w:before="60" w:after="40" w:line="276" w:lineRule="auto"/>' % nsdecls('w')
        pPr.append(parse_xml(sp_xml))
    elif "HẾT" in text:
        for jc in pPr.findall(qn('w:jc')):
            pPr.remove(jc)
        pPr.append(parse_xml(r'<w:jc %s w:val="center"/>' % nsdecls('w')))
        for sp in pPr.findall(qn('w:spacing')):
            pPr.remove(sp)
        sp_xml = r'<w:spacing %s w:before="140" w:after="140" w:line="276" w:lineRule="auto"/>' % nsdecls('w')
        pPr.append(parse_xml(sp_xml))

    for r in p.runs:
        r.font.name = 'Times New Roman'
        r.font.size = Pt(12)
        r_rPr = r._r.get_or_add_rPr()
        f = parse_xml(r'<w:rFonts %s w:ascii="Times New Roman" w:hAnsi="Times New Roman" w:cs="Times New Roman"/>' % nsdecls('w'))
        r_rPr.append(f)

def insert_header_and_code(doc, target_p, ma_de, is_solution=False):
    table = doc.add_table(rows=2, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    row0 = table.rows[0]
    cell_left = row0.cells[0]
    cell_right = row0.cells[1]
    cell_right.merge(row0.cells[2])
    
    cell_left.width = Cm(8.5)
    cell_right.width = Cm(9.0)
    
    p_l = cell_left.paragraphs[0]
    p_l.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_l.paragraph_format.line_spacing = 1.15
    p_l.paragraph_format.space_before = Pt(3)
    p_l.paragraph_format.space_after = Pt(3)
    r1 = p_l.add_run("LỚP TOÁN CÔ THÚY\n")
    r1.bold = True
    r1.font.name = "Times New Roman"
    r1.font.size = Pt(12)
    r1.font.color.rgb = RGBColor(31, 73, 125)
    
    r2 = p_l.add_run("SĐT: 0935.322.328\nĐịa chỉ: 50/2C Phạm Thị Liên")
    r2.font.name = "Times New Roman"
    r2.font.size = Pt(10.5)
    
    p_r = cell_right.paragraphs[0]
    p_r.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_r.paragraph_format.line_spacing = 1.15
    p_r.paragraph_format.space_before = Pt(3)
    p_r.paragraph_format.space_after = Pt(3)
    
    if is_solution:
        r_top = p_r.add_run("HƯỚNG DẪN GIẢI CHI TIẾT\n")
        r_top.bold = True
        r_top.font.name = "Times New Roman"
        r_top.font.size = Pt(12)
        r_top.font.color.rgb = RGBColor(192, 0, 0)
        
        r3 = p_r.add_run("ĐỀ KHẢO SÁT CHẤT LƯỢNG TOÁN 12\n")
        r3.bold = True
        r3.font.name = "Times New Roman"
        r3.font.size = Pt(11)
        r3.font.color.rgb = RGBColor(0, 0, 0)
        
        r_src = p_r.add_run("THPT Số 3 Bảo Thắng (Năm học 2026 - 2027)")
        r_src.italic = True
        r_src.font.name = "Times New Roman"
        r_src.font.size = Pt(10.5)
    else:
        r3 = p_r.add_run("ĐỀ KHẢO SÁT CHẤT LƯỢNG TOÁN 12\n")
        r3.bold = True
        r3.font.name = "Times New Roman"
        r3.font.size = Pt(12)
        r3.font.color.rgb = RGBColor(192, 0, 0)
        
        r_src = p_r.add_run("THPT Số 3 Bảo Thắng (2026 - 2027)\n")
        r_src.bold = True
        r_src.font.name = "Times New Roman"
        r_src.font.size = Pt(10.5)
        
        r4 = p_r.add_run("Thời gian làm bài: 25 phút")
        r4.italic = True
        r4.font.name = "Times New Roman"
        r4.font.size = Pt(10.5)
    
    row1 = table.rows[1]
    if is_solution:
        c_hoten = row1.cells[0]
        c_hoten.merge(row1.cells[1])
        c_hoten.width = Cm(14.0)
        c_made = row1.cells[2]
        c_made.width = Cm(3.5)
        
        p_hoten = c_hoten.paragraphs[0]
        p_hoten.paragraph_format.space_before = Pt(3)
        p_hoten.paragraph_format.space_after = Pt(3)
        r_ht = p_hoten.add_run("TÀI LIỆU DÀNH CHO GIÁO VIÊN / HỌC SINH THAM KHẢO")
        r_ht.bold = True
        r_ht.font.size = Pt(10)
        r_ht.font.name = "Times New Roman"
        
        p_md = c_made.paragraphs[0]
        p_md.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_md.paragraph_format.space_before = Pt(3)
        p_md.paragraph_format.space_after = Pt(3)
        r_md = p_md.add_run(f"Mã đề: {ma_de}")
        r_md.bold = True
        r_md.font.name = "Times New Roman"
        r_md.font.size = Pt(11)
    else:
        c_hoten = row1.cells[0]
        c_sbd = row1.cells[1]
        c_made = row1.cells[2]
        c_hoten.width = Cm(8.5)
        c_sbd.width = Cm(5.5)
        c_made.width = Cm(3.5)
        
        p_hoten = c_hoten.paragraphs[0]
        p_hoten.paragraph_format.space_before = Pt(3)
        p_hoten.paragraph_format.space_after = Pt(3)
        r_ht = p_hoten.add_run("Họ và tên: ................................................................")
        r_ht.font.size = Pt(11)
        r_ht.font.name = "Times New Roman"
        
        p_sbd = c_sbd.paragraphs[0]
        p_sbd.paragraph_format.space_before = Pt(3)
        p_sbd.paragraph_format.space_after = Pt(3)
        r_sb = p_sbd.add_run("Số báo danh: ................")
        r_sb.font.name = "Times New Roman"
        r_sb.font.size = Pt(11)
        
        p_md = c_made.paragraphs[0]
        p_md.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_md.paragraph_format.space_before = Pt(3)
        p_md.paragraph_format.space_after = Pt(3)
        r_md = p_md.add_run(f"Mã đề {ma_de}")
        r_md.bold = True
        r_md.font.name = "Times New Roman"
        r_md.font.size = Pt(11)
    
    for row in table.rows:
        for cell in row.cells:
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            
    tblPr = table._tbl.tblPr
    for b in tblPr.findall(qn('w:tblBorders')):
        tblPr.remove(b)
    tblBorders = parse_xml(
        r'<w:tblBorders %s>'
        r'<w:top w:val="single" w:sz="8" w:space="0" w:color="000000"/>'
        r'<w:left w:val="single" w:sz="8" w:space="0" w:color="000000"/>'
        r'<w:bottom w:val="single" w:sz="8" w:space="0" w:color="000000"/>'
        r'<w:right w:val="single" w:sz="8" w:space="0" w:color="000000"/>'
        r'<w:insideH w:val="single" w:sz="6" w:space="0" w:color="000000"/>'
        r'<w:insideV w:val="single" w:sz="6" w:space="0" w:color="000000"/>'
        r'</w:tblBorders>' % nsdecls('w')
    )
    tblPr.append(tblBorders)
    
    target_p._p.addprevious(table._tbl)
    target_p._p.getparent().remove(target_p._p)

def insert_section_header(doc, target_p, sec_title, sec_desc):
    p_sec = doc.add_paragraph()
    p_sec.paragraph_format.space_before = Pt(8)
    p_sec.paragraph_format.space_after = Pt(2)
    r_title = p_sec.add_run(sec_title)
    r_title.bold = True
    r_title.font.name = "Times New Roman"
    r_title.font.size = Pt(12)
    
    if sec_desc:
        r_desc = p_sec.add_run(" " + sec_desc)
        r_desc.font.name = "Times New Roman"
        r_desc.font.size = Pt(11)
        r_desc.italic = True
        
    p_sec.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    target_p._p.addprevious(p_sec._p)
    target_p._p.getparent().remove(target_p._p)
