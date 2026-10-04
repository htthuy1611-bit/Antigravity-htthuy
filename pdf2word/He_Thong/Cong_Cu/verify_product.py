# -*- coding: utf-8 -*-
"""Kiểm tra sản phẩm: header/footer Word + ảnh xem trước PDF. python verify.py VL10_BTX 207"""
import os, sys, re
import docx, pymupdf as fitz
sys.stdout.reconfigure(encoding="utf-8")
BASE = r"c:\AnTiGraViTy-htthuy\pdf2word\San_Pham"
OUT = r"C:\Users\user\.gemini\antigravity\brain\bd035255-37c1-4545-9eea-d9028737b8e8\scratch"
prefix, made = sys.argv[1], sys.argv[2]
BAD = ["BÙI THỊ HỒNG", "Phan Đình Quân", "Thầy", "Bùi Thị Xuân --- MÃ", "KHỐI 11"]
for kind in ("De", "HDG"):
    name = f"{prefix}_{kind}_{made}"
    d = docx.Document(os.path.join(BASE, name + ".docx"))
    t0 = d.tables[0]
    hdr = " | ".join(c.text.replace("\n", " / ") for r in t0.rows for c in r.cells)
    ftr = d.sections[0].footer.paragraphs[0].text
    body = "\n".join(p.text for p in d.paragraphs) + hdr
    bad = [b for b in BAD if b in body + ftr]
    print(f"== {name}.docx\n  HEADER: {hdr}\n  FOOTER: {ftr}\n  Cấm: {bad or 'OK'}")
    pdf = fitz.open(os.path.join(BASE, name + ".pdf"))
    for i in (0, len(pdf) - 1):
        pdf[i].get_pixmap(dpi=70).save(os.path.join(OUT, f"{name}_p{i+1}.png"))
    txt = "".join(p.get_text() for p in pdf)
    print(f"  PDF: {len(pdf)} trang; 'HỒ THỊ THÚY' x{txt.count('HỒ THỊ THÚY')}; 'Lớp Toán Cô Thúy' x{txt.count('Lớp Toán Cô Thúy')}; cấm: {[b for b in BAD if b in txt] or 'OK'}")
