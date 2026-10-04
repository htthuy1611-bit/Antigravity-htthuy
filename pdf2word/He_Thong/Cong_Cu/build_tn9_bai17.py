import os
import sys
import subprocess
import shutil
import docx
from docx.shared import Pt, Cm, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn

if sys.stdout.encoding.lower() != 'utf-8':
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except Exception:
        pass

CURRENT_DIR = os.path.dirname(os.path.abspath(__file__))
BASE_DIR = os.path.abspath(os.path.join(CURRENT_DIR, "..", ".."))
HINH_ANH_DIR = os.path.join(BASE_DIR, "He_Thong", "Hinh_Anh", "TN9_Chuong5_Bai17")
LATEX_DIR = os.path.join(BASE_DIR, "He_Thong", "LaTeX", "TN9_Chuong5_Bai17")
SAN_PHAM_DIR = os.path.join(BASE_DIR, "San_Pham")
pandoc_exe = os.path.expandvars(r"%LOCALAPPDATA%\Pandoc\pandoc.exe")

os.makedirs(SAN_PHAM_DIR, exist_ok=True)

OUTPUT_DE_DOCX = os.path.join(SAN_PHAM_DIR, "TN9_Chuong5_Bai17_De.docx")
OUTPUT_HDG_DOCX = os.path.join(SAN_PHAM_DIR, "TN9_Chuong5_Bai17_HDG.docx")
OUTPUT_DE_PDF = os.path.join(SAN_PHAM_DIR, "TN9_Chuong5_Bai17_De.pdf")
OUTPUT_HDG_PDF = os.path.join(SAN_PHAM_DIR, "TN9_Chuong5_Bai17_HDG.pdf")

def insert_header_and_code(doc, target_p, is_solution=False):
    table = doc.add_table(rows=2, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    row0 = table.rows[0]
    cell_left = row0.cells[0]
    cell_right = row0.cells[1]
    cell_right.merge(row0.cells[2])
    
    cell_left.width = Cm(7.5)
    cell_right.width = Cm(10.0)
    
    p_l = cell_left.paragraphs[0]
    p_l.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_l.paragraph_format.line_spacing = 1.15
    p_l.paragraph_format.space_before = Pt(3)
    p_l.paragraph_format.space_after = Pt(3)
    r1 = p_l.add_run("LỚP TOÁN CÔ THÚY\n")
    r1.bold = True
    r1.font.name = "Times New Roman"
    r1.font.size = Pt(12)
    r1.font.color.rgb = RGBColor(31, 73, 125)
    
    r2 = p_l.add_run("SĐT: 0935.322.328\nĐịa chỉ: 50/2C Phạm Thị Liên")
    r2.font.name = "Times New Roman"
    r2.font.size = Pt(10)
    
    p_r = cell_right.paragraphs[0]
    p_r.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_r.paragraph_format.line_spacing = 1.15
    p_r.paragraph_format.space_before = Pt(3)
    p_r.paragraph_format.space_after = Pt(3)
    
    if is_solution:
        r_top = p_r.add_run("HƯỚNG DẪN GIẢI CHI TIẾT TRẮC NGHIỆM\n")
        r_top.bold = True
        r_top.font.name = "Times New Roman"
        r_top.font.size = Pt(12)
        r_top.font.color.rgb = RGBColor(192, 0, 0)
        
        r3 = p_r.add_run("Môn: TOÁN 9 – KẾT NỐI TRI THỨC VỚI CUỘC SỐNG\n")
        r3.bold = True
        r3.font.name = "Times New Roman"
        r3.font.size = Pt(11)
        r3.font.color.rgb = RGBColor(0, 0, 0)
        
        r_src = p_r.add_run("Chương V: Đường tròn – Bài 17: Vị trí tương đối của hai đường tròn")
        r_src.italic = True
        r_src.font.name = "Times New Roman"
        r_src.font.size = Pt(10)
    else:
        r3 = p_r.add_run("PHIẾU BÀI TẬP TRẮC NGHIỆM TOÁN 9\n")
        r3.bold = True
        r3.font.name = "Times New Roman"
        r3.font.size = Pt(12)
        r3.font.color.rgb = RGBColor(192, 0, 0)
        
        r_src = p_r.add_run("Chương V: Đường tròn – Bài 17: Vị trí tương đối của hai đường tròn\n")
        r_src.bold = True
        r_src.font.name = "Times New Roman"
        r_src.font.size = Pt(10.5)
        
        r4 = p_r.add_run("Thời gian làm bài: 45 phút (Bộ sách Kết nối tri thức)")
        r4.italic = True
        r4.font.name = "Times New Roman"
        r4.font.size = Pt(10)
    
    row1 = table.rows[1]
    if is_solution:
        c_hoten = row1.cells[0]
        c_hoten.merge(row1.cells[1])
        c_hoten.width = Cm(13.5)
        c_made = row1.cells[2]
        c_made.width = Cm(4.0)
        
        p_info = c_hoten.paragraphs[0]
        p_info.paragraph_format.space_before = Pt(3)
        p_info.paragraph_format.space_after = Pt(3)
        r_ht = p_info.add_run(" Giáo viên: ")
        r_ht.bold = True
        r_ht.font.name = "Times New Roman"
        r_ht.font.size = Pt(10.5)
        r_ht_v = p_info.add_run("HỒ THỊ THÚY          ")
        r_ht_v.font.name = "Times New Roman"
        r_ht_v.font.size = Pt(10.5)
        
        r_mh = p_info.add_run("Môn học: ")
        r_mh.bold = True
        r_mh.font.name = "Times New Roman"
        r_mh.font.size = Pt(10.5)
        r_mh_v = p_info.add_run("TOÁN 9 (KNTT)")
        r_mh_v.font.name = "Times New Roman"
        r_mh_v.font.size = Pt(10.5)
        
        p_code = c_made.paragraphs[0]
        p_code.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_code.paragraph_format.space_before = Pt(3)
        p_code.paragraph_format.space_after = Pt(3)
        rc = p_code.add_run("Mã đề: 101")
        rc.bold = True
        rc.font.name = "Times New Roman"
        rc.font.size = Pt(11)
    else:
        c_hoten = row1.cells[0]
        c_sbd = row1.cells[1]
        c_made = row1.cells[2]
        c_hoten.width = Cm(8.0)
        c_sbd.width = Cm(5.5)
        c_made.width = Cm(4.0)
        
        p_ht = c_hoten.paragraphs[0]
        p_ht.paragraph_format.space_before = Pt(3)
        p_ht.paragraph_format.space_after = Pt(3)
        r_ht = p_ht.add_run("Họ và tên học sinh: ............................")
        r_ht.font.name = "Times New Roman"
        r_ht.font.size = Pt(10)
        
        p_sbd = c_sbd.paragraphs[0]
        p_sbd.paragraph_format.space_before = Pt(3)
        p_sbd.paragraph_format.space_after = Pt(3)
        r_sbd = p_sbd.add_run("Lớp: ......................")
        r_sbd.font.name = "Times New Roman"
        r_sbd.font.size = Pt(10)
        
        p_code = c_made.paragraphs[0]
        p_code.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_code.paragraph_format.space_before = Pt(3)
        p_code.paragraph_format.space_after = Pt(3)
        rc = p_code.add_run("Mã đề: 101")
        rc.bold = True
        rc.font.name = "Times New Roman"
        rc.font.size = Pt(11)
        
    for row in table.rows:
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(r'<w:cantSplit %s/>' % nsdecls('w')))
        for cell in row.cells:
            tcPr = cell._tc.get_or_add_tcPr()
            for b in tcPr.findall(qn('w:tcBorders')):
                tcPr.remove(b)
            borders = parse_xml(
                r'<w:tcBorders %s>'
                r'<w:top w:val="single" w:sz="6" w:space="0" w:color="808080"/>'
                r'<w:left w:val="single" w:sz="6" w:space="0" w:color="808080"/>'
                r'<w:bottom w:val="single" w:sz="6" w:space="0" w:color="808080"/>'
                r'<w:right w:val="single" w:sz="6" w:space="0" w:color="808080"/>'
                r'</w:tcBorders>' % nsdecls('w')
            )
            tcPr.append(borders)
            
    target_p._p.addprevious(table._tbl)
    target_p._p.getparent().remove(target_p._p)

def insert_section_header(doc, target_p, title, note=""):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    cell = table.rows[0].cells[0]
    cell.width = Cm(17.5)
    
    tcPr = cell._tc.get_or_add_tcPr()
    for b in tcPr.findall(qn('w:tcBorders')):
        tcPr.remove(b)
    borders = parse_xml(
        r'<w:tcBorders %s>'
        r'<w:top w:val="none"/>'
        r'<w:left w:val="single" w:sz="24" w:space="0" w:color="1F497D"/>'
        r'<w:bottom w:val="none"/>'
        r'<w:right w:val="none"/>'
        r'</w:tcBorders>' % nsdecls('w')
    )
    tcPr.append(borders)
    shd = parse_xml(r'<w:shd %s w:fill="F2F5F9"/>' % nsdecls('w'))
    tcPr.append(shd)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    
    r_title = p.add_run(title)
    r_title.bold = True
    r_title.font.name = "Times New Roman"
    r_title.font.size = Pt(11.5)
    r_title.font.color.rgb = RGBColor(31, 73, 125)
    
    if note:
        r_note = p.add_run(f"\n{note}")
        r_note.italic = True
        r_note.font.name = "Times New Roman"
        r_note.font.size = Pt(10)
        r_note.font.color.rgb = RGBColor(80, 80, 80)
        
    target_p._p.addprevious(table._tbl)
    target_p._p.getparent().remove(target_p._p)

def insert_answer_key_table(doc, target_p):
    tbl = doc.add_table(rows=2, cols=15)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    keys = [
        ("1", "D"), ("2", "C"), ("3", "C"), ("4", "D"), ("5", "D"),
        ("6", "A"), ("7", "D"), ("8", "A"), ("9", "A"), ("10", "B"),
        ("11", "A"), ("12", "A"), ("13", "B"), ("14", "B"), ("15", "D")
    ]
    
    for i, (q, a) in enumerate(keys):
        cell_q = tbl.rows[0].cells[i]
        cell_q.width = Cm(1.1)
        cell_q.paragraphs[0].text = q
        cell_q.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        cell_q.paragraphs[0].paragraph_format.space_before = Pt(2)
        cell_q.paragraphs[0].paragraph_format.space_after = Pt(2)
        cell_q.paragraphs[0].runs[0].bold = True
        cell_q.paragraphs[0].runs[0].font.name = "Times New Roman"
        cell_q.paragraphs[0].runs[0].font.size = Pt(10.5)
        
        tcPr = cell_q._tc.get_or_add_tcPr()
        shd = parse_xml(r'<w:shd %s w:fill="E9EEF4"/>' % nsdecls('w'))
        tcPr.append(shd)
        
        cell_a = tbl.rows[1].cells[i]
        cell_a.width = Cm(1.1)
        cell_a.paragraphs[0].text = a
        cell_a.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        cell_a.paragraphs[0].paragraph_format.space_before = Pt(2)
        cell_a.paragraphs[0].paragraph_format.space_after = Pt(2)
        cell_a.paragraphs[0].runs[0].bold = True
        cell_a.paragraphs[0].runs[0].font.name = "Times New Roman"
        cell_a.paragraphs[0].runs[0].font.size = Pt(11)
        cell_a.paragraphs[0].runs[0].font.color.rgb = RGBColor(192, 0, 0)
        
    for row in tbl.rows:
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(r'<w:cantSplit %s/>' % nsdecls('w')))
        for cell in row.cells:
            tcPr = cell._tc.get_or_add_tcPr()
            borders = parse_xml(
                r'<w:tcBorders %s>'
                r'<w:top w:val="single" w:sz="6" w:space="0" w:color="808080"/>'
                r'<w:left w:val="single" w:sz="6" w:space="0" w:color="808080"/>'
                r'<w:bottom w:val="single" w:sz="6" w:space="0" w:color="808080"/>'
                r'<w:right w:val="single" w:sz="6" w:space="0" w:color="808080"/>'
                r'</w:tcBorders>' % nsdecls('w')
            )
            tcPr.append(borders)
            
    target_p._p.addprevious(tbl._tbl)
    target_p._p.getparent().remove(target_p._p)

def format_doc_paragraph(p):
    p_xml = p._p.xml
    text = p.text.strip()
    
    if "@@TAB@@" in p_xml:
        tab_count = p_xml.count("@@TAB@@")
        new_xml = p_xml.replace("@@TAB@@", '</w:t><w:tab/><w:t>')
        new_p = parse_xml(new_xml)
        p._p.getparent().replace(p._p, new_p)
        
        pPr = new_p.get_or_add_pPr()
        for tabs in pPr.findall(qn('w:tabs')):
            pPr.remove(tabs)
            
        tabs_xml = r'<w:tabs %s>' % nsdecls('w')
        if tab_count == 4:
            tabs_xml += r'<w:tab w:val="left" w:pos="300"/>'
            tabs_xml += r'<w:tab w:val="left" w:pos="2700"/>'
            tabs_xml += r'<w:tab w:val="left" w:pos="5100"/>'
            tabs_xml += r'<w:tab w:val="left" w:pos="7500"/>'
        elif tab_count == 2:
            tabs_xml += r'<w:tab w:val="left" w:pos="300"/>'
            tabs_xml += r'<w:tab w:val="left" w:pos="5000"/>'
        else:
            tabs_xml += r'<w:tab w:val="left" w:pos="300"/>'
        tabs_xml += r'</w:tabs>'
        pPr.append(parse_xml(tabs_xml))
        
        for sp in pPr.findall(qn('w:spacing')):
            pPr.remove(sp)
        sp_xml = r'<w:spacing %s w:before="20" w:after="30" w:line="276" w:lineRule="auto"/>' % nsdecls('w')
        pPr.append(parse_xml(sp_xml))
        return

    pPr = p._p.get_or_add_pPr()
    if text.startswith("Câu "):
        for jc in pPr.findall(qn('w:jc')):
            pPr.remove(jc)
        pPr.append(parse_xml(r'<w:jc %s w:val="both"/>' % nsdecls('w')))
        for sp in pPr.findall(qn('w:spacing')):
            pPr.remove(sp)
        sp_xml = r'<w:spacing %s w:before="60" w:after="25" w:line="276" w:lineRule="auto"/>' % nsdecls('w')
        pPr.append(parse_xml(sp_xml))
    elif text.startswith("Lời giải:"):
        for jc in pPr.findall(qn('w:jc')):
            pPr.remove(jc)
        pPr.append(parse_xml(r'<w:jc %s w:val="both"/>' % nsdecls('w')))
        for sp in pPr.findall(qn('w:spacing')):
            pPr.remove(sp)
        sp_xml = r'<w:spacing %s w:before="40" w:after="30" w:line="276" w:lineRule="auto"/>' % nsdecls('w')
        pPr.append(parse_xml(sp_xml))
    elif "HẾT" in text:
        for jc in pPr.findall(qn('w:jc')):
            pPr.remove(jc)
        pPr.append(parse_xml(r'<w:jc %s w:val="center"/>' % nsdecls('w')))
        for sp in pPr.findall(qn('w:spacing')):
            pPr.remove(sp)
        sp_xml = r'<w:spacing %s w:before="120" w:after="120" w:line="276" w:lineRule="auto"/>' % nsdecls('w')
        pPr.append(parse_xml(sp_xml))

    for r in p.runs:
        r.font.name = 'Times New Roman'
        r.font.size = Pt(11)
        r_rPr = r._r.get_or_add_rPr()
        f = parse_xml(r'<w:rFonts %s w:ascii="Times New Roman" w:hAnsi="Times New Roman" w:cs="Times New Roman"/>' % nsdecls('w'))
        r_rPr.append(f)

# LaTeX content for ĐỀ BÀI (Exam paper)
tex_de = r"""\documentclass[11pt,a4paper]{article}
\usepackage[utf8]{vietnam}
\usepackage{amsmath,amssymb}
\usepackage{graphicx}

\begin{document}

@@DOCUMENT_HEADER@@

@@SECTION_1_HEADER@@

\textbf{Câu 1.} Nếu hai đường tròn không cắt nhau thì số điểm chung của hai đường tròn là

@@TAB@@\textbf{A.} $1$.@@TAB@@\textbf{B.} $2$.@@TAB@@\textbf{C.} $3$.@@TAB@@\textbf{D.} $0$.

\textbf{Câu 2.} Cho hai đường tròn $(O_1)$ và $(O_2)$ tiếp xúc ngoài tại $A$ và một đường thẳng $\mathrm{d}$ tiếp xúc với $(O_1)$ và $(O_2)$ lần lượt tại $B, C$. Tam giác $ABC$ là

@@TAB@@\textbf{A.} Tam giác cân.@@TAB@@\textbf{B.} Tam giác đều.@@TAB@@\textbf{C.} Tam giác vuông.@@TAB@@\textbf{D.} Tam giác vuông cân.

\textbf{Câu 3.} Cho đoạn thẳng $OO'$ và điểm $A$ nằm trên đoạn $OO'$ sao cho $OA = 2O'A$. Vị trí tương đối của đường tròn tâm $O$ bán kính $OA$ và đường tròn tâm $O'$ bán kính $O'A$ là

@@TAB@@\textbf{A.} Nằm ngoài nhau.@@TAB@@\textbf{B.} Cắt nhau.@@TAB@@\textbf{C.} Tiếp xúc ngoài.@@TAB@@\textbf{D.} Tiếp xúc trong.

\textbf{Câu 4.} Cho hai đường tròn $(O), (O')$ cắt nhau tại $A, B$ trong đó $O' \in (O)$. Kẻ đường kính $O'C$ của đường tròn $(O)$. Chọn khẳng định \textbf{sai}.

@@TAB@@\textbf{A.} $AC = CB$.@@TAB@@\textbf{B.} $\widehat{CBO'} = 90^\circ$.

@@TAB@@\textbf{C.} $CA, CB$ là hai tiếp tuyến của $(O')$.@@TAB@@\textbf{D.} $CA, CB$ là hai cát tuyến của $(O')$.

\textbf{Câu 5.} Cho hai đường tròn tiếp xúc ngoài $(O; R)$ và $(O'; r)$ với $R > r$ và khoảng cách giữa hai tâm $OO' = \mathrm{d}$. Khi đó

@@TAB@@\textbf{A.} $\mathrm{d} = R - r$.@@TAB@@\textbf{B.} $\mathrm{d} > R + r$.@@TAB@@\textbf{C.} $R - r < \mathrm{d} < R + r$.@@TAB@@\textbf{D.} $\mathrm{d} = R + r$.

\textbf{Câu 6.} Cho đường tròn $(O_1; 3\text{ cm})$ tiếp xúc ngoài với đường tròn $(O_2; 1\text{ cm})$. Vẽ bán kính $O_1B$ và $O_2C$ song song với nhau cùng thuộc nửa mặt phẳng bờ $O_1O_2$. Gọi $D$ là giao điểm của $BC$ và $O_1O_2$. Tính số đo góc $\widehat{BAC}$ (với $A$ là tiếp điểm của $(O_1)$ và $(O_2)$).

@@TAB@@\textbf{A.} $90^\circ$.@@TAB@@\textbf{B.} $60^\circ$.@@TAB@@\textbf{C.} $100^\circ$.@@TAB@@\textbf{D.} $80^\circ$.

\textbf{Câu 7.} Cho hai đường tròn $(O; 10\text{ cm})$ và $(O'; 5\text{ cm})$ cắt nhau tại $A, B$. Tính độ dài đoạn thẳng $OO'$ biết $AB = 8\text{ cm}$ và $O, O'$ nằm cùng phía đối với $AB$ (làm tròn kết quả đến chữ số thập phân thứ nhất).

@@TAB@@\textbf{A.} $OO' \approx 6{,}5\text{ cm}$.@@TAB@@\textbf{B.} $OO' \approx 6{,}1\text{ cm}$.@@TAB@@\textbf{C.} $OO' \approx 6{,}3\text{ cm}$.@@TAB@@\textbf{D.} $OO' \approx 6{,}2\text{ cm}$.

\textbf{Câu 8.} Cho hai đường tròn $(O; 20\text{ cm})$ và $(O'; 15\text{ cm})$ cắt nhau tại $A, B$. Tính độ dài đoạn thẳng $OO'$ biết $AB = 24\text{ cm}$ và $O, O'$ nằm cùng phía đối với $AB$.

@@TAB@@\textbf{A.} $OO' = 7\text{ cm}$.@@TAB@@\textbf{B.} $OO' = 8\text{ cm}$.@@TAB@@\textbf{C.} $OO' = 9\text{ cm}$.@@TAB@@\textbf{D.} $OO' = 25\text{ cm}$.

\textbf{Câu 9.} Cho hai đường tròn $(O), (O')$ tiếp xúc ngoài tại $A$. Kẻ tiếp tuyến chung ngoài $MN$ với $M \in (O), N \in (O')$. Gọi $P, Q$ lần lượt là điểm đối xứng với $M, N$ qua đường nối tâm $OO'$. Khi đó tứ giác $MNPQ$ là hình gì?

@@TAB@@\textbf{A.} Hình thang cân.@@TAB@@\textbf{B.} Hình thang.@@TAB@@\textbf{C.} Hình thang vuông.@@TAB@@\textbf{D.} Hình bình hành.

\textbf{Câu 10.} Cho đường tròn $(O; 6\text{ cm})$ và $(O'; 2\text{ cm})$ cắt nhau tại $A, B$ sao cho $OA$ là tiếp tuyến của $(O')$. Độ dài dây chung $AB$ bằng

@@TAB@@\textbf{A.} $AB = 3\sqrt{10}\text{ cm}$.@@TAB@@\textbf{B.} $AB = \dfrac{6\sqrt{10}}{5}\text{ cm}$.

@@TAB@@\textbf{C.} $AB = \dfrac{3\sqrt{10}}{5}\text{ cm}$.@@TAB@@\textbf{D.} $AB = \dfrac{\sqrt{10}}{5}\text{ cm}$.

\textbf{Câu 11.} Cho các đường tròn $(O_1; 10\text{ cm})$, $(O_2; 15\text{ cm})$ và $(O_3; 15\text{ cm})$ tiếp xúc ngoài với nhau đôi một. Hai đường tròn $(O_2)$ và $(O_3)$ tiếp xúc nhau tại điểm $A$. Đường tròn $(O_1)$ tiếp xúc với $(O_2)$ và $(O_3)$ lần lượt tại $C$ và $B$. Khẳng định nào sau đây là đúng?

@@TAB@@\textbf{A.} $O_1A$ là tiếp tuyến chung của $(O_2)$ và $(O_3)$.@@TAB@@\textbf{B.} $O_1A = 25\text{ cm}$.

@@TAB@@\textbf{C.} $O_1A = 15\text{ cm}$.@@TAB@@\textbf{D.} Cả A và B đều đúng.

\textbf{Câu 12.} Cho hai đường tròn $(O), (O')$ tiếp xúc ngoài tại $A$. Kẻ các đường kính $AB$ của $(O)$ và $AC$ của $(O')$, gọi $DE$ là tiếp tuyến chung ngoài của hai đường tròn $(D \in (O), E \in (O'))$. Gọi $M$ là giao điểm của $BD$ và $CE$. Tính diện tích tứ giác $ADME$ biết $\widehat{DOA} = 60^\circ$ và $OA = 6\text{ cm}$.

@@TAB@@\textbf{A.} $12\sqrt{3}\text{ cm}^2$.@@TAB@@\textbf{B.} $12\text{ cm}^2$.@@TAB@@\textbf{C.} $16\text{ cm}^2$.@@TAB@@\textbf{D.} $24\text{ cm}^2$.

\textbf{Câu 13.} Cho hai đường tròn $(O), (O')$ tiếp xúc ngoài tại $A$. Kẻ các đường kính $AB$ của $(O)$ và $AC$ của $(O')$, gọi $DE$ là tiếp tuyến chung ngoài của hai đường tròn $(D \in (O), E \in (O'))$. Gọi $M$ là giao điểm của $BD$ và $CE$. Tính diện tích tứ giác $ADME$ biết $\widehat{DOA} = 60^\circ$ và $OA = 8\text{ cm}$.

@@TAB@@\textbf{A.} $12\sqrt{3}\text{ cm}^2$.@@TAB@@\textbf{B.} $\dfrac{64\sqrt{3}}{3}\text{ cm}^2$.@@TAB@@\textbf{C.} $\dfrac{32\sqrt{3}}{3}\text{ cm}^2$.@@TAB@@\textbf{D.} $36\text{ cm}^2$.

\textbf{Câu 14.} Cho các đường tròn $(O_1; 10\text{ cm})$, $(O_2; 15\text{ cm})$ và $(O_3; 15\text{ cm})$ tiếp xúc ngoài với nhau đôi một. Hai đường tròn $(O_2)$ và $(O_3)$ tiếp xúc nhau tại điểm $A$. Đường tròn $(O_1)$ tiếp xúc với $(O_2)$ và $(O_3)$ lần lượt tại $C$ và $B$. Tính diện tích tam giác $ABC$.

@@TAB@@\textbf{A.} $36\text{ cm}^2$.@@TAB@@\textbf{B.} $72\text{ cm}^2$.@@TAB@@\textbf{C.} $144\text{ cm}^2$.@@TAB@@\textbf{D.} $96\text{ cm}^2$.

\textbf{Câu 15.} Cho hai đường tròn $(O; R)$ và $(O'; r)$ với $R > r$ tiếp xúc ngoài tại $A$. Vẽ các bán kính $OB \parallel O'D$ với $B, D$ cùng thuộc nửa mặt phẳng bờ $OO'$. Đường thẳng $DB$ và $OO'$ cắt nhau tại $I$. Tiếp tuyến chung ngoài $GH$ của $(O)$ và $(O')$ với $G, H$ nằm ở nửa mặt phẳng bờ $OO'$ không chứa $B, D$. Tính $OI$ theo $R$ và $r$.

@@TAB@@\textbf{A.} $OI = \dfrac{R+r}{R-r}$.@@TAB@@\textbf{B.} $OI = \dfrac{R-r}{R+r}$.

@@TAB@@\textbf{C.} $OI = \dfrac{R(R-r)}{R+r}$.@@TAB@@\textbf{D.} $OI = \dfrac{R(R+r)}{R-r}$.

\vspace{0.4cm}
\begin{center}
\textbf{BẢNG ĐÁP ÁN TRẮC NGHIỆM}
\end{center}

@@ANSWER_KEY_TABLE@@

\vspace{0.4cm}
\begin{center}
\textbf{--------- HẾT ---------}
\end{center}

\end{document}
"""

# LaTeX content for HƯỚNG DẪN GIẢI CHI TIẾT
tex_hdg = r"""\documentclass[11pt,a4paper]{article}
\usepackage[utf8]{vietnam}
\usepackage{amsmath,amssymb}
\usepackage{graphicx}

\begin{document}

@@DOCUMENT_HEADER@@

@@SECTION_1_HEADER@@

\textbf{Câu 1.} Nếu hai đường tròn không cắt nhau thì số điểm chung của hai đường tròn là

@@TAB@@\textbf{A.} $1$.@@TAB@@\textbf{B.} $2$.@@TAB@@\textbf{C.} $3$.@@TAB@@\textbf{D.} $0$.

\textbf{Lời giải:}

Theo định nghĩa về vị trí tương đối giữa hai đường tròn:
- Hai đường tròn cắt nhau: có đúng $2$ điểm chung.
- Hai đường tròn tiếp xúc nhau (tiếp xúc trong hoặc tiếp xúc ngoài): có đúng $1$ điểm chung.
- Hai đường tròn không cắt nhau (ở ngoài nhau, đựng nhau hoặc đồng tâm): không có điểm chung nào (số điểm chung bằng $0$).

Vậy nếu hai đường tròn không cắt nhau thì số điểm chung là $0$.

\textbf{Chọn đáp án D.}

\textbf{Câu 2.} Cho hai đường tròn $(O_1)$ và $(O_2)$ tiếp xúc ngoài tại $A$ và một đường thẳng $\mathrm{d}$ tiếp xúc với $(O_1)$ và $(O_2)$ lần lượt tại $B, C$. Tam giác $ABC$ là

@@TAB@@\textbf{A.} Tam giác cân.@@TAB@@\textbf{B.} Tam giác đều.@@TAB@@\textbf{C.} Tam giác vuông.@@TAB@@\textbf{D.} Tam giác vuông cân.

\textbf{Lời giải:}

\begin{center}
\includegraphics[width=9cm]{fig_c2.png}
\end{center}

Kẻ tiếp tuyến chung trong của hai đường tròn tại tiếp điểm $A$, tiếp tuyến này cắt đường thẳng $\mathrm{d}$ tại $M$.
Theo tính chất của hai tiếp tuyến cắt nhau:
- Đối với đường tròn $(O_1)$: $MA$ và $MB$ là hai tiếp tuyến cắt nhau tại $M \Rightarrow MA = MB$.
- Đối với đường tròn $(O_2)$: $MA$ và $MC$ là hai tiếp tuyến cắt nhau tại $M \Rightarrow MA = MC$.

Suy ra:
\[MA = MB = MC = \dfrac{1}{2}BC\]

Tam giác $ABC$ có đường trung tuyến $AM$ ứng với cạnh $BC$ và có độ dài bằng nửa cạnh $BC$, do đó tam giác $ABC$ vuông tại $A$.

\textbf{Chọn đáp án C.}

\textbf{Câu 3.} Cho đoạn thẳng $OO'$ và điểm $A$ nằm trên đoạn $OO'$ sao cho $OA = 2O'A$. Vị trí tương đối của đường tròn tâm $O$ bán kính $OA$ và đường tròn tâm $O'$ bán kính $O'A$ là

@@TAB@@\textbf{A.} Nằm ngoài nhau.@@TAB@@\textbf{B.} Cắt nhau.@@TAB@@\textbf{C.} Tiếp xúc ngoài.@@TAB@@\textbf{D.} Tiếp xúc trong.

\textbf{Lời giải:}

Vì điểm $A$ nằm trên đoạn thẳng nối hai tâm $OO'$ nên:
\[OO' = OA + O'A\]

Bán kính của đường tròn tâm $O$ là $R = OA$ và bán kính của đường tròn tâm $O'$ là $R' = O'A$.
Khoảng cách giữa hai tâm là $\mathrm{d} = OO' = R + R'$.

Theo hệ thức liên hệ giữa khoảng cách hai tâm và các bán kính, điều kiện $\mathrm{d} = R + R'$ chứng tỏ hai đường tròn $(O; OA)$ và $(O'; O'A)$ tiếp xúc ngoài nhau tại điểm $A$.

\textbf{Chọn đáp án C.}

\textbf{Câu 4.} Cho hai đường tròn $(O), (O')$ cắt nhau tại $A, B$ trong đó $O' \in (O)$. Kẻ đường kính $O'C$ của đường tròn $(O)$. Chọn khẳng định \textbf{sai}.

@@TAB@@\textbf{A.} $AC = CB$.@@TAB@@\textbf{B.} $\widehat{CBO'} = 90^\circ$.

@@TAB@@\textbf{C.} $CA, CB$ là hai tiếp tuyến của $(O')$.@@TAB@@\textbf{D.} $CA, CB$ là hai cát tuyến của $(O')$.

\textbf{Lời giải:}

\begin{center}
\includegraphics[width=8.5cm]{fig_c4.png}
\end{center}

Vì $O'C$ là đường kính của đường tròn $(O)$ và các điểm $A, B \in (O)$ nên:
\[\widehat{CAO'} = 90^\circ \quad \text{và} \quad \widehat{CBO'} = 90^\circ\]
(các góc nội tiếp chắn nửa đường tròn $(O)$).

- Xét đường tròn $(O')$: Bán kính $O'A \perp CA$ tại $A \in (O') \Rightarrow CA$ là tiếp tuyến của $(O')$.
- Tương tự, bán kính $O'B \perp CB$ tại $B \in (O') \Rightarrow CB$ là tiếp tuyến của $(O')$.

Do đó $CA$ và $CB$ là hai tiếp tuyến kẻ từ $C$ đến đường tròn $(O')$, suy ra $CA = CB$ (hay $AC = CB$).
Vậy các khẳng định A, B, C đều đúng.
Khẳng định sai là D vì $CA, CB$ là hai tiếp tuyến chứ không phải cát tuyến của $(O')$.

\textbf{Chọn đáp án D.}

\textbf{Câu 5.} Cho hai đường tròn tiếp xúc ngoài $(O; R)$ và $(O'; r)$ với $R > r$ và khoảng cách giữa hai tâm $OO' = \mathrm{d}$. Khi đó

@@TAB@@\textbf{A.} $\mathrm{d} = R - r$.@@TAB@@\textbf{B.} $\mathrm{d} > R + r$.@@TAB@@\textbf{C.} $R - r < \mathrm{d} < R + r$.@@TAB@@\textbf{D.} $\mathrm{d} = R + r$.

\textbf{Lời giải:}

Theo bảng hệ thức vị trí tương đối giữa hai đường tròn phân biệt:
- Ở ngoài nhau: $\mathrm{d} > R + r$.
- Tiếp xúc ngoài: $\mathrm{d} = R + r$.
- Cắt nhau: $R - r < \mathrm{d} < R + r$.
- Tiếp xúc trong: $\mathrm{d} = R - r > 0$.
- Đựng nhau: $\mathrm{d} < R - r$.

Do đó, khi hai đường tròn $(O; R)$ và $(O'; r)$ tiếp xúc ngoài thì $\mathrm{d} = R + r$.

*(Ghi chú: Trong đề gốc sưu tầm có thể xuất hiện lỗi gõ nhầm dấu thành $\mathrm{d} < R+r$, bản chuẩn xác tương ứng với đáp án là $\mathrm{d} = R+r$).*

\textbf{Chọn đáp án D.}

\textbf{Câu 6.} Cho đường tròn $(O_1; 3\text{ cm})$ tiếp xúc ngoài với đường tròn $(O_2; 1\text{ cm})$. Vẽ bán kính $O_1B$ và $O_2C$ song song với nhau cùng thuộc nửa mặt phẳng bờ $O_1O_2$. Gọi $D$ là giao điểm của $BC$ và $O_1O_2$. Tính số đo góc $\widehat{BAC}$ (với $A$ là tiếp điểm của $(O_1)$ và $(O_2)$).

@@TAB@@\textbf{A.} $90^\circ$.@@TAB@@\textbf{B.} $60^\circ$.@@TAB@@\textbf{C.} $100^\circ$.@@TAB@@\textbf{D.} $80^\circ$.

\textbf{Lời giải:}

\begin{center}
\includegraphics[width=8.5cm]{fig_c6.png}
\end{center}

Vì hai đường tròn $(O_1)$ và $(O_2)$ tiếp xúc ngoài tại $A$ nên tiếp điểm $A$ thuộc đoạn nối tâm $O_1O_2$.
Vì $O_1B \parallel O_2C$ và cùng thuộc nửa mặt phẳng bờ $O_1O_2$ nên hai góc trong cùng phía bù nhau:
\[\widehat{BO_1A} + \widehat{CO_2A} = 180^\circ\]

Xét các tam giác cân tại các tâm:
- $\triangle O_1AB$ cân tại $O_1$ ($O_1A = O_1B = R_1$) $\Rightarrow \widehat{O_1AB} = \dfrac{180^\circ - \widehat{BO_1A}}{2}$.
- $\triangle O_2AC$ cân tại $O_2$ ($O_2A = O_2C = R_2$) $\Rightarrow \widehat{O_2AC} = \dfrac{180^\circ - \widehat{CO_2A}}{2}$.

Cộng vế với vế, ta được:
\[\widehat{O_1AB} + \widehat{O_2AC} = \dfrac{360^\circ - (\widehat{BO_1A} + \widehat{CO_2A})}{2} = \dfrac{360^\circ - 180^\circ}{2} = 90^\circ\]

Vì ba điểm $O_1, A, O_2$ thẳng hàng nên góc bẹt $\widehat{O_1AO_2} = 180^\circ$, suy ra:
\[\widehat{BAC} = 180^\circ - (\widehat{O_1AB} + \widehat{O_2AC}) = 180^\circ - 90^\circ = 90^\circ\]

\textbf{Chọn đáp án A.}

\textbf{Câu 7.} Cho hai đường tròn $(O; 10\text{ cm})$ và $(O'; 5\text{ cm})$ cắt nhau tại $A, B$. Tính độ dài đoạn thẳng $OO'$ biết $AB = 8\text{ cm}$ và $O, O'$ nằm cùng phía đối với $AB$ (làm tròn kết quả đến chữ số thập phân thứ nhất).

@@TAB@@\textbf{A.} $OO' \approx 6{,}5\text{ cm}$.@@TAB@@\textbf{B.} $OO' \approx 6{,}1\text{ cm}$.@@TAB@@\textbf{C.} $OO' \approx 6{,}3\text{ cm}$.@@TAB@@\textbf{D.} $OO' \approx 6{,}2\text{ cm}$.

\textbf{Lời giải:}

Gọi $H$ là giao điểm của đường nối tâm $OO'$ và dây chung $AB$.
Theo tính chất đường nối tâm của hai đường tròn cắt nhau, $OO'$ là đường trung trực của $AB$, do đó $OO' \perp AB$ tại trung điểm $H$ của $AB$:
\[AH = \dfrac{AB}{2} = \dfrac{8}{2} = 4\text{ cm}\]

Áp dụng định lý Pythagore trong các tam giác vuông $\triangle OHA$ và $\triangle O'HA$:
- Trong $\triangle OHA$ vuông tại $H$:
\[OH = \sqrt{OA^2 - AH^2} = \sqrt{10^2 - 4^2} = \sqrt{100 - 16} = \sqrt{84} \approx 9{,}165\text{ cm}\]
- Trong $\triangle O'HA$ vuông tại $H$:
\[O'H = \sqrt{O'A^2 - AH^2} = \sqrt{5^2 - 4^2} = \sqrt{25 - 16} = \sqrt{9} = 3\text{ cm}\]

Vì hai tâm $O$ và $O'$ nằm cùng phía đối với dây chung $AB$ nên độ dài đoạn nối tâm $OO'$ bằng hiệu khoảng cách:
\[OO' = |OH - O'H| = \sqrt{84} - 3 \approx 9{,}165 - 3 = 6{,}165\text{ cm}\]

Làm tròn đến chữ số thập phân thứ nhất, ta được $OO' \approx 6{,}2\text{ cm}$.

*(Ghi chú: Trong bản Word gốc sưu tầm có sự nhầm lẫn chép số liệu $AB=24\text{ cm}$ của Câu 8 sang. Với đường tròn có bán kính $r=5\text{ cm}$, đường kính lớn nhất chỉ là $10\text{ cm}$ nên dây cung $AB$ không thể dài $24\text{ cm}$. Số liệu chuẩn mực tương ứng với phương án trắc nghiệm là $AB = 8\text{ cm}$).*

\textbf{Chọn đáp án D.}

\textbf{Câu 8.} Cho hai đường tròn $(O; 20\text{ cm})$ và $(O'; 15\text{ cm})$ cắt nhau tại $A, B$. Tính độ dài đoạn thẳng $OO'$ biết $AB = 24\text{ cm}$ và $O, O'$ nằm cùng phía đối với $AB$.

@@TAB@@\textbf{A.} $OO' = 7\text{ cm}$.@@TAB@@\textbf{B.} $OO' = 8\text{ cm}$.@@TAB@@\textbf{C.} $OO' = 9\text{ cm}$.@@TAB@@\textbf{D.} $OO' = 25\text{ cm}$.

\textbf{Lời giải:}

\begin{center}
\includegraphics[width=8cm]{fig_c8.png}
\end{center}

Gọi $H$ là giao điểm của $OO'$ và $AB$.
Vì $OO'$ là đường trung trực của dây chung $AB$ nên $OO' \perp AB$ tại trung điểm $H$ của $AB$:
\[AH = \dfrac{AB}{2} = \dfrac{24}{2} = 12\text{ cm}\]

Áp dụng định lý Pythagore trong các tam giác vuông:
- Trong $\triangle OHA$ vuông tại $H$:
\[OH = \sqrt{OA^2 - AH^2} = \sqrt{20^2 - 12^2} = \sqrt{400 - 144} = \sqrt{256} = 16\text{ cm}\]
- Trong $\triangle O'HA$ vuông tại $H$:
\[O'H = \sqrt{O'A^2 - AH^2} = \sqrt{15^2 - 12^2} = \sqrt{225 - 144} = \sqrt{81} = 9\text{ cm}\]

Vì $O$ và $O'$ nằm cùng phía đối với $AB$ nên khoảng cách giữa hai tâm là:
\[OO' = OH - O'H = 16 - 9 = 7\text{ cm}\]

\textbf{Chọn đáp án A.}

\textbf{Câu 9.} Cho hai đường tròn $(O), (O')$ tiếp xúc ngoài tại $A$. Kẻ tiếp tuyến chung ngoài $MN$ với $M \in (O), N \in (O')$. Gọi $P, Q$ lần lượt là điểm đối xứng với $M, N$ qua đường nối tâm $OO'$. Khi đó tứ giác $MNPQ$ là hình gì?

@@TAB@@\textbf{A.} Hình thang cân.@@TAB@@\textbf{B.} Hình thang.@@TAB@@\textbf{C.} Hình thang vuông.@@TAB@@\textbf{D.} Hình bình hành.

\textbf{Lời giải:}

\begin{center}
\includegraphics[width=8cm]{fig_c9.png}
\end{center}

Vì $P$ đối xứng với $M$ qua đường nối tâm $OO'$ nên $OO'$ là đường trung trực của đoạn thẳng $MP \Rightarrow MP \perp OO'$ và $P \in (O)$.
Tương tự, $Q$ đối xứng với $N$ qua $OO'$ nên $OO'$ là đường trung trực của $NQ \Rightarrow NQ \perp OO'$ và $Q \in (O')$.
Từ $MP \perp OO'$ và $NQ \perp OO'$ suy ra:
\[MP \parallel NQ\]

Do đó, tứ giác $MNPQ$ có hai cạnh đối song song nên là một hình thang.
Hơn nữa, phép đối xứng trục qua đường thẳng $OO'$ biến đoạn thẳng $MN$ thành đoạn thẳng $PQ$, suy ra $MN = PQ$.
Hình thang $MNPQ$ nhận đường thẳng $OO'$ làm trục đối xứng và có hai cạnh bên bằng nhau nên $MNPQ$ là \textbf{hình thang cân}.

\textbf{Chọn đáp án A.}

\textbf{Câu 10.} Cho đường tròn $(O; 6\text{ cm})$ và $(O'; 2\text{ cm})$ cắt nhau tại $A, B$ sao cho $OA$ là tiếp tuyến của $(O')$. Độ dài dây chung $AB$ bằng

@@TAB@@\textbf{A.} $AB = 3\sqrt{10}\text{ cm}$.@@TAB@@\textbf{B.} $AB = \dfrac{6\sqrt{10}}{5}\text{ cm}$.

@@TAB@@\textbf{C.} $AB = \dfrac{3\sqrt{10}}{5}\text{ cm}$.@@TAB@@\textbf{D.} $AB = \dfrac{\sqrt{10}}{5}\text{ cm}$.

\textbf{Lời giải:}

\begin{center}
\includegraphics[width=8cm]{fig_c10.png}
\end{center}

Vì $OA$ là tiếp tuyến của $(O')$ tại $A$ nên $OA \perp O'A$ tại $A$.
Xét tam giác $OAO'$ vuông tại $A$, có $OA = 6\text{ cm}$, $O'A = 2\text{ cm}$.
Theo định lý Pythagore:
\[OO' = \sqrt{OA^2 + O'A^2} = \sqrt{6^2 + 2^2} = \sqrt{36 + 4} = \sqrt{40} = 2\sqrt{10}\text{ cm}\]

Gọi $H$ là giao điểm của $OO'$ và dây chung $AB$. Khi đó $OO' \perp AB$ tại trung điểm $H$ của $AB$.
Trong tam giác vuông $OAO'$, đường cao $AH$ thỏa mãn hệ thức lượng:
\[AH \cdot OO' = OA \cdot O'A \implies AH = \dfrac{OA \cdot O'A}{OO'} = \dfrac{6 \cdot 2}{2\sqrt{10}} = \dfrac{6}{\sqrt{10}} = \dfrac{3\sqrt{10}}{5}\text{ cm}\]

Do $H$ là trung điểm của $AB$ nên:
\[AB = 2AH = 2 \cdot \dfrac{3\sqrt{10}}{5} = \dfrac{6\sqrt{10}}{5}\text{ cm}\]

\textbf{Chọn đáp án B.}

\textbf{Câu 11.} Cho các đường tròn $(O_1; 10\text{ cm})$, $(O_2; 15\text{ cm})$ và $(O_3; 15\text{ cm})$ tiếp xúc ngoài với nhau đôi một. Hai đường tròn $(O_2)$ và $(O_3)$ tiếp xúc nhau tại điểm $A$. Đường tròn $(O_1)$ tiếp xúc với $(O_2)$ và $(O_3)$ lần lượt tại $C$ và $B$. Khẳng định nào sau đây là đúng?

@@TAB@@\textbf{A.} $O_1A$ là tiếp tuyến chung của $(O_2)$ và $(O_3)$.@@TAB@@\textbf{B.} $O_1A = 25\text{ cm}$.

@@TAB@@\textbf{C.} $O_1A = 15\text{ cm}$.@@TAB@@\textbf{D.} Cả A và B đều đúng.

\textbf{Lời giải:}

\begin{center}
\includegraphics[width=8cm]{fig_c11.png}
\end{center}

Khoảng cách giữa các tâm:
- $O_2O_3 = R_2 + R_3 = 15 + 15 = 30\text{ cm}$.
- $O_1O_2 = R_1 + R_2 = 10 + 15 = 25\text{ cm}$.
- $O_1O_3 = R_1 + R_3 = 10 + 15 = 25\text{ cm}$.

Vì $O_1O_2 = O_1O_3 = 25\text{ cm}$ nên $\triangle O_1O_2O_3$ cân tại $O_1$.
Điểm $A$ là tiếp điểm của $(O_2)$ và $(O_3)$ nên $O_2A = R_2 = 15\text{ cm}$ và $O_3A = R_3 = 15\text{ cm}$, suy ra $A$ là trung điểm của cạnh $O_2O_3$.
Trong tam giác cân $O_1O_2O_3$, đoạn nối đỉnh với trung điểm đáy $O_1A$ đồng thời là đường cao:
\[O_1A \perp O_2O_3 \text{ tại tiếp điểm } A\]

Vì $O_1A \perp O_2O_3$ tại tiếp điểm chung $A$ nên $O_1A$ là tiếp tuyến chung của hai đường tròn $(O_2)$ và $(O_3)$ (khẳng định A đúng).

Tính độ dài $O_1A$:
\[O_1A = \sqrt{O_1O_2^2 - O_2A^2} = \sqrt{25^2 - 15^2} = \sqrt{625 - 225} = \sqrt{400} = 20\text{ cm}\]

Vì $O_1A = 20\text{ cm} \neq 25\text{ cm}$ và $\neq 15\text{ cm}$ nên các phương án B, C, D đều sai.

\textbf{Chọn đáp án A.}

\textbf{Câu 12.} Cho hai đường tròn $(O), (O')$ tiếp xúc ngoài tại $A$. Kẻ các đường kính $AB$ của $(O)$ và $AC$ của $(O')$, gọi $DE$ là tiếp tuyến chung ngoài của hai đường tròn $(D \in (O), E \in (O'))$. Gọi $M$ là giao điểm của $BD$ và $CE$. Tính diện tích tứ giác $ADME$ biết $\widehat{DOA} = 60^\circ$ và $OA = 6\text{ cm}$.

@@TAB@@\textbf{A.} $12\sqrt{3}\text{ cm}^2$.@@TAB@@\textbf{B.} $12\text{ cm}^2$.@@TAB@@\textbf{C.} $16\text{ cm}^2$.@@TAB@@\textbf{D.} $24\text{ cm}^2$.

\textbf{Lời giải:}

\begin{center}
\includegraphics[width=8cm]{fig_c12.png}
\end{center}

Vì $AB$ là đường kính của $(O)$ nên $\widehat{ADB} = 90^\circ$ (góc nội tiếp chắn nửa đường tròn) $\Rightarrow AD \perp BD$ tại $D$, hay $\widehat{ADM} = 90^\circ$.
Tương tự, $AC$ là đường kính của $(O')$ nên $\widehat{AEC} = 90^\circ \Rightarrow AE \perp CE$ tại $E$, hay $\widehat{AEM} = 90^\circ$.
Lại có $DE$ là tiếp tuyến chung ngoài của hai đường tròn tiếp xúc ngoài tại $A$, theo tính chất tiếp tuyến chung:
\[\widehat{DAE} = 90^\circ\]

Tứ giác $ADME$ có ba góc vuông ($\widehat{DAE} = \widehat{ADM} = \widehat{AEM} = 90^\circ$) nên là một \textbf{hình chữ nhật}.
Diện tích hình chữ nhật $ADME$ là:
\[S_{ADME} = AD \cdot AE\]

Xét tam giác $OAD$: Ta có $OA = OD = R = 6\text{ cm} \Rightarrow \triangle OAD$ cân tại $O$.
Mà $\widehat{DOA} = 60^\circ$ nên $\triangle OAD$ là tam giác đều:
\[AD = OA = 6\text{ cm} \quad \text{và} \quad \widehat{ODA} = 60^\circ\]

Vì $DE$ là tiếp tuyến của $(O)$ tại $D$ nên $OD \perp DE$, do đó:
\[\widehat{ADE} = 90^\circ - \widehat{ODA} = 90^\circ - 60^\circ = 30^\circ\]

Trong tam giác vuông $ADE$ tại $A$:
\[AE = AD \cdot \tan \widehat{ADE} = 6 \cdot \tan 30^\circ = 6 \cdot \dfrac{1}{\sqrt{3}} = 2\sqrt{3}\text{ cm}\]

Vậy diện tích hình chữ nhật $ADME$ là:
\[S_{ADME} = AD \cdot AE = 6 \cdot 2\sqrt{3} = 12\sqrt{3}\text{ cm}^2\]

\textbf{Chọn đáp án A.}

\textbf{Câu 13.} Cho hai đường tròn $(O), (O')$ tiếp xúc ngoài tại $A$. Kẻ các đường kính $AB$ của $(O)$ và $AC$ của $(O')$, gọi $DE$ là tiếp tuyến chung ngoài của hai đường tròn $(D \in (O), E \in (O'))$. Gọi $M$ là giao điểm của $BD$ và $CE$. Tính diện tích tứ giác $ADME$ biết $\widehat{DOA} = 60^\circ$ và $OA = 8\text{ cm}$.

@@TAB@@\textbf{A.} $12\sqrt{3}\text{ cm}^2$.@@TAB@@\textbf{B.} $\dfrac{64\sqrt{3}}{3}\text{ cm}^2$.@@TAB@@\textbf{C.} $\dfrac{32\sqrt{3}}{3}\text{ cm}^2$.@@TAB@@\textbf{D.} $36\text{ cm}^2$.

\textbf{Lời giải:}

Chứng minh tương tự Câu 12:
Tứ giác $ADME$ là hình chữ nhật.
Tam giác $OAD$ cân tại $O$ và có $\widehat{DOA} = 60^\circ$ nên là tam giác đều, suy ra:
\[AD = OA = 8\text{ cm} \quad \text{và} \quad \widehat{ODA} = 60^\circ\]

Vì $OD \perp DE$ nên:
\[\widehat{ADE} = 90^\circ - \widehat{ODA} = 90^\circ - 60^\circ = 30^\circ\]

Trong tam giác vuông $ADE$ tại $A$:
\[AE = AD \cdot \tan 30^\circ = 8 \cdot \dfrac{\sqrt{3}}{3} = \dfrac{8\sqrt{3}}{3}\text{ cm}\]

Vậy diện tích hình chữ nhật $ADME$ là:
\[S_{ADME} = AD \cdot AE = 8 \cdot \dfrac{8\sqrt{3}}{3} = \dfrac{64\sqrt{3}}{3}\text{ cm}^2\]

\textbf{Chọn đáp án B.}

\textbf{Câu 14.} Cho các đường tròn $(O_1; 10\text{ cm})$, $(O_2; 15\text{ cm})$ và $(O_3; 15\text{ cm})$ tiếp xúc ngoài với nhau đôi một. Hai đường tròn $(O_2)$ và $(O_3)$ tiếp xúc nhau tại điểm $A$. Đường tròn $(O_1)$ tiếp xúc với $(O_2)$ và $(O_3)$ lần lượt tại $C$ và $B$. Tính diện tích tam giác $ABC$.

@@TAB@@\textbf{A.} $36\text{ cm}^2$.@@TAB@@\textbf{B.} $72\text{ cm}^2$.@@TAB@@\textbf{C.} $144\text{ cm}^2$.@@TAB@@\textbf{D.} $96\text{ cm}^2$.

\textbf{Lời giải:}

Chọn hệ trục tọa độ vuông góc $Oxy$ với gốc đặt tại tiếp điểm $A(0; 0)$ trên đường nối tâm $O_2O_3$:
- Trục hoành $Ox$ trùng với đường nối tâm $O_2O_3$, khi đó $O_2(-15; 0)$ và $O_3(15; 0)$.
- Trục tung $Oy$ vuông góc với $O_2O_3$ tại $A$, đi qua tâm $O_1$.

Vì $O_1O_2 = R_1 + R_2 = 10 + 15 = 25\text{ cm}$ nên tung độ của điểm $O_1$ là:
\[y_{O_1} = \sqrt{O_1O_2^2 - O_2A^2} = \sqrt{25^2 - 15^2} = \sqrt{400} = 20 \implies O_1(0; 20)\]

Tìm tọa độ các tiếp điểm:
- Điểm $C$ thuộc đoạn thẳng $O_2O_1$ và chia đoạn $O_2O_1$ theo tỉ số $\dfrac{O_2C}{O_2O_1} = \dfrac{15}{25} = \dfrac{3}{5}$:
\[x_C = x_{O_2} + \dfrac{3}{5}(x_{O_1} - x_{O_2}) = -15 + \dfrac{3}{5}(0 - (-15)) = -6\]
\[y_C = y_{O_2} + \dfrac{3}{5}(y_{O_1} - y_{O_2}) = 0 + \dfrac{3}{5}(20 - 0) = 12 \implies C(-6; 12)\]

- Tương tự, điểm $B$ thuộc đoạn thẳng $O_3O_1$ chia đoạn $O_3O_1$ theo tỉ số $\dfrac{O_3B}{O_3O_1} = \dfrac{3}{5}$:
\[x_B = x_{O_3} + \dfrac{3}{5}(x_{O_1} - x_{O_3}) = 15 + \dfrac{3}{5}(0 - 15) = 6\]
\[y_B = y_{O_3} + \dfrac{3}{5}(y_{O_1} - y_{O_3}) = 0 + \dfrac{3}{5}(20 - 0) = 12 \implies B(6; 12)\]

Tam giác $ABC$ có các đỉnh $A(0; 0)$, $B(6; 12)$, $C(-6; 12)$:
- Cạnh đáy $BC$ nằm ngang trên đường thẳng $y = 12$, có độ dài $BC = x_B - x_C = 6 - (-6) = 12\text{ cm}$.
- Chiều cao $h$ hạ từ đỉnh $A(0; 0)$ đến đường thẳng chứa cạnh $BC$ ($y = 12$) là $h = 12\text{ cm}$.

Diện tích tam giác $ABC$ là:
\[S_{\triangle ABC} = \dfrac{1}{2} \cdot BC \cdot h = \dfrac{1}{2} \cdot 12 \cdot 12 = 72\text{ cm}^2\]

\textbf{Chọn đáp án B.}

\textbf{Câu 15.} Cho hai đường tròn $(O; R)$ và $(O'; r)$ với $R > r$ tiếp xúc ngoài tại $A$. Vẽ các bán kính $OB \parallel O'D$ với $B, D$ cùng thuộc nửa mặt phẳng bờ $OO'$. Đường thẳng $DB$ và $OO'$ cắt nhau tại $I$. Tiếp tuyến chung ngoài $GH$ của $(O)$ và $(O')$ với $G, H$ nằm ở nửa mặt phẳng bờ $OO'$ không chứa $B, D$. Tính $OI$ theo $R$ và $r$.

@@TAB@@\textbf{A.} $OI = \dfrac{R+r}{R-r}$.@@TAB@@\textbf{B.} $OI = \dfrac{R-r}{R+r}$.

@@TAB@@\textbf{C.} $OI = \dfrac{R(R-r)}{R+r}$.@@TAB@@\textbf{D.} $OI = \dfrac{R(R+r)}{R-r}$.

\textbf{Lời giải:}

\begin{center}
\includegraphics[width=8cm]{fig_c15.png}
\end{center}

Xét tam giác $IOB$ có $O'D \parallel OB$ ($D \in IB, O' \in IO$).
Theo định lý Ta-lét trong tam giác:
\[\dfrac{IO'}{IO} = \dfrac{O'D}{OB} = \dfrac{r}{R}\]

Do hai đường tròn $(O; R)$ và $(O'; r)$ tiếp xúc ngoài tại $A$ nên khoảng cách hai tâm là:
\[OO' = OA + O'A = R + r\]

Mặt khác, vì $R > r$ nên điểm $O'$ nằm giữa hai điểm $I$ và $O$, ta có:
\[IO - IO' = OO' = R + r\]

Từ $\dfrac{IO'}{IO} = \dfrac{r}{R}$, áp dụng tính chất của dãy tỉ số bằng nhau:
\[\dfrac{IO - IO'}{IO} = \dfrac{R - r}{R}\]

Thay $IO - IO' = R + r$ vào hệ thức:
\[\dfrac{R + r}{OI} = \dfrac{R - r}{R} \implies OI = \dfrac{R(R + r)}{R - r}\]

\textbf{Chọn đáp án D.}

\vspace{0.4cm}
\begin{center}
\textbf{BẢNG ĐÁP ÁN TRẮC NGHIỆM}
\end{center}

@@ANSWER_KEY_TABLE@@

\vspace{0.4cm}
\begin{center}
\textbf{--------- HẾT ---------}
\end{center}

\end{document}
"""

def build_tn9_doc(tex_content, output_path, is_solution=False):
    scratch_dir = os.path.join(CURRENT_DIR, "scratch_tn9_bai17")
    os.makedirs(scratch_dir, exist_ok=True)
    
    file_type = "hdg" if is_solution else "de"
    temp_tex = os.path.join(scratch_dir, f"temp_{file_type}.tex")
    temp_docx = os.path.join(scratch_dir, f"temp_{file_type}.docx")
    
    with open(temp_tex, "w", encoding="utf-8") as f:
        f.write(tex_content)
        
    cmd = [pandoc_exe, temp_tex, "-o", temp_docx, "--from=latex", "--to=docx", f"--resource-path={LATEX_DIR};{HINH_ANH_DIR}"]
    print(f"[*] Đang xuất bản Pandoc: {os.path.basename(output_path)}...")
    res = subprocess.run(cmd, cwd=LATEX_DIR, capture_output=True, text=True, encoding="utf-8")
    if res.returncode != 0:
        print("[LỖI PANDOC]:", res.stderr)
        return
        
    doc = docx.Document(temp_docx)
    
    # Thiết lập lề trang chuẩn Macro BTPro
    for s in doc.sections:
        s.top_margin = Cm(1.6)
        s.bottom_margin = Cm(1.6)
        s.left_margin = Cm(2.0)
        s.right_margin = Cm(1.5)
        s.page_width = Cm(21.0)
        s.page_height = Cm(29.7)
        
        # Footer
        footer = s.footer
        f_p = footer.paragraphs[0]
        f_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        f_p.text = ""
        r_l = f_p.add_run("Lớp Toán Cô Thúy  •  SĐT: 0935.322.328  •  50/2C Phạm Thị Liên                                   ")
        r_l.font.name = "Times New Roman"
        r_l.font.size = Pt(9.5)
        r_l.italic = True
        r_l.font.color.rgb = RGBColor(100, 100, 100)
        
        r_r = f_p.add_run("Bài 17: Vị trí tương đối của hai đường tròn")
        r_r.font.name = "Times New Roman"
        r_r.font.size = Pt(9.5)
        r_r.italic = True
        r_r.font.color.rgb = RGBColor(80, 80, 80)
        
    style = doc.styles['Normal']
    style.font.name = 'Times New Roman'
    style.font.size = Pt(11)
    style.font.color.rgb = RGBColor(0, 0, 0)
    rPr = style.element.get_or_add_rPr()
    rFonts = parse_xml(r'<w:rFonts %s w:ascii="Times New Roman" w:hAnsi="Times New Roman" w:cs="Times New Roman"/>' % nsdecls('w'))
    rPr.append(rFonts)
    
    # Duyệt và thay thế Header sections & Answer table
    for p in list(doc.paragraphs):
        if "@@DOCUMENT_HEADER@@" in p.text:
            insert_header_and_code(doc, p, is_solution)
        elif "@@SECTION_1_HEADER@@" in p.text:
            insert_section_header(doc, p, "PHẦN I. Câu trắc nghiệm nhiều phương án lựa chọn.", "Thí sinh trả lời từ câu 1 đến câu 15. Mỗi câu hỏi thí sinh chỉ chọn một phương án.")
        elif "@@ANSWER_KEY_TABLE@@" in p.text:
            insert_answer_key_table(doc, p)
            
    # Định dạng các đoạn văn bản
    for p in doc.paragraphs:
        format_doc_paragraph(p)
        
    doc.save(output_path)
    print(f"[THÀNH CÔNG] Đã xuất bản file Word: {output_path}")

print("=== BẮT ĐẦU XUẤT BẢN WORD CHUẨN MACRO BTPRO ===")
build_tn9_doc(tex_de, OUTPUT_DE_DOCX, is_solution=False)
build_tn9_doc(tex_hdg, OUTPUT_HDG_DOCX, is_solution=True)

# Copy PDFs to San_Pham
def safe_copy(src, dst):
    try:
        shutil.copyfile(src, dst)
        print(f"[THÀNH CÔNG] Đã sao chép PDF: {os.path.basename(dst)}")
    except PermissionError:
        print(f"[CHÚ Ý] File {os.path.basename(dst)} đang được mở trong Acrobat Reader. Bản cập nhật mới nằm tại {src}.")

safe_copy(os.path.join(LATEX_DIR, "TN9_Chuong5_Bai17_De.pdf"), OUTPUT_DE_PDF)
safe_copy(os.path.join(LATEX_DIR, "TN9_Chuong5_Bai17_HDG.pdf"), OUTPUT_HDG_PDF)
