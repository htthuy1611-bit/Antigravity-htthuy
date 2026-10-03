import os
import sys
import subprocess
import docx
from docx.shared import Pt, Cm, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn

if sys.stdout.encoding.lower() != 'utf-8':
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except Exception:
        pass

sys.path.append(os.path.dirname(__file__))
from bao_thang_core import (
    create_answer_box_table,
    format_doc_paragraph,
    insert_header_and_code,
    insert_section_header
)

CURRENT_DIR = os.path.dirname(os.path.abspath(__file__))
BASE_DIR = os.path.abspath(os.path.join(CURRENT_DIR, "..", ".."))
HINH_ANH_0104 = os.path.join(BASE_DIR, "He_Thong", "Hinh_Anh", "Bao_Thang_3", "0104")
SAN_PHAM_DIR = os.path.join(BASE_DIR, "San_Pham")
pandoc_exe = os.path.expandvars(r"%LOCALAPPDATA%\Pandoc\pandoc.exe")

OUTPUT_DE_0104 = os.path.join(SAN_PHAM_DIR, "de-khao-sat-toan-12-nam-2026-2027-truong-thpt-bao-thang-3-lao-cai_De_0104.docx")
OUTPUT_HDG_0104 = os.path.join(SAN_PHAM_DIR, "de-khao-sat-toan-12-nam-2026-2027-truong-thpt-bao-thang-3-lao-cai_HDG_0104.docx")

tex_de_0104 = r"""\documentclass[12pt,a4paper]{article}
\usepackage[utf8]{vietnam}
\usepackage{amsmath,amssymb}
\usepackage{graphicx}
\usepackage[top=1.6cm,bottom=1.6cm,left=2.0cm,right=1.5cm]{geometry}

\begin{document}

@@DOCUMENT_HEADER@@

@@SECTION_1_HEADER@@

% Câu 1
\textbf{Câu 1.} Cho hàm số $y = f(x)$ có đạo hàm $f'(x) = 3(x - 1)(x - 3)$ với mọi $x \in \mathbb{R}$. Khoảng nghịch biến của hàm số là

@@TAB@@\textbf{A.} $(-\infty; 1)$.@@TAB@@\textbf{B.} $(3; +\infty)$.@@TAB@@\textbf{C.} $(-\infty; 3)$.@@TAB@@\textbf{D.} $(1; 3)$.

% Câu 2 (BBT)
\textbf{Câu 2.} Cho hàm số $y = f(x)$ có bảng biến thiên như dưới đây.

@@CENTER_IMAGE_c2_bbt@@

Giá trị cực đại của hàm số là

@@TAB@@\textbf{A.} $0$.@@TAB@@\textbf{B.} $7$.@@TAB@@\textbf{C.} $4$.@@TAB@@\textbf{D.} $-1$.

% Câu 3
\textbf{Câu 3.} Cho hàm số $y = \dfrac{3x - 2}{x + 4}$. Phương trình tiệm cận đứng là

@@TAB@@\textbf{A.} $x = -4$.@@TAB@@\textbf{B.} $x = 4$.@@TAB@@\textbf{C.} $y = 3$.@@TAB@@\textbf{D.} $y = -4$.

% Câu 4
\textbf{Câu 4.} Cho hàm số $y = \dfrac{-2x + 5}{x - 1}$. Phương trình tiệm cận ngang là

@@TAB@@\textbf{A.} $y = 2$.@@TAB@@\textbf{B.} $x = 1$.@@TAB@@\textbf{C.} $y = -2$.@@TAB@@\textbf{D.} $y = 5$.

% Câu 5
\textbf{Câu 5.} Cho hàm số $y = x - 1 + \dfrac{2}{x - 2}$. Phương trình đường tiệm cận xiên của đồ thị hàm số là

@@TAB@@\textbf{A.} $y = x + 1$.@@TAB@@\textbf{B.} $y = x - 1$.@@TAB@@\textbf{C.} $y = -x + 1$.@@TAB@@\textbf{D.} $x = 2$.

% Câu 6
\textbf{Câu 6.} Đồ thị hàm số $y = -3 + \dfrac{5}{x + 2}$ có tâm đối xứng là

@@TAB@@\textbf{A.} $I(2; -3)$.@@TAB@@\textbf{B.} $I(-3; 2)$.@@TAB@@\textbf{C.} $I(-2; 3)$.@@TAB@@\textbf{D.} $I(-2; -3)$.

% Câu 7
\textbf{Câu 7.} Cho hàm số $y = f(x)$ có đạo hàm $f'(x) = (x + 1)^2(x - 2)(x - 5)$ với mọi $x \in \mathbb{R}$. Số điểm cực trị của hàm số là

@@TAB@@\textbf{A.} $1$.@@TAB@@\textbf{B.} $3$.@@TAB@@\textbf{C.} $4$.@@TAB@@\textbf{D.} $2$.

% Câu 8
\textbf{Câu 8.} Cho hàm số $y = x + 1 + \dfrac{3}{x - 1}$. Giao điểm hai đường tiệm cận của đồ thị hàm số là

@@TAB@@\textbf{A.} $I(1; 1)$.@@TAB@@\textbf{B.} $I(1; 2)$.@@TAB@@\textbf{C.} $I(-1; 0)$.@@TAB@@\textbf{D.} $I(2; 1)$.

% Câu 9 (Side-by-side)
@@START_SIDE_BY_SIDE_C9@@

\textbf{Câu 9.} Cho hàm số phân thức $y = \dfrac{ax + b}{cx + d}$ ($ac \ne 0$, $ad - bc \ne 0$) có đồ thị như hình bên. Hàm số nào dưới đây có đồ thị như hình đã cho?

@@TAB@@\textbf{A.} $y = 1 + \dfrac{2}{x + 1}$.@@TAB@@\textbf{B.} $y = 1 - \dfrac{2}{x + 1}$.

@@TAB@@\textbf{C.} $y = 1 + \dfrac{2}{x - 1}$.@@TAB@@\textbf{D.} $y = -1 + \dfrac{2}{x + 1}$.

@@END_SIDE_BY_SIDE_C9@@

% Câu 10
\textbf{Câu 10.} Cho hàm số $y = f(x)$ xác định trên $\mathbb{R} \setminus \{0\}$ và có đạo hàm $f'(x) = \dfrac{x^2 - 1}{x^2}$. Số điểm cực trị của hàm số là

@@TAB@@\textbf{A.} $0$.@@TAB@@\textbf{B.} $1$.@@TAB@@\textbf{C.} $2$.@@TAB@@\textbf{D.} $3$.

% Câu 11
\textbf{Câu 11.} Cho hàm số $y = f(x) = x^3 - 12x + 1$. Khẳng định nào sau đây đúng?

@@TAB@@\textbf{A.} Hàm số đạt cực trị tại $x = -2$ và $x = 2$.

@@TAB@@\textbf{B.} Hàm số không có cực trị.

@@TAB@@\textbf{C.} Hàm số có một điểm cực trị.

@@TAB@@\textbf{D.} Tâm đối xứng của đồ thị là $I(1; 0)$.

% Câu 12 (BBT)
\textbf{Câu 12.} Cho hàm số bậc ba $y = f(x)$ có bảng biến thiên như dưới đây:

@@CENTER_IMAGE_c12_bbt@@

Hàm số nào dưới đây có bảng biến thiên đã cho?

@@TAB@@\textbf{A.} $y = -x^3 + 3x$.@@TAB@@\textbf{B.} $y = x^3 + 3x$.@@TAB@@\textbf{C.} $y = x^3 - 3x$.@@TAB@@\textbf{D.} $y = x^3 - 3x + 1$.

@@SECTION_2_HEADER@@

% Phần 2 Câu 1 (BBT)
\textbf{Câu 1.} Cho hàm số $y = \dfrac{(x + 1)^2}{x + 2}$. Xét tính đúng, sai của các phát biểu sau:

@@TAB@@\textbf{a)} Đạo hàm là $y' = \dfrac{(x + 3)(x + 1)}{(x + 2)^2}$.

@@TAB@@\textbf{b)} Hàm số nghịch biến trên từng khoảng $(-3; -2)$ và $(-2; -1)$.

@@TAB@@\textbf{c)} Giá trị cực đại của hàm số bằng $0$.

@@TAB@@\textbf{d)} Bảng biến thiên dưới đây là bảng biến thiên của hàm số đã cho.

@@CENTER_IMAGE_p2_c1_bbt@@

% Phần 2 Câu 2
\textbf{Câu 2.} Cho hàm số $y = f(x) = x^3 - 6x^2 + 9x - 2$. Xét tính đúng, sai của các phát biểu sau:

@@TAB@@\textbf{a)} Tập xác định là $\mathbb{R}$.

@@TAB@@\textbf{b)} Hàm số đồng biến trên khoảng $(1; 3)$.

@@TAB@@\textbf{c)} Giá trị cực tiểu của hàm số bằng $-2$.

@@TAB@@\textbf{d)} Điểm $I(2; 1)$ là tâm đối xứng của đồ thị.

% Phần 2 Câu 3 (Side-by-side)
@@START_SIDE_BY_SIDE_P2C3@@

\textbf{Câu 3.} Cho hàm số phân thức $y = \dfrac{ax + b}{cx + d}$ ($ac \ne 0$, $ad - bc \ne 0$) có đồ thị như hình bên. Xét tính đúng, sai của các phát biểu sau:

@@TAB@@\textbf{a)} Đồ thị có tiệm cận đứng $x = -2$.

@@TAB@@\textbf{b)} Đồ thị có tiệm cận ngang $y = 1$.

@@TAB@@\textbf{c)} Tâm đối xứng của đồ thị là $I(-2; -1)$.

@@TAB@@\textbf{d)} Hàm số đồng biến trên từng khoảng xác định.

@@END_SIDE_BY_SIDE_P2C3@@

% Phần 2 Câu 4 (Side-by-side)
@@START_SIDE_BY_SIDE_P2C4@@

\textbf{Câu 4.} Cho hàm số $y = \dfrac{-2x + 1}{x + 1}$. Xét tính đúng, sai của các phát biểu sau:

@@TAB@@\textbf{a)} Đạo hàm là $y' = -\dfrac{3}{(x + 1)^2}$.

@@TAB@@\textbf{b)} Hàm số nghịch biến trên từng khoảng xác định.

@@TAB@@\textbf{c)} Đồ thị có tiệm cận ngang $y = 2$.

@@TAB@@\textbf{d)} Hình bên là đồ thị của hàm số đã cho.

@@END_SIDE_BY_SIDE_P2C4@@

@@SECTION_3_HEADER@@

% Phần 3 Câu 1
\textbf{Câu 1.} Cho hàm số $y = f(x) = -x^3 + 3x + 1$. Tính tổng hai giá trị cực trị của hàm số.

@@ANSWER_BOX_1@@

% Phần 3 Câu 2
\textbf{Câu 2.} Cho hàm số $y = \dfrac{x^2 - x + 2}{x - 2}$. Gọi $I(a; b)$ là giao điểm hai đường tiệm cận của đồ thị. Tính $a + b$.

@@ANSWER_BOX_2@@

% Phần 3 Câu 3
\textbf{Câu 3.} Cho hàm số $y = f(x)$ có đạo hàm $f'(x) = (x + 2)^2(x - 1)(x - 3)$ với mọi $x \in \mathbb{R}$. Tính tổng các giá trị của $x$ tại đó hàm số đạt cực trị.

@@ANSWER_BOX_3@@

% Phần 3 Câu 4
\textbf{Câu 4.} Cho hàm số $y = \dfrac{x^2 - 2x + 5}{x - 1}$. Gọi $A, B$ là hai điểm cực trị của đồ thị hàm số. Tính độ dài đoạn thẳng $AB$ (kết quả làm tròn đến hàng phần trăm).

@@ANSWER_BOX_4@@

% Phần 3 Câu 5
\textbf{Câu 5.} Cho hàm số $y = f(x) = \dfrac{x^2 - x + 4}{x}$. Tính hiệu giữa giá trị cực tiểu và giá trị cực đại của hàm số.

@@ANSWER_BOX_5@@

% Phần 3 Câu 6 (BBT)
\textbf{Câu 6.} Cho hàm số $y = \dfrac{ax + b}{cx + d}$ ($ac \ne 0$, $ad - bc \ne 0$) có bảng biến thiên như dưới đây:

@@CENTER_IMAGE_p3_c6_bbt@@

Gọi $I(u; v)$ là tâm đối xứng của đồ thị hàm số. Tính $u + v$.

@@ANSWER_BOX_6@@

\vspace{0.4cm}
\begin{center}
\textbf{------ HẾT ------}
\end{center}

\end{document}
"""

tex_hdg_0104 = r"""\documentclass[12pt,a4paper]{article}
\usepackage[utf8]{vietnam}
\usepackage{amsmath,amssymb}
\usepackage{graphicx}
\usepackage[top=1.6cm,bottom=1.6cm,left=2.0cm,right=1.5cm]{geometry}

\begin{document}

@@DOCUMENT_HEADER@@

@@SECTION_1_HEADER@@

% Câu 1
\textbf{Câu 1.} Cho hàm số $y = f(x)$ có đạo hàm $f'(x) = 3(x - 1)(x - 3)$ với mọi $x \in \mathbb{R}$. Khoảng nghịch biến của hàm số là

@@TAB@@\textbf{A.} $(-\infty; 1)$.@@TAB@@\textbf{B.} $(3; +\infty)$.@@TAB@@\textbf{C.} $(-\infty; 3)$.@@TAB@@\textbf{D.} $(1; 3)$.

\textbf{Lời giải.} Ta có $f'(x) < 0 \iff 1 < x < 3$, nên hàm số nghịch biến trên khoảng $(1; 3)$.

\textbf{==> Chọn đáp án D.}

% Câu 2 (BBT)
\textbf{Câu 2.} Cho hàm số $y = f(x)$ có bảng biến thiên như dưới đây.

@@CENTER_IMAGE_c2_bbt@@

Giá trị cực đại của hàm số là

@@TAB@@\textbf{A.} $0$.@@TAB@@\textbf{B.} $7$.@@TAB@@\textbf{C.} $4$.@@TAB@@\textbf{D.} $-1$.

\textbf{Lời giải.} Dựa vào bảng biến thiên, tại điểm $x = 0$ hàm số đạt cực đại và giá trị cực đại bằng $7$.

\textbf{==> Chọn đáp án B.}

% Câu 3
\textbf{Câu 3.} Cho hàm số $y = \dfrac{3x - 2}{x + 4}$. Phương trình tiệm cận đứng là

@@TAB@@\textbf{A.} $x = -4$.@@TAB@@\textbf{B.} $x = 4$.@@TAB@@\textbf{C.} $y = 3$.@@TAB@@\textbf{D.} $y = -4$.

\textbf{Lời giải.} Ta có $\lim_{x \to -4^+} \dfrac{3x - 2}{x + 4} = -\infty$ (mẫu số bằng $0$ tại $x = -4$), do đó đường tiệm cận đứng là $x = -4$.

\textbf{==> Chọn đáp án A.}

% Câu 4
\textbf{Câu 4.} Cho hàm số $y = \dfrac{-2x + 5}{x - 1}$. Phương trình tiệm cận ngang là

@@TAB@@\textbf{A.} $y = 2$.@@TAB@@\textbf{B.} $x = 1$.@@TAB@@\textbf{C.} $y = -2$.@@TAB@@\textbf{D.} $y = 5$.

\textbf{Lời giải.} Ta có $\lim_{x \to \pm\infty} \dfrac{-2x + 5}{x - 1} = -2$, do đó đường tiệm cận ngang là $y = -2$.

\textbf{==> Chọn đáp án C.}

% Câu 5
\textbf{Câu 5.} Cho hàm số $y = x - 1 + \dfrac{2}{x - 2}$. Phương trình đường tiệm cận xiên của đồ thị hàm số là

@@TAB@@\textbf{A.} $y = x + 1$.@@TAB@@\textbf{B.} $y = x - 1$.@@TAB@@\textbf{C.} $y = -x + 1$.@@TAB@@\textbf{D.} $x = 2$.

\textbf{Lời giải.} Vì $\lim_{x \to \pm\infty} [y - (x - 1)] = \lim_{x \to \pm\infty} \dfrac{2}{x - 2} = 0$, nên đường tiệm cận xiên của đồ thị là $y = x - 1$.

\textbf{==> Chọn đáp án B.}

% Câu 6
\textbf{Câu 6.} Đồ thị hàm số $y = -3 + \dfrac{5}{x + 2}$ có tâm đối xứng là

@@TAB@@\textbf{A.} $I(2; -3)$.@@TAB@@\textbf{B.} $I(-3; 2)$.@@TAB@@\textbf{C.} $I(-2; 3)$.@@TAB@@\textbf{D.} $I(-2; -3)$.

\textbf{Lời giải.} Đồ thị có tiệm cận đứng $x = -2$ và tiệm cận ngang $y = -3$. Giao điểm của hai tiệm cận là tâm đối xứng $I(-2; -3)$.

\textbf{==> Chọn đáp án D.}

% Câu 7
\textbf{Câu 7.} Cho hàm số $y = f(x)$ có đạo hàm $f'(x) = (x + 1)^2(x - 2)(x - 5)$ với mọi $x \in \mathbb{R}$. Số điểm cực trị của hàm số là

@@TAB@@\textbf{A.} $1$.@@TAB@@\textbf{B.} $3$.@@TAB@@\textbf{C.} $4$.@@TAB@@\textbf{D.} $2$.

\textbf{Lời giải.} Nghiệm $x = -1$ là nghiệm bội chẵn nên $f'(x)$ không đổi dấu qua $x = -1$. Đạo hàm $f'(x)$ đổi dấu khi qua hai nghiệm đơn $x = 2$ và $x = 5$. Vậy hàm số có $2$ điểm cực trị.

\textbf{==> Chọn đáp án D.}

% Câu 8
\textbf{Câu 8.} Cho hàm số $y = x + 1 + \dfrac{3}{x - 1}$. Giao điểm hai đường tiệm cận của đồ thị hàm số là

@@TAB@@\textbf{A.} $I(1; 1)$.@@TAB@@\textbf{B.} $I(1; 2)$.@@TAB@@\textbf{C.} $I(-1; 0)$.@@TAB@@\textbf{D.} $I(2; 1)$.

\textbf{Lời giải.} Tiệm cận đứng là $x = 1$, tiệm cận xiên là $y = x + 1$. Thay $x = 1$ vào phương trình tiệm cận xiên ta được $y = 2$. Giao điểm là $I(1; 2)$.

\textbf{==> Chọn đáp án B.}

% Câu 9 (Side-by-side)
@@START_SIDE_BY_SIDE_C9@@

\textbf{Câu 9.} Cho hàm số phân thức $y = \dfrac{ax + b}{cx + d}$ ($ac \ne 0$, $ad - bc \ne 0$) có đồ thị như hình bên. Hàm số nào dưới đây có đồ thị như hình đã cho?

@@TAB@@\textbf{A.} $y = 1 + \dfrac{2}{x + 1}$.@@TAB@@\textbf{B.} $y = 1 - \dfrac{2}{x + 1}$.

@@TAB@@\textbf{C.} $y = 1 + \dfrac{2}{x - 1}$.@@TAB@@\textbf{D.} $y = -1 + \dfrac{2}{x + 1}$.

\textbf{Lời giải.} Đồ thị có tiệm cận đứng $x = -1$, tiệm cận ngang $y = 1$. Với $x > -1$, đồ thị nằm phía trên đường tiệm cận ngang $y = 1$, do đó phần dư $\dfrac{2}{x + 1} > 0$. Vậy hàm số là $y = 1 + \dfrac{2}{x + 1}$.

\textbf{==> Chọn đáp án A.}

@@END_SIDE_BY_SIDE_C9@@

% Câu 10
\textbf{Câu 10.} Cho hàm số $y = f(x)$ xác định trên $\mathbb{R} \setminus \{0\}$ và có đạo hàm $f'(x) = \dfrac{x^2 - 1}{x^2}$. Số điểm cực trị của hàm số là

@@TAB@@\textbf{A.} $0$.@@TAB@@\textbf{B.} $1$.@@TAB@@\textbf{C.} $2$.@@TAB@@\textbf{D.} $3$.

\textbf{Lời giải.} Phương trình $f'(x) = 0 \iff x^2 - 1 = 0 \iff x = \pm 1$. Cả hai nghiệm đều thuộc tập xác định $\mathbb{R} \setminus \{0\}$ và là các nghiệm đơn làm $f'(x)$ đổi dấu. Vậy hàm số có $2$ điểm cực trị.

\textbf{==> Chọn đáp án C.}

% Câu 11
\textbf{Câu 11.} Cho hàm số $y = f(x) = x^3 - 12x + 1$. Khẳng định nào sau đây đúng?

@@TAB@@\textbf{A.} Hàm số đạt cực trị tại $x = -2$ và $x = 2$.

@@TAB@@\textbf{B.} Hàm số không có cực trị.

@@TAB@@\textbf{C.} Hàm số có một điểm cực trị.

@@TAB@@\textbf{D.} Tâm đối xứng của đồ thị là $I(1; 0)$.

\textbf{Lời giải.} Ta có $f'(x) = 3x^2 - 12 = 3(x^2 - 4) = 0 \iff x = \pm 2$. Do đó hàm số đạt cực trị tại $x = -2$ và $x = 2$ (tâm đối xứng của đồ thị là $I(0; 1)$).

\textbf{==> Chọn đáp án A.}

% Câu 12 (BBT)
\textbf{Câu 12.} Cho hàm số bậc ba $y = f(x)$ có bảng biến thiên như dưới đây:

@@CENTER_IMAGE_c12_bbt@@

Hàm số nào dưới đây có bảng biến thiên đã cho?

@@TAB@@\textbf{A.} $y = -x^3 + 3x$.@@TAB@@\textbf{B.} $y = x^3 + 3x$.@@TAB@@\textbf{C.} $y = x^3 - 3x$.@@TAB@@\textbf{D.} $y = x^3 - 3x + 1$.

\textbf{Lời giải.} Đồ thị đi lên khi $x \to +\infty \implies a > 0$. Hàm số đạt cực đại tại $x = -1$ với $y_{CĐ} = 2$, đạt cực tiểu tại $x = 1$ với $y_{CT} = -2$. Xét hàm số $y = x^3 - 3x$: $y' = 3x^2 - 3 = 0 \iff x = \pm 1$; $y(-1) = 2$ và $y(1) = -2$, hoàn toàn phù hợp.

\textbf{==> Chọn đáp án C.}

@@SECTION_2_HEADER@@

% Phần 2 Câu 1 (BBT)
\textbf{Câu 1.} Cho hàm số $y = \dfrac{(x + 1)^2}{x + 2}$. Xét tính đúng, sai của các phát biểu sau:

@@TAB@@\textbf{a)} Đạo hàm là $y' = \dfrac{(x + 3)(x + 1)}{(x + 2)^2}$.

@@TAB@@\textbf{b)} Hàm số nghịch biến trên từng khoảng $(-3; -2)$ và $(-2; -1)$.

@@TAB@@\textbf{c)} Giá trị cực đại của hàm số bằng $0$.

@@TAB@@\textbf{d)} Bảng biến thiên dưới đây là bảng biến thiên của hàm số đã cho.

@@CENTER_IMAGE_p2_c1_bbt@@

\textbf{Lời giải.} Biến đổi: $y = \dfrac{x^2 + 2x + 1}{x + 2} = x + \dfrac{1}{x + 2} \implies y' = 1 - \dfrac{1}{(x + 2)^2} = \dfrac{x^2 + 4x + 3}{(x + 2)^2} = \dfrac{(x + 3)(x + 1)}{(x + 2)^2}$.
- Ý a) Đúng.
- Ý b) Đúng vì $y' < 0$ với mọi $x \in (-3; -2)$ và $x \in (-2; -1)$.
- Ý c) Sai vì hàm số đạt cực đại tại $x = -3$, giá trị cực đại bằng $y(-3) = -4$ (tại $x = -1$ hàm số đạt cực tiểu với $y_{CT} = 0$).
- Ý d) Đúng vì bảng biến thiên hoàn toàn trùng khớp với các giới hạn và điểm cực trị.

\textbf{==> Chọn đáp án: a) Đúng  |  b) Đúng  |  c) Sai  |  d) Đúng.}

% Phần 2 Câu 2
\textbf{Câu 2.} Cho hàm số $y = f(x) = x^3 - 6x^2 + 9x - 2$. Xét tính đúng, sai của các phát biểu sau:

@@TAB@@\textbf{a)} Tập xác định là $\mathbb{R}$.

@@TAB@@\textbf{b)} Hàm số đồng biến trên khoảng $(1; 3)$.

@@TAB@@\textbf{c)} Giá trị cực tiểu của hàm số bằng $-2$.

@@TAB@@\textbf{d)} Điểm $I(2; 1)$ là tâm đối xứng của đồ thị.

\textbf{Lời giải.} Ta có $f'(x) = 3x^2 - 12x + 9 = 3(x - 1)(x - 3)$.
- Ý a) Đúng vì hàm số đa thức xác định trên $\mathbb{R}$.
- Ý b) Sai vì trên khoảng $(1; 3)$ thì $f'(x) < 0$ nên hàm số nghịch biến.
- Ý c) Đúng vì hàm số đạt cực tiểu tại $x = 3$ với giá trị cực tiểu $f(3) = 27 - 54 + 27 - 2 = -2$.
- Ý d) Sai vì hoành độ tâm đối xứng là nghiệm của $f''(x) = 6x - 12 = 0 \iff x = 2$; tung độ $y(2) = 8 - 24 + 18 - 2 = 0$. Tâm đối xứng là $I(2; 0)$, không phải $I(2; 1)$.

\textbf{==> Chọn đáp án: a) Đúng  |  b) Sai  |  c) Đúng  |  d) Sai.}

% Phần 2 Câu 3 (Side-by-side)
@@START_SIDE_BY_SIDE_P2C3@@

\textbf{Câu 3.} Cho hàm số phân thức $y = \dfrac{ax + b}{cx + d}$ ($ac \ne 0$, $ad - bc \ne 0$) có đồ thị như hình bên. Xét tính đúng, sai của các phát biểu sau:

@@TAB@@\textbf{a)} Đồ thị có tiệm cận đứng $x = -2$.

@@TAB@@\textbf{b)} Đồ thị có tiệm cận ngang $y = 1$.

@@TAB@@\textbf{c)} Tâm đối xứng của đồ thị là $I(-2; -1)$.

@@TAB@@\textbf{d)} Hàm số đồng biến trên từng khoảng xác định.

\textbf{Lời giải.} Quan sát đồ thị:
- Đường tiệm cận đứng là $x = -2 \implies$ a) Đúng.
- Đường tiệm cận ngang là đường $y = -1$, không phải $y = 1 \implies$ b) Sai.
- Giao điểm hai đường tiệm cận là $I(-2; -1)$, là tâm đối xứng của đồ thị $\implies$ c) Đúng.
- Đồ thị đi xuống từ trái sang phải trên từng khoảng $(-\infty; -2)$ và $(-2; +\infty)$ nên hàm số nghịch biến trên từng khoảng xác định $\implies$ d) Sai.

\textbf{==> Chọn đáp án: a) Đúng  |  b) Sai  |  c) Đúng  |  d) Sai.}

@@END_SIDE_BY_SIDE_P2C3@@

% Phần 2 Câu 4 (Side-by-side)
@@START_SIDE_BY_SIDE_P2C4@@

\textbf{Câu 4.} Cho hàm số $y = \dfrac{-2x + 1}{x + 1}$. Xét tính đúng, sai của các phát biểu sau:

@@TAB@@\textbf{a)} Đạo hàm là $y' = -\dfrac{3}{(x + 1)^2}$.

@@TAB@@\textbf{b)} Hàm số nghịch biến trên từng khoảng xác định.

@@TAB@@\textbf{c)} Đồ thị có tiệm cận ngang $y = 2$.

@@TAB@@\textbf{d)} Hình bên là đồ thị của hàm số đã cho.

\textbf{Lời giải.} Ta có $y = \dfrac{-2(x + 1) + 3}{x + 1} = -2 + \dfrac{3}{x + 1}$.
- Đạo hàm $y' = \dfrac{(-2)\cdot 1 - 1\cdot 1}{(x + 1)^2} = -\dfrac{3}{(x + 1)^2} \implies$ a) Đúng.
- Vì $y' < 0$ với mọi $x \ne -1$ nên hàm số nghịch biến trên từng khoảng xác định $(-\infty; -1)$ và $(-1; +\infty) \implies$ b) Đúng.
- Tiệm cận ngang là $y = \lim_{x \to \pm\infty} \dfrac{-2x + 1}{x + 1} = -2 \implies$ c) Sai.
- Đồ thị có tiệm cận đứng $x = -1$, tiệm cận ngang $y = -2$, cắt trục tung tại $(0; 1)$ và trục hoành tại $(1/2; 0)$, hoàn toàn trùng khớp với hình vẽ $\implies$ d) Đúng.

\textbf{==> Chọn đáp án: a) Đúng  |  b) Đúng  |  c) Sai  |  d) Đúng.}

@@END_SIDE_BY_SIDE_P2C4@@

@@SECTION_3_HEADER@@

% Phần 3 Câu 1
\textbf{Câu 1.} Cho hàm số $y = f(x) = -x^3 + 3x + 1$. Tính tổng hai giá trị cực trị của hàm số.

\textbf{Lời giải.} Ta có $f'(x) = -3x^2 + 3 = 0 \iff x = \pm 1$.
- Tại $x = 1$, giá trị cực đại là $y(1) = -1 + 3 + 1 = 3$.
- Tại $x = -1$, giá trị cực tiểu là $y(-1) = 1 - 3 + 1 = -1$.
Tổng hai giá trị cực trị là: $3 + (-1) = 2$.

\textbf{==> Đáp án: 2.}

% Phần 3 Câu 2
\textbf{Câu 2.} Cho hàm số $y = \dfrac{x^2 - x + 2}{x - 2}$. Gọi $I(a; b)$ là giao điểm hai đường tiệm cận của đồ thị. Tính $a + b$.

\textbf{Lời giải.} Chia đa thức: $y = \dfrac{x(x - 2) + (x - 2) + 4}{x - 2} = x + 1 + \dfrac{4}{x - 2}$.
- Tiệm cận đứng: $x = 2 \implies a = 2$.
- Tiệm cận xiên: $y = x + 1$. Thay $x = 2 \implies b = 2 + 1 = 3$.
Giao điểm hai đường tiệm cận là $I(2; 3)$. Vậy $a + b = 2 + 3 = 5$.

\textbf{==> Đáp án: 5.}

% Phần 3 Câu 3
\textbf{Câu 3.} Cho hàm số $y = f(x)$ có đạo hàm $f'(x) = (x + 2)^2(x - 1)(x - 3)$ với mọi $x \in \mathbb{R}$. Tính tổng các giá trị của $x$ tại đó hàm số đạt cực trị.

\textbf{Lời giải.} Nghiệm $x = -2$ là nghiệm bội chẵn nên không phải điểm cực trị. Chỉ có hai nghiệm đơn $x = 1$ và $x = 3$ làm đạo hàm đổi dấu khi qua chúng, nên hàm số đạt cực trị tại $x = 1$ và $x = 3$.
Tổng các giá trị của $x$ tại đó hàm số đạt cực trị là: $1 + 3 = 4$.

\textbf{==> Đáp án: 4.}

% Phần 3 Câu 4
\textbf{Câu 4.} Cho hàm số $y = \dfrac{x^2 - 2x + 5}{x - 1}$. Gọi $A, B$ là hai điểm cực trị của đồ thị hàm số. Tính độ dài đoạn thẳng $AB$ (kết quả làm tròn đến hàng phần trăm).

\textbf{Lời giải.} Ta có $y = x - 1 + \dfrac{4}{x - 1}$.
Đạo hàm: $y' = 1 - \dfrac{4}{(x - 1)^2} = \dfrac{(x - 1)^2 - 4}{(x - 1)^2} = \dfrac{x^2 - 2x - 3}{(x - 1)^2} = \dfrac{(x + 1)(x - 3)}{(x - 1)^2}$.
Phương trình $y' = 0 \iff x = -1$ hoặc $x = 3$.
- Với $x = -1 \implies y = -1 - 1 + \dfrac{4}{-2} = -4 \implies A(-1; -4)$.
- Với $x = 3 \implies y = 3 - 1 + \dfrac{4}{2} = 4 \implies B(3; 4)$.
Độ dài đoạn thẳng $AB$ là:
\[
AB = \sqrt{(3 - (-1))^2 + (4 - (-4))^2} = \sqrt{4^2 + 8^2} = \sqrt{16 + 64} = \sqrt{80} = 4\sqrt{5} \approx 8{,}944 \approx 8{,}94.
\]

\textbf{==> Đáp án: 8,94.}

% Phần 3 Câu 5
\textbf{Câu 5.} Cho hàm số $y = f(x) = \dfrac{x^2 - x + 4}{x}$. Tính hiệu giữa giá trị cực tiểu và giá trị cực đại của hàm số.

\textbf{Lời giải.} Biến đổi: $f(x) = x - 1 + \dfrac{4}{x}$.
Đạo hàm: $f'(x) = 1 - \dfrac{4}{x^2} = \dfrac{x^2 - 4}{x^2} = 0 \iff x = \pm 2$.
- Tại $x = -2$, hàm số đạt cực đại và giá trị cực đại là $y_{CĐ} = f(-2) = -2 - 1 - 2 = -5$.
- Tại $x = 2$, hàm số đạt cực tiểu và giá trị cực tiểu là $y_{CT} = f(2) = 2 - 1 + 2 = 3$.
Hiệu giữa giá trị cực tiểu và giá trị cực đại là:
\[
y_{CT} - y_{CĐ} = 3 - (-5) = 8.
\]

\textbf{==> Đáp án: 8.}

% Phần 3 Câu 6 (BBT)
\textbf{Câu 6.} Cho hàm số $y = \dfrac{ax + b}{cx + d}$ ($ac \ne 0$, $ad - bc \ne 0$) có bảng biến thiên như dưới đây:

@@CENTER_IMAGE_p3_c6_bbt@@

Gọi $I(u; v)$ là tâm đối xứng của đồ thị hàm số. Tính $u + v$.

\textbf{Lời giải.} Từ bảng biến thiên:
- Tiệm cận đứng: $x = -2 \implies u = -2$.
- Tiệm cận ngang: $y = -1 \implies v = -1$.
Tâm đối xứng của đồ thị hàm phân thức bậc nhất trên bậc nhất là giao điểm của hai đường tiệm cận, tức $I(-2; -1)$.
Do đó:
\[
u + v = -2 + (-1) = -3.
\]

\textbf{==> Đáp án: -3.}

\vspace{0.4cm}
\begin{center}
\textbf{------ HẾT ------}
\end{center}

\end{document}
"""

def build_0104_doc(tex_content, output_path, is_solution=False):
    temp_tex = os.path.join(CURRENT_DIR, f"temp_{'hdg' if is_solution else 'de'}_0104.tex")
    temp_docx = os.path.join(CURRENT_DIR, f"temp_{'hdg' if is_solution else 'de'}_0104.docx")
    
    with open(temp_tex, "w", encoding="utf-8") as f:
        f.write(tex_content)
        
    cmd = [pandoc_exe, temp_tex, "-o", temp_docx, "--from=latex", "--to=docx"]
    print(f"[*] Đang xuất bản: {os.path.basename(output_path)}...")
    res = subprocess.run(cmd, capture_output=True, text=True, encoding="utf-8")
    if res.returncode != 0:
        print("[LỖI PANDOC]:", res.stderr)
        return
        
    doc = docx.Document(temp_docx)
    
    # Thiết lập lề trang
    for s in doc.sections:
        s.top_margin = Cm(1.6)
        s.bottom_margin = Cm(1.6)
        s.left_margin = Cm(2.0)
        s.right_margin = Cm(1.5)
        s.page_width = Cm(21.0)
        s.page_height = Cm(29.7)
        
        # Footer
        footer = s.footer
        f_p = footer.paragraphs[0]
        f_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        f_p.text = ""
        r_l = f_p.add_run("Lớp Toán Cô Thúy  •  SĐT: 0935.322.328  •  50/2C Phạm Thị Liên                                   ")
        r_l.font.name = "Times New Roman"
        r_l.font.size = Pt(9.5)
        r_l.italic = True
        r_l.font.color.rgb = RGBColor(100, 100, 100)
        
        r_r = f_p.add_run("Trang Mã đề 0104")
        r_r.font.name = "Times New Roman"
        r_r.font.size = Pt(9.5)
        r_r.italic = True
        r_r.font.color.rgb = RGBColor(80, 80, 80)
        
    style = doc.styles['Normal']
    style.font.name = 'Times New Roman'
    style.font.size = Pt(12)
    style.font.color.rgb = RGBColor(0, 0, 0)
    rPr = style.element.get_or_add_rPr()
    rFonts = parse_xml(r'<w:rFonts %s w:ascii="Times New Roman" w:hAnsi="Times New Roman" w:cs="Times New Roman"/>' % nsdecls('w'))
    rPr.append(rFonts)
    
    # Duyệt và thay thế Header sections
    for p in list(doc.paragraphs):
        if "@@DOCUMENT_HEADER@@" in p.text:
            insert_header_and_code(doc, p, "0104", is_solution)
        elif "@@SECTION_1_HEADER@@" in p.text:
            insert_section_header(doc, p, "PHẦN I. Câu trắc nghiệm nhiều phương án lựa chọn.", "Thí sinh trả lời từ câu 1 đến câu 12. Mỗi câu hỏi thí sinh chỉ chọn một phương án.")
        elif "@@SECTION_2_HEADER@@" in p.text:
            insert_section_header(doc, p, "PHẦN II. Câu trắc nghiệm đúng sai.", "Thí sinh trả lời từ câu 1 đến câu 4. Trong mỗi ý a), b), c), d) ở mỗi câu, thí sinh chọn đúng hoặc sai.")
        elif "@@SECTION_3_HEADER@@" in p.text:
            insert_section_header(doc, p, "PHẦN III. Câu trắc nghiệm yêu cầu trả lời ngắn.", "Thí sinh trả lời từ câu 1 đến câu 6.")
            
    # Xử lý các khối câu hỏi Side-by-side
    sbs_configs = [
        ("C9", "c9_hinh.png", 4.8),
        ("P2C3", "p2_c3_hinh.png", 4.8),
        ("P2C4", "p2_c4_hinh.png", 4.8),
    ]
    for tag, img_name, img_w in sbs_configs:
        start_tag = f"@@START_SIDE_BY_SIDE_{tag}@@"
        end_tag = f"@@END_SIDE_BY_SIDE_{tag}@@"
        start_p, end_p, inner_ps = None, None, []
        found = False
        for p in doc.paragraphs:
            if start_tag in p.text:
                start_p = p; found = True; continue
            if end_tag in p.text:
                end_p = p; found = False; continue
            if found:
                inner_ps.append(p)
                
        if start_p and end_p:
            tbl = doc.add_table(rows=1, cols=2)
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            
            tblPr = tbl._tbl.tblPr
            for look in tblPr.findall(qn('w:tblLook')): tblPr.remove(look)
            tblLook = parse_xml(r'<w:tblLook %s w:val="0000" w:firstRow="0" w:lastRow="0" w:firstColumn="0" w:lastColumn="0" w:noHBand="0" w:noVBand="0"/>' % nsdecls('w'))
            for b in tblPr.findall(qn('w:tblBorders')): tblPr.remove(b)
            tblBorders = parse_xml(
                r'<w:tblBorders %s>'
                r'<w:top w:val="none"/><w:left w:val="none"/><w:bottom w:val="none"/><w:right w:val="none"/>'
                r'<w:insideH w:val="none"/><w:insideV w:val="none"/>'
                r'</w:tblBorders>' % nsdecls('w')
            )
            tblPr.append(tblBorders)
            tblPr.append(tblLook)
            
            cell_0 = tbl.rows[0].cells[0]
            cell_1 = tbl.rows[0].cells[1]
            cell_0.width = Cm(11.8)
            cell_1.width = Cm(5.7)
            cell_0.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            cell_1.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            
            for c in [cell_0, cell_1]:
                tcPr = c._tc.get_or_add_tcPr()
                for b in tcPr.findall(qn('w:tcBorders')): tcPr.remove(b)
                tcPr.append(parse_xml(
                    r'<w:tcBorders %s><w:top w:val="none"/><w:left w:val="none"/><w:bottom w:val="none"/><w:right w:val="none"/></w:tcBorders>' % nsdecls('w')
                ))
                
            start_p._p.addprevious(tbl._tbl)
            cell_0._tc.remove(cell_0.paragraphs[0]._p)
            for p in inner_ps:
                cell_0._tc.append(p._p)
                
            img_path = os.path.join(HINH_ANH_0104, img_name)
            cell_1_p = cell_1.paragraphs[0]
            cell_1_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cell_1_p.paragraph_format.space_before = Pt(0)
            cell_1_p.paragraph_format.space_after = Pt(0)
            if os.path.exists(img_path):
                cell_1_p.add_run().add_picture(img_path, width=Cm(img_w))
                
            start_p._p.getparent().remove(start_p._p)
            end_p._p.getparent().remove(end_p._p)
            
    # Duyệt các ảnh căn giữa & ô trả lời Phần III
    for p in list(doc.paragraphs):
        for img_tag in ["c2_bbt", "c12_bbt", "p2_c1_bbt", "p3_c6_bbt"]:
            token = f"@@CENTER_IMAGE_{img_tag}@@"
            if token in p.text:
                img_path = os.path.join(HINH_ANH_0104, f"{img_tag}.png")
                p.text = ""
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.space_after = Pt(6)
                if os.path.exists(img_path):
                    p.add_run().add_picture(img_path, width=Cm(9.2))
                break
                
        if not is_solution:
            for num in range(1, 7):
                tag = f"@@ANSWER_BOX_{num}@@"
                if tag in p.text:
                    box_tbl = create_answer_box_table(doc)
                    p._p.addprevious(box_tbl._tbl)
                    p._p.getparent().remove(p._p)
                    break
                    
    # Định dạng văn bản và Tab stops BTPro
    for p in list(doc.paragraphs):
        format_doc_paragraph(p, inside_table_cell=False)
        if is_solution:
            p_text = p.text.strip()
            if "==> Chọn đáp án" in p_text or "==> Đáp án:" in p_text:
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(8)
                for r in p.runs:
                    r.font.bold = True
                    r.font.color.rgb = RGBColor(192, 0, 0)
                    
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in list(cell.paragraphs):
                    format_doc_paragraph(p, inside_table_cell=True)
                    if is_solution:
                        p_text = p.text.strip()
                        if "==> Chọn đáp án" in p_text or "==> Đáp án:" in p_text:
                            for r in p.runs:
                                r.font.bold = True
                                r.font.color.rgb = RGBColor(192, 0, 0)
                                
    # Loại bỏ tblStylePr gây gạch ngang dưới
    for s in doc.styles:
        if s.type == docx.enum.style.WD_STYLE_TYPE.TABLE:
            for child in list(s._element):
                if child.tag.endswith('tblStylePr'):
                    s._element.remove(child)
                    
    # Chế độ tương thích Word 2013+
    settings = doc.settings.element
    compat = parse_xml(
        r'<w:compat %s><w:compatSetting w:name="compatibilityMode" w:uri="http://schemas.microsoft.com/office/word" w:val="15"/></w:compat>' % nsdecls('w')
    )
    settings.append(compat)
    
    doc.save(output_path)
    print(f"[THÀNH CÔNG] Đã xuất bản: {output_path}")
    if os.path.exists(temp_tex): os.remove(temp_tex)
    if os.path.exists(temp_docx): os.remove(temp_docx)

if __name__ == "__main__":
    print("=== TIẾN HÀNH XUẤT BẢN MÃ ĐỀ 0104 (ĐỀ BÀI + LỜI GIẢI CHI TIẾT) ===")
    build_0104_doc(tex_de_0104, OUTPUT_DE_0104, is_solution=False)
    build_0104_doc(tex_hdg_0104, OUTPUT_HDG_0104, is_solution=True)
